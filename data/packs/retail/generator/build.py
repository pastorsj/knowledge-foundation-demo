# SPDX-FileCopyrightText: Copyright (c) 2026, NVIDIA CORPORATION & AFFILIATES. All rights reserved.
# SPDX-License-Identifier: Apache-2.0
#
# /// script
# requires-python = ">=3.12"
# dependencies = [
#     "numpy",
#     "pandas",
#     "openpyxl",
#     "reportlab",
#     "python-docx",
#     "python-pptx",
#     "pillow",
#     "pypdfium2",
# ]
# ///
"""Build the Lumen Retail Group industry pack.

    uv run data/packs/retail/generator/build.py

Writes, under data/packs/retail/files/, the CSV and XLSX tables, the PDF/DOCX/PPTX documents and the scanned PNG
memo, and regenerates ../README.md from content/readme.md. The run is seeded and deterministic: the same inputs give
identical bytes. Every figure quoted in a document or in the README answer key is computed here from the generated
tables, so documents and tables agree.
"""

from __future__ import annotations

import sys
from datetime import date
from datetime import datetime
from pathlib import Path

import pandas as pd

HERE = Path(__file__).resolve().parent
PACK = HERE.parent
FILES = PACK / "files"
CONTENT = HERE / "content"
sys.path.insert(0, str(HERE))

import datagen  # noqa: E402
import render  # noqa: E402

Q3S, Q4S = pd.Timestamp("2026-07-01"), pd.Timestamp("2026-10-01")
Q3PS, Q4PS = pd.Timestamp("2025-07-01"), pd.Timestamp("2025-10-01")
H1S, H1E = pd.Timestamp("2026-01-01"), pd.Timestamp("2026-07-01")
SCAN_SEED = 20260812


def usd(v: float, dec: int = 0) -> str:
    return f"${v:,.{dec}f}"


def pct(v: float, dec: int = 1, sign: bool = False) -> str:
    return f"{v * 100:+.{dec}f}%" if sign else f"{v * 100:.{dec}f}%"


def short(name: str) -> str:
    return name.removeprefix("Lumen ")


def join_and(items: list[str]) -> str:
    if len(items) <= 1:
        return "".join(items)
    return ", ".join(items[:-1]) + " and " + items[-1]


def between(s: pd.Series, lo: pd.Timestamp, hi: pd.Timestamp) -> pd.Series:
    return (s >= lo) & (s < hi)


def md_date(d: pd.Timestamp | object) -> str:
    return f"{d:%b} {d.day}"


# --------------------------------------------------------------------------- tables to disk


def write_csvs(T: dict[str, pd.DataFrame]) -> None:
    out = FILES / "sales"
    out.mkdir(parents=True, exist_ok=True)
    stores = T["stores"].copy()
    stores["square_feet"] = stores["square_feet"].astype("Int64")
    customers = T["customers"].copy()
    customers["email_opt_in"] = customers["email_opt_in"].map({True: "true", False: "false"})
    for name, df in (
        ("stores", stores),
        ("products", T["products"]),
        ("customers", customers),
        ("orders", T["orders"]),
        ("order_items", T["order_items"]),
        ("returns", T["returns"]),
    ):
        df.to_csv(out / f"{name}.csv", index=False, lineterminator="\n", date_format="%Y-%m-%d %H:%M:%S")


def write_xlsx(T: dict[str, pd.DataFrame]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment
    from openpyxl.styles import Font
    from openpyxl.styles import PatternFill

    wb = Workbook()
    sheets = [
        (
            "category_targets",
            T["category_targets"],
            {
                "target_net_sales": "#,##0",
                "target_gross_margin_pct": "0.0%",
                "target_units": "#,##0",
                "quarter_start": "yyyy-mm-dd",
            },
        ),
        (
            "promo_calendar",
            T["promo_calendar"],
            {"start_date": "yyyy-mm-dd", "end_date": "yyyy-mm-dd", "discount_pct": "0%", "budget_usd": "#,##0"},
        ),
    ]
    for i, (name, df, fmts) in enumerate(sheets):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = name
        ws.append(list(df.columns))
        for rec in df.itertuples(index=False):
            ws.append([v.date() if isinstance(v, pd.Timestamp) else v for v in rec])
        for c, col in enumerate(df.columns, start=1):
            head = ws.cell(row=1, column=c)
            head.font = Font(bold=True, color="FFFFFF")
            head.fill = PatternFill("solid", fgColor="1B2A49")
            head.alignment = Alignment(horizontal="center")
            width = max(len(str(col)), *(len(str(v)) for v in df[col])) + 3
            ws.column_dimensions[head.column_letter].width = min(width, 34)
            if col in fmts:
                for r in range(2, len(df) + 2):
                    ws.cell(row=r, column=c).number_format = fmts[col]
        ws.freeze_panes = "A2"
    wb.properties.creator = "Lumen Retail Group - Merchandising and Planning"
    wb.properties.lastModifiedBy = "Lumen Retail Group - Merchandising and Planning"
    wb.properties.title = "Merchandising plan, fiscal 2026"
    wb.properties.created = render.FIXED_TIME
    wb.properties.modified = render.FIXED_TIME
    path = FILES / "sales" / "merchandising_plan.xlsx"
    wb.save(path)
    render.normalize_zip(path)


# --------------------------------------------------------------------------- analysis


def analyze(T: dict[str, pd.DataFrame]) -> tuple[dict[str, str], dict[str, dict], dict[str, dict], dict[str, list]]:
    stores, products, customers = T["stores"], T["products"], T["customers"]
    orders, items, returns = T["orders"].copy(), T["order_items"].copy(), T["returns"].copy()
    targets, promos = T["category_targets"], T["promo_calendar"]
    smap = stores.set_index("store_id")
    pmap = products.set_index("product_id")
    if abs(items["line_amount"].sum() - orders["net_amount"].sum()) > 0.5:
        raise AssertionError("order lines do not add up to order net amounts")
    orders["region"] = orders["store_id"].map(smap["region"])
    orders["store_type"] = orders["store_id"].map(smap["store_type"])
    returns["region"] = returns["store_id"].map(smap["region"])
    returns["store_type"] = returns["store_id"].map(smap["store_type"])
    returns["category"] = returns["product_id"].map(pmap["category"])
    items = items.merge(orders[["order_id", "ordered_at", "store_id", "region"]], on="order_id")
    items["category"] = items["product_id"].map(pmap["category"])
    items["supplier"] = items["product_id"].map(pmap["supplier"])

    f: dict[str, str] = {}
    a: dict[str, str] = {}
    tables: dict[str, dict] = {}
    charts: dict[str, dict] = {}
    kpis: dict[str, list] = {}
    regions = [*datagen.REGIONS, "E-commerce"]

    o26 = orders[between(orders["ordered_at"], Q3S, Q4S)]
    o25 = orders[between(orders["ordered_at"], Q3PS, Q4PS)]
    r26 = returns[between(returns["returned_at"], Q3S, Q4S)]
    r25 = returns[between(returns["returned_at"], Q3PS, Q4PS)]
    i26 = items[between(items["ordered_at"], Q3S, Q4S)]

    # ---- chain level
    net26, net25 = o26["net_amount"].sum(), o25["net_amount"].sum()
    f["q3_net_sales"] = usd(net26)
    f["q3_orders"] = f"{len(o26):,}"
    f["q3_yoy"] = pct(net26 / net25 - 1, 1, True)
    f["q3_aov"] = usd(net26 / len(o26))
    f["q3_returned"] = usd(r26["returned_amount"].sum())
    f["q3_returned_py"] = usd(r25["returned_amount"].sum())
    f["q3_return_count"] = f"{len(r26):,}"
    chain_rate = r26["returned_amount"].sum() / net26
    f["q3_return_rate"] = pct(chain_rate)

    # ---- regions
    reg = {}
    for name in regions:
        s26 = o26[o26["region"] == name]["net_amount"].sum()
        s25 = o25[o25["region"] == name]["net_amount"].sum()
        n = int((o26["region"] == name).sum())
        rv = r26[r26["region"] == name]["returned_amount"].sum()
        reg[name] = {"s26": s26, "s25": s25, "yoy": s26 / s25 - 1, "n": n, "ret": rv, "rate": rv / s26}
    n_stores = stores[stores["store_type"] != "ecommerce"].groupby("region").size().to_dict()
    rows = []
    for name in regions:
        g = reg[name]
        rows.append(
            [
                name,
                str(n_stores.get(name, "Online")),
                usd(g["s26"]),
                usd(g["s25"]),
                pct(g["yoy"], 1, True),
                f"{g['n']:,}",
                usd(g["s26"] / g["n"]),
                usd(g["ret"]),
                pct(g["rate"]),
            ]
        )
    rows.append(
        [
            "Total",
            "24",
            usd(net26),
            usd(net25),
            pct(net26 / net25 - 1, 1, True),
            f"{len(o26):,}",
            usd(net26 / len(o26)),
            usd(r26["returned_amount"].sum()),
            pct(chain_rate),
        ]
    )
    tables["region_summary"] = {
        "caption": "Table 1. Q3 2026 results by region",
        "header": [
            "Region",
            "Stores",
            "Q3 2026 net sales",
            "Q3 2025 net sales",
            "YoY",
            "Orders",
            "Avg order value",
            "Returned value",
            "Return rate",
        ],
        "rows": rows,
        "align": ["L", "R", "R", "R", "R", "R", "R", "R", "R"],
        "widths": [0.14, 0.07, 0.13, 0.13, 0.09, 0.08, 0.1, 0.13, 0.13],
        "bold_last": True,
        "note": "Net sales are after discounts and before returns. Return rate is returned value divided by net sales.",
    }
    charts["region_sales"] = {
        "title": "Net sales by region, $ thousands",
        "labels": regions,
        "series": [
            ("Q3 2025", [round(reg[r]["s25"] / 1000, 1) for r in regions]),
            ("Q3 2026", [round(reg[r]["s26"] / 1000, 1) for r in regions]),
        ],
        "label_fmt": lambda v: f"{v:,.1f}",
        "axis_fmt": "%d",
    }
    by_yoy = sorted(regions, key=lambda r: reg[r]["yoy"], reverse=True)
    top, second, third, low = by_yoy[0], by_yoy[1], by_yoy[2], by_yoy[-1]
    if top == "E-commerce" or reg[top]["yoy"] - reg[second]["yoy"] < 0.045:
        raise AssertionError("the story needs one geographic region clearly ahead on year-over-year growth")
    for key, name in (("top", top), ("second", second), ("third", third), ("low", low)):
        f[f"{key}_region"] = name
        f[f"{key}_region_yoy"] = pct(reg[name]["yoy"], 1, True)
    f["top_region_sales"] = usd(reg[top]["s26"])
    largest = max(regions, key=lambda r: reg[r]["s26"])
    f["largest_region"] = largest
    f["largest_region_sales"] = usd(reg[largest]["s26"])
    f["largest_region_share"] = pct(reg[largest]["s26"] / net26)
    f["ecom_share"] = pct(reg["E-commerce"]["s26"] / net26)
    hr = max(regions, key=lambda r: reg[r]["rate"])
    f["hr_region"], f["hr_region_rate"] = hr, pct(reg[hr]["rate"])

    # ---- physical stores
    phys = o26[o26["store_type"] != "ecommerce"]
    st_sales = phys.groupby("store_id")["net_amount"].sum().sort_values(ascending=False)
    st_orders = phys.groupby("store_id").size()
    st_ret = r26.groupby("store_id")["returned_amount"].sum()
    top5 = st_sales.index[:5].tolist()
    if any(st_sales.iloc[i] < st_sales.iloc[i + 1] * 1.04 for i in range(5)):
        raise AssertionError("two of the top six stores are within 4 percent of each other")
    rows = []
    for rank, sid in enumerate(top5, start=1):
        rows.append(
            [
                str(rank),
                sid,
                short(smap.at[sid, "store_name"]),
                smap.at[sid, "region"],
                usd(st_sales[sid]),
                f"{st_orders[sid]:,}",
                pct(st_ret.get(sid, 0.0) / st_sales[sid]),
            ]
        )
    tables["top_stores"] = {
        "caption": "Table 2. Five physical stores with the highest Q3 2026 net sales",
        "header": ["Rank", "Store ID", "Store", "Region", "Net sales", "Orders", "Return rate"],
        "rows": rows,
        "align": ["R", "L", "L", "L", "R", "R", "R"],
        "widths": [0.07, 0.1, 0.29, 0.13, 0.15, 0.11, 0.15],
    }
    f["top_store_id"] = top5[0]
    f["top_store_name"] = short(smap.at[top5[0], "store_name"])
    f["top_store_sales"] = usd(st_sales[top5[0]])
    flag = smap.index[smap["store_type"] == "flagship"]
    f["flagship_share"] = pct(st_sales[st_sales.index.isin(flag)].sum() / st_sales.sum())
    f["s24_sales"] = usd(st_sales.get("S24", 0.0))
    f["s24_sales_py"] = usd(o25[o25["store_id"] == "S24"]["net_amount"].sum())
    for kind in ("flagship", "mall", "standalone", "outlet"):
        f[f"stores_{kind}"] = ", ".join(sorted(smap.index[smap["store_type"] == kind]))
    a["top5"] = "; ".join(
        f"({i}) {s} {short(smap.at[s, 'store_name'])} {usd(st_sales[s])}" for i, s in enumerate(top5, 1)
    )
    a["web_rank_note"] = usd(o26[o26["store_id"] == "WEB"]["net_amount"].sum())

    # ---- categories against plan
    cat_sales = i26.groupby("category")["line_amount"].sum()
    cat_ret = r26.groupby("category")["returned_amount"].sum()
    plan = targets[targets["quarter"] == "2026-Q3"].set_index("category")["target_net_sales"]
    cat_rate = (cat_ret / cat_sales).sort_values(ascending=False)
    var = cat_sales / plan - 1
    rows = []
    for cat in datagen.CATEGORIES:
        rows.append([cat, usd(cat_sales[cat]), usd(plan[cat]), pct(var[cat], 1, True), pct(cat_rate[cat])])
    rows.append(
        [
            "Total",
            usd(cat_sales.sum()),
            usd(plan.sum()),
            pct(cat_sales.sum() / plan.sum() - 1, 1, True),
            pct(chain_rate),
        ]
    )
    tables["category_summary"] = {
        "caption": "Table 3. Q3 2026 net sales against plan, and return rate, by category",
        "header": ["Category", "Q3 net sales", "Q3 plan", "Variance", "Return rate"],
        "rows": rows,
        "align": ["L", "R", "R", "R", "R"],
        "widths": [0.28, 0.18, 0.18, 0.16, 0.2],
        "bold_last": True,
        "note": "Variance is net sales against the Q3 2026 target (negative: below plan).",
    }
    missed = [c for c in var.sort_values().index if var[c] < 0]
    beat = [c for c in var.sort_values(ascending=False).index if var[c] >= 0]
    f["n_missed"], f["n_beat"] = str(len(missed)), str(len(beat))
    if len(missed) != 3:
        raise AssertionError("the story needs three categories below plan")
    f["missed_summary"] = (
        f"{join_and([f'{c} ({pct(var[c], 1, True)})' for c in missed])} finished below plan; {join_and(beat)} beat it."
    )
    f["elec_plan_variance"] = pct(var["Electronics"], 1, True)
    f["hr_category"], f["hr_category_rate"] = cat_rate.index[0], pct(cat_rate.iloc[0])
    f["elec_return_rate"] = pct(cat_rate["Electronics"])
    if cat_rate.index[0] != "Electronics":
        raise AssertionError("the documents say Electronics is the highest-returning category")
    a["missed"] = "; ".join(
        f"{c} {pct(var[c], 1, True)} ({usd(cat_sales[c])} against a {usd(plan[c])} target)" for c in missed
    )
    a["beat"] = "; ".join(f"{c} {pct(var[c], 1, True)}" for c in beat)
    a["cat_rate_rank"] = "; ".join(f"{c} {pct(v)}" for c, v in cat_rate.items())
    a["cat_top"] = (
        f"{cat_rate.index[0]} at {pct(cat_rate.iloc[0])} ({usd(cat_ret[cat_rate.index[0]])} returned "
        f"on {usd(cat_sales[cat_rate.index[0]])} net sales)"
    )

    # ---- return conditions
    cond_names = {"opened": "Opened or used", "unopened": "Unopened or unused", "defective": "Defective or damaged"}
    cm = r26.groupby("condition").agg(n=("return_id", "count"), v=("returned_amount", "sum"))
    rows = [
        [cond_names[c], f"{int(cm.at[c, 'n']):,}", usd(cm.at[c, "v"]), pct(cm.at[c, "v"] / cm["v"].sum())]
        for c in ("opened", "unopened", "defective")
    ]
    tables["condition_mix"] = {
        "caption": "Table 4. Q3 2026 returns by condition",
        "header": ["Condition", "Return records", "Returned value", "Share of returned value"],
        "rows": rows,
        "align": ["L", "R", "R", "R"],
        "widths": [0.34, 0.2, 0.22, 0.24],
    }
    f["web_return_share"] = pct(r26[r26["store_id"] == "WEB"]["returned_amount"].sum() / r26["returned_amount"].sum())

    # ---- electronics returns by physical store (hybrid question)
    el = r26[(r26["category"] == "Electronics") & (r26["store_type"] != "ecommerce")]
    ec = el.groupby("store_id").size().sort_values(ascending=False)
    es = ec.index[0]
    if ec.iloc[0] < ec.iloc[1] + 4 or es != datagen.ELECTRONICS_SPIKE_STORE:
        raise AssertionError("the electronics story needs a clear winner at the planted store")
    opened = el[(el["store_id"] == es) & (el["condition"] == "opened")]
    f["elec_spike_id"], f["elec_spike_name"] = es, short(smap.at[es, "store_name"])
    f["elec_spike_count"] = str(int(ec.iloc[0]))
    f["elec_spike_multiple"] = f"{ec.iloc[0] / ec.iloc[1]:.1f}"
    f["elec_spike_opened"] = f"{len(opened)} of {int(ec.iloc[0])}"
    fee = 0.15 * opened["returned_amount"].sum()
    tied = [str(i) for i, n in ec.items() if n == ec.iloc[1]]
    a["elec"] = (
        f"{es} {short(smap.at[es, 'store_name'])} with {int(ec.iloc[0])} electronics returns in Q3 2026 "
        f"(next highest: {int(ec.iloc[1])}, at {join_and(tied)}); "
        f"{len(opened)} of them were opened (non-defective) with {usd(opened['returned_amount'].sum(), 2)} returned, "
        f"so the 15% restocking fee is {usd(fee, 2)}"
    )
    a["elec_cond"] = ", ".join(f"{k} {int(v)}" for k, v in el[el["store_id"] == es].groupby("condition").size().items())
    a["elec_runner"] = ", ".join(f"{s} {int(n)}" for s, n in ec.head(4).items())
    a["elec_web"] = str(int(r26[(r26["category"] == "Electronics") & (r26["store_id"] == "WEB")].shape[0]))

    # ---- loss prevention memo and the monitored stores
    h_items = items[between(items["ordered_at"], H1S, H1E)]
    h_ret = returns[between(returns["returned_at"], H1S, H1E)]
    h_sales = h_items.groupby("store_id")["line_amount"].sum()
    h_rv = h_ret.groupby("store_id")["returned_amount"].sum()
    chain_h = h_rv.sum() / h_sales.sum()
    h_rate = h_rv.reindex(h_sales.index).fillna(0) / h_sales
    h_rate = h_rate[h_rate.index != "WEB"].sort_values(ascending=False)
    lp = h_rate.index[:3].tolist()
    if set(lp) != set(datagen.LP_MONITORED_STORES) or h_rate.iloc[2] / chain_h < 2.0 or h_rate.iloc[3] / chain_h > 1.9:
        raise AssertionError(f"the monitored stores are not the top three at twice the chain rate: {h_rate.head(4)}")
    rows = [
        [s, short(smap.at[s, "store_name"]), smap.at[s, "region"], pct(h_rate[s]), f"{h_rate[s] / chain_h:.1f}x"]
        for s in lp
    ]
    tables["lp_stores"] = {
        "header": ["Store ID", "Store", "Region", "Jan-Jun 2026 return rate", "Multiple of chain rate"],
        "rows": rows,
        "align": ["L", "L", "L", "R", "R"],
        "widths": [0.12, 0.34, 0.14, 0.22, 0.18],
    }
    f["lp_store_ids"] = join_and(sorted(lp))
    f["lp_stores"] = join_and([f"{s} {short(smap.at[s, 'store_name'])}" for s in sorted(lp)])
    lp_q3 = r26[r26["store_id"].isin(lp)]
    f["lp_q3_count"] = str(len(lp_q3))
    f["lp_q3_value"] = usd(lp_q3["returned_amount"].sum())
    f["lp_q3_share"] = pct(lp_q3["returned_amount"].sum() / r26["returned_amount"].sum())
    a["lp_rates"] = "; ".join(
        f"{s} {pct(h_rate[s])} ({h_rate[s] / chain_h:.1f}x the chain rate of {pct(chain_h)})" for s in lp
    )
    a["lp_q3"] = (
        f"{len(lp_q3)} returns worth {usd(lp_q3['returned_amount'].sum(), 2)} "
        f"({', '.join(f'{s} {int((lp_q3["store_id"] == s).sum())}' for s in sorted(lp))}); "
        f"that is {pct(lp_q3['returned_amount'].sum() / r26['returned_amount'].sum())} of the chain's "
        f"{usd(r26['returned_amount'].sum(), 2)} returned in Q3 2026"
    )

    # ---- promotions in Q3
    rows = []
    for p in promos.itertuples(index=False):
        if not (p.start_date <= date(2026, 9, 30) and p.end_date >= date(2026, 7, 1)):
            continue
        sub = o26[o26["promo_id"] == p.promo_id]
        scope = p.category if p.channel == "All" else f"{p.category} ({p.channel.lower()})"
        rows.append(
            [
                p.promo_name,
                f"{md_date(p.start_date)} to {md_date(p.end_date)}",
                scope,
                pct(p.discount_pct, 0),
                f"{len(sub):,}",
                usd(sub["net_amount"].sum()),
            ]
        )
    tables["promo_q3"] = {
        "header": ["Promotion", "Dates", "Scope", "Discount", "Orders", "Net sales"],
        "rows": rows,
        "align": ["L", "L", "L", "R", "R", "R"],
        "widths": [0.19, 0.17, 0.2, 0.15, 0.12, 0.17],
    }
    f["promo_order_share"] = pct(o26["promo_id"].notna().mean())
    f["promo_sales_share"] = pct(o26[o26["promo_id"].notna()]["net_amount"].sum() / net26)

    # ---- returned products, suppliers
    el_all = r26[r26["category"] == "Electronics"]
    prod = (
        el_all.groupby("product_id")
        .agg(n=("return_id", "count"), v=("returned_amount", "sum"))
        .sort_values(["n", "v"], ascending=False)
        .head(5)
    )
    rows = [[pmap.at[p, "product_name"], pmap.at[p, "supplier"], str(int(r.n)), usd(r.v)] for p, r in prod.iterrows()]
    tables["top_returned_products"] = {
        "header": ["Electronics product", "Supplier", "Q3 returns", "Returned value"],
        "rows": rows,
        "align": ["L", "L", "R", "R"],
        "widths": [0.42, 0.26, 0.14, 0.18],
    }
    sup = i26.groupby("supplier")["line_amount"].sum().sort_values(ascending=False)
    sup_cats = products.groupby("supplier")["category"].agg(lambda s: ", ".join(sorted(set(s))))
    rows = [[s, sup_cats[s], usd(v), pct(v / net26)] for s, v in sup.head(5).items()]
    tables["supplier_q3"] = {
        "header": ["Supplier", "Categories", "Q3 net sales", "Share"],
        "rows": rows,
        "align": ["L", "L", "R", "R"],
        "widths": [0.32, 0.34, 0.2, 0.14],
    }
    f["top_supplier"], f["top_supplier_sales"], f["top_supplier_share"] = (
        sup.index[0],
        usd(sup.iloc[0]),
        pct(sup.iloc[0] / net26),
    )
    elec_sup = i26[i26["category"] == "Electronics"].groupby("supplier")["line_amount"].sum()
    if elec_sup.get("Voltara Devices", 0) / elec_sup.sum() <= 0.5:
        raise AssertionError("the supplier terms say Voltara has the largest Electronics lines")

    # ---- Q4 plan
    q4 = targets[targets["quarter"] == "2026-Q4"].set_index("category")
    rows = [
        [
            c,
            usd(q4.at[c, "target_net_sales"]),
            pct(q4.at[c, "target_gross_margin_pct"]),
            f"{int(q4.at[c, 'target_units']):,}",
            pct(q4.at[c, "target_net_sales"] / cat_sales[c] - 1, 0, True),
        ]
        for c in datagen.CATEGORIES
    ]
    tables["q4_targets"] = {
        "header": ["Category", "Q4 net sales target", "Target margin", "Target units", "vs Q3 actual"],
        "rows": rows,
        "align": ["L", "R", "R", "R", "R"],
        "widths": [0.26, 0.24, 0.17, 0.16, 0.17],
    }
    q4_budget = sum(
        p.budget_usd for p in promos.itertuples(index=False) if date(2026, 10, 1) <= p.start_date <= date(2026, 12, 31)
    )
    f["q4_promo_budget"] = usd(q4_budget)

    kpis["q3_glance"] = [
        (usd(net26), "Net sales", f"{pct(net26 / net25 - 1, 1, True)} vs Q3 2025"),
        (f"{len(o26):,}", "Orders", f"{pct(len(o26) / len(o25) - 1, 1, True)} vs Q3 2025"),
        (
            usd(net26 / len(o26)),
            "Average order value",
            f"{pct((net26 / len(o26)) / (net25 / len(o25)) - 1, 1, True)} vs Q3 2025",
        ),
        (pct(chain_rate), "Return rate", "returned value / net sales"),
    ]

    # ---- answer key: regional report and predictions
    a["region_growth"] = (
        f"{top} grew {pct(reg[top]['yoy'], 1, True)} to {usd(reg[top]['s26'])} (Q3 2025: {usd(reg[top]['s25'])})"
    )
    a["region_rank"] = "; ".join(f"{r} {pct(reg[r]['yoy'], 1, True)}" for r in by_yoy)
    a["region_note"] = (
        f"among the four geographic regions the fastest is {max(datagen.REGIONS, key=lambda r: reg[r]['yoy'])}"
    )
    anchor = pd.Timestamp("2026-10-01")  # the Kumo anchor time: the day after the last data day
    gold = customers[customers["tier"] == "gold"]
    n_gold = len(gold)
    last_order = orders.groupby("customer_id")["ordered_at"].max()
    recent = last_order[last_order >= anchor - pd.Timedelta(days=90)].index
    a["n_gold"] = str(n_gold)
    a["n_tier"] = ", ".join(f"{k} {v}" for k, v in customers["tier"].value_counts().sort_index().items())
    a["gold_lapsed"] = str(int((~gold["customer_id"].isin(recent)).sum()))
    b_anchor = pd.Timestamp("2026-07-01")
    nxt = orders[between(orders["ordered_at"], b_anchor, b_anchor + pd.Timedelta(days=90))]["customer_id"].unique()
    hist = customers[customers["joined_at"] < b_anchor.date()]
    for t in ("gold", "silver", "bronze"):
        sub = hist[hist["tier"] == t]
        a[f"churn_{t}"] = pct((~sub["customer_id"].isin(nxt)).mean())
    a["churn_all"] = pct((~hist["customer_id"].isin(nxt)).mean())
    r_anchor = pd.Timestamp("2026-08-30")
    ret_next = returns[between(returns["returned_at"], r_anchor, r_anchor + pd.Timedelta(days=30))][
        "customer_id"
    ].nunique()
    a["return_backtest"] = f"{ret_next} of {len(customers):,} customers ({pct(ret_next / len(customers))})"
    o30 = (
        orders[between(orders["ordered_at"], r_anchor, r_anchor + pd.Timedelta(days=30))]
        .groupby("customer_id")["net_amount"]
        .sum()
    )
    a["hv_backtest"] = (
        f"{int((o30 > 300).sum())} of {len(customers):,} customers ({pct((o30 > 300).sum() / len(customers))})"
    )
    ret_hist = returns.groupby("customer_id").size()
    a["return_repeat"] = str(int((ret_hist >= 3).sum()))

    # return-risk-30d asks about one tier, so the population is under Kumo's limit of 1,000 entities per request
    silver = customers[customers["tier"] == "silver"]
    a["n_silver"] = str(len(silver))
    sid = silver[silver["joined_at"] < r_anchor.date()]["customer_id"]
    prior_ret = returns[returns["returned_at"] < r_anchor].groupby("customer_id").size()
    last_before = orders[orders["ordered_at"] < r_anchor].groupby("customer_id")["ordered_at"].max()
    ret_win = set(returns[between(returns["returned_at"], r_anchor, r_anchor + pd.Timedelta(days=30))]["customer_id"])
    hit = sid.isin(ret_win)
    hot_b = sid.map(prior_ret).fillna(0).ge(2) & (sid.map(last_before) >= r_anchor - pd.Timedelta(days=35))
    a["silver_return_backtest"] = f"{int(hit.sum())} of {len(sid)} silver customers ({pct(hit.mean())})"
    a["silver_hot_n"] = str(int(hot_b.sum()))
    a["silver_hot_rate"] = pct(hit[hot_b].mean())
    a["silver_rest_rate"] = pct(hit[~hot_b].mean())
    sn = silver["customer_id"]
    hot_now = sn.map(ret_hist).fillna(0).ge(2) & (sn.map(last_order) >= anchor - pd.Timedelta(days=35))
    a["silver_hot_now_n"] = str(int(hot_now.sum()))
    a["silver_hot_now_ids"] = ", ".join(sorted(sn[hot_now]))

    for name, df in T.items():
        f[f"rows_{name}"] = f"{len(df):,}"
    f["total_rows"] = f"{sum(len(df) for df in T.values()):,}"
    f["seed"] = str(datagen.SEED)
    f["n_products"] = str(len(products))
    f["n_customers"] = f"{len(customers):,}"
    f["first_order"] = orders["ordered_at"].min().strftime("%Y-%m-%d")
    f["last_order"] = orders["ordered_at"].max().strftime("%Y-%m-%d")
    f["gold_count"] = a["n_gold"]
    return {**f, **{f"a_{k}": v for k, v in a.items()}}, tables, charts, kpis


# --------------------------------------------------------------------------- documents


def build_documents(
    facts: dict[str, str], tables: dict[str, dict], charts: dict[str, dict], kpis: dict[str, list]
) -> None:
    pol, rep = FILES / "policies", FILES / "reports"
    pol.mkdir(parents=True, exist_ok=True)
    rep.mkdir(parents=True, exist_ok=True)

    meta, blocks = render.load_markdown(CONTENT / "return-refund-policy.md", facts)
    render.render_pdf(pol / "return-refund-policy.pdf", meta, blocks, tables, charts)

    meta, blocks = render.load_markdown(CONTENT / "store-operations-sop.md", facts)
    render.render_pdf(pol / "store-operations-sop.pdf", meta, blocks, tables, charts)

    meta, blocks = render.load_markdown(CONTENT / "supplier-terms.md", facts)
    render.render_docx(pol / "supplier-terms.docx", meta, blocks, tables)

    meta, blocks = render.load_markdown(CONTENT / "loss-prevention-memo.md", facts)
    memo = render.render_memo_pdf(meta, blocks, tables)
    scan = render.degrade_scan(render.rasterize_pdf(memo, dpi=160), SCAN_SEED)
    scan.save(pol / "loss-prevention-memo.png", optimize=True)

    meta, blocks = render.load_markdown(CONTENT / "regional-performance-q3-2026.md", facts)
    render.render_pdf(rep / "regional-performance-q3-2026.pdf", meta, blocks, tables, charts)

    meta, blocks = render.load_markdown(CONTENT / "q3-2026-merchandising-review.md", facts)
    render.render_pptx(rep / "q3-2026-merchandising-review.pptx", meta, blocks, tables, charts, kpis)


def build_readme(facts: dict[str, str]) -> None:
    text = (CONTENT / "readme.md").read_text(encoding="utf-8")
    import re

    # Drop the authoring note, keep the SPDX header.
    text = re.sub(r"<!--(?!\s*SPDX).*?-->\n?", "", text, flags=re.DOTALL)
    (PACK / "README.md").write_text(render.fill(text, facts), encoding="utf-8")


def main() -> None:
    started = datetime.now()
    T = datagen.generate()
    write_csvs(T)
    write_xlsx(T)
    facts, tables, charts, kpis = analyze(T)
    build_documents(facts, tables, charts, kpis)
    build_readme(facts)
    total = sum(p.stat().st_size for p in FILES.rglob("*") if p.is_file())
    print(
        f"wrote {sum(len(df) for df in T.values()):,} rows and {total / 1e6:.2f} MB of files "
        f"in {(datetime.now() - started).total_seconds():.1f}s"
    )
    for key in sorted(k for k in facts if k.startswith("a_")):
        print(f"{key[2:]}: {facts[key]}")


if __name__ == "__main__":
    main()
