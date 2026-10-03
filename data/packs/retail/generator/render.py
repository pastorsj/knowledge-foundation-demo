# SPDX-FileCopyrightText: Copyright (c) 2026, NVIDIA CORPORATION & AFFILIATES. All rights reserved.
# SPDX-License-Identifier: Apache-2.0
"""Render the authored Markdown in content/ into PDF, DOCX, PPTX and a scanned PNG.

The Markdown subset: front matter (`key: value` between `---` lines), `#`/`##`/`###` headings, paragraphs, `-` and
`1.` lists, pipe tables, `>` call-outs, `\\pagebreak`, `**bold**`/`*italic*`, `{{fact}}` placeholders and the block
directives `{{table:key}}`, `{{chart:key}}` and `{{kpis:key}}` that build.py fills from the generated tables.

Every renderer is deterministic: reportlab runs with `invariant=1`, the Office files get fixed core properties and
a normalized zip container (fixed timestamps), and the scan's noise comes from a seeded generator.
"""

# ruff: noqa: PLR0917  (drawing helpers take coordinates and sizes positionally)
from __future__ import annotations

import io
import math
import re
import zipfile
from dataclasses import dataclass
from dataclasses import field
from datetime import datetime
from pathlib import Path

import numpy as np
from PIL import Image
from PIL import ImageFilter
from reportlab import rl_config

rl_config.invariant = 1

FIXED_TIME = datetime(2026, 9, 30, 12, 0, 0)
ZIP_TIME = (2026, 9, 30, 12, 0, 0)

NAVY = "#1B2A49"
AMBER = "#E8A33D"
TEAL = "#1F7A8C"
SLATE = "#5B6577"
LIGHT = "#F2F4F8"
GRID = "#C5CDD9"

DIRECTIVE = re.compile(r"^\{\{(table|chart|kpis):([a-z0-9_]+)\}\}$")
PLACEHOLDER = re.compile(r"\{\{([a-z0-9_]+)\}\}")


# --------------------------------------------------------------------------- markdown subset


@dataclass
class Block:
    kind: str  # h1 h2 h3 p ul ol table callout directive pagebreak
    text: str = ""
    items: list[str] = field(default_factory=list)
    header: list[str] = field(default_factory=list)
    rows: list[list[str]] = field(default_factory=list)
    align: list[str] = field(default_factory=list)
    directive: str = ""
    key: str = ""


def fill(text: str, facts: dict[str, str]) -> str:
    def sub(m: re.Match) -> str:
        key = m.group(1)
        if key not in facts:
            raise KeyError(f"unknown placeholder {{{{{key}}}}}")
        return str(facts[key])

    return PLACEHOLDER.sub(sub, text)


def load_markdown(path: Path, facts: dict[str, str]) -> tuple[dict[str, str], list[Block]]:
    raw = path.read_text(encoding="utf-8")
    raw = re.sub(r"<!--.*?-->", "", raw, flags=re.DOTALL).lstrip("\n")
    meta: dict[str, str] = {}
    if raw.startswith("---\n"):
        end = raw.index("\n---\n", 4)
        for line in raw[4:end].splitlines():
            if line.strip():
                k, _, v = line.partition(":")
                meta[k.strip()] = fill(v.strip().strip('"'), facts)
        raw = raw[end + 5 :]
    blocks: list[Block] = []
    para: list[str] = []
    lines = raw.splitlines()

    def flush() -> None:
        if para:
            blocks.append(Block("p", text=fill(" ".join(s.strip() for s in para), facts)))
            para.clear()

    i = 0
    while i < len(lines):
        line = lines[i]
        s = line.strip()
        if not s:
            flush()
        elif s == "\\pagebreak":
            flush()
            blocks.append(Block("pagebreak"))
        elif DIRECTIVE.match(s):
            flush()
            m = DIRECTIVE.match(s)
            blocks.append(Block("directive", directive=m.group(1), key=m.group(2)))
        elif s.startswith("### "):
            flush()
            blocks.append(Block("h3", text=fill(s[4:], facts)))
        elif s.startswith("## "):
            flush()
            blocks.append(Block("h2", text=fill(s[3:], facts)))
        elif s.startswith("# "):
            flush()
            blocks.append(Block("h1", text=fill(s[2:], facts)))
        elif s.startswith("> "):
            flush()
            quote = [s[2:]]
            while i + 1 < len(lines) and lines[i + 1].strip().startswith("> "):
                i += 1
                quote.append(lines[i].strip()[2:])
            blocks.append(Block("callout", text=fill(" ".join(quote), facts)))
        elif re.match(r"^[-*] ", s) or re.match(r"^\d+\. ", s):
            flush()
            ordered = bool(re.match(r"^\d+\. ", s))
            items: list[str] = []
            while i < len(lines):
                cur = lines[i]
                cs = cur.strip()
                if re.match(r"^[-*] ", cs) and not ordered:
                    items.append(cs[2:])
                elif re.match(r"^\d+\. ", cs) and ordered:
                    items.append(re.sub(r"^\d+\. ", "", cs))
                elif cs and cur.startswith("  ") and items:
                    items[-1] += " " + cs
                else:
                    break
                i += 1
            i -= 1
            blocks.append(Block("ol" if ordered else "ul", items=[fill(t, facts) for t in items]))
        elif s.startswith("|"):
            flush()
            rows: list[list[str]] = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            i -= 1
            align = []
            for cell in rows[1]:
                c = cell.replace(" ", "")
                align.append("C" if c.startswith(":") and c.endswith(":") else "R" if c.endswith(":") else "L")
            blocks.append(
                Block(
                    "table",
                    header=[fill(c, facts) for c in rows[0]],
                    rows=[[fill(c, facts) for c in r] for r in rows[2:]],
                    align=align,
                )
            )
        else:
            para.append(s)
        i += 1
    flush()
    return meta, blocks


def esc(text: str) -> str:
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def inline_rl(text: str) -> str:
    """Markdown inline to reportlab paragraph markup."""
    t = esc(text)
    t = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", t)
    t = re.sub(r"(?<!\*)\*(?!\s)(.+?)(?<!\s)\*(?!\*)", r"<i>\1</i>", t)
    return t


def inline_runs(text: str) -> list[tuple[str, bool, bool]]:
    """Markdown inline to (text, bold, italic) runs."""
    out: list[tuple[str, bool, bool]] = []
    for part in re.split(r"(\*\*.+?\*\*|(?<!\*)\*(?!\s).+?(?<!\s)\*(?!\*))", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            out.append((part[2:-2], True, False))
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            out.append((part[1:-1], False, True))
        else:
            out.append((part, False, False))
    return out


# --------------------------------------------------------------------------- zip normalization


def normalize_zip(path: Path) -> None:
    """Rewrite an OOXML container with fixed timestamps so identical content gives identical bytes."""
    with zipfile.ZipFile(path) as zin:
        entries = [(info.filename, zin.read(info.filename)) for info in zin.infolist()]
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
        for name, raw in entries:
            data = raw
            if name == "docProps/core.xml":  # openpyxl stamps `modified` with the save time
                data = re.sub(
                    rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)",
                    rb"\g<1>" + FIXED_TIME.strftime("%Y-%m-%dT%H:%M:%SZ").encode() + rb"\g<2>",
                    raw,
                )
            zi = zipfile.ZipInfo(name, date_time=ZIP_TIME)
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o600 << 16
            zout.writestr(zi, data)
    path.write_bytes(buf.getvalue())


# --------------------------------------------------------------------------- PDF


def _pdf_styles():
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_RIGHT
    from reportlab.lib.styles import ParagraphStyle

    base = dict(fontName="Helvetica", fontSize=9.6, leading=13.4, textColor=colors.HexColor("#1F2430"))
    return {
        "body": ParagraphStyle("body", spaceAfter=6, **base),
        "title": ParagraphStyle(
            "title", fontName="Helvetica-Bold", fontSize=22, leading=26, textColor=colors.HexColor(NAVY), spaceAfter=4
        ),
        "subtitle": ParagraphStyle(
            "subtitle", fontName="Helvetica", fontSize=11, leading=15, textColor=colors.HexColor(SLATE), spaceAfter=10
        ),
        "h1": ParagraphStyle(
            "h1",
            fontName="Helvetica-Bold",
            fontSize=13.5,
            leading=17,
            textColor=colors.HexColor(NAVY),
            spaceBefore=12,
            spaceAfter=5,
            keepWithNext=1,
        ),
        "h2": ParagraphStyle(
            "h2",
            fontName="Helvetica-Bold",
            fontSize=11,
            leading=14,
            textColor=colors.HexColor(TEAL),
            spaceBefore=8,
            spaceAfter=3,
            keepWithNext=1,
        ),
        "h3": ParagraphStyle(
            "h3",
            fontName="Helvetica-BoldOblique",
            fontSize=10,
            leading=13,
            textColor=colors.HexColor(NAVY),
            spaceBefore=6,
            spaceAfter=2,
            keepWithNext=1,
        ),
        "bullet": ParagraphStyle("bullet", leftIndent=0, spaceAfter=2, **base),
        "cell": ParagraphStyle("cell", **{**base, "fontSize": 8.8, "leading": 11.2}),
        "cell_r": ParagraphStyle("cell_r", alignment=TA_RIGHT, **{**base, "fontSize": 8.8, "leading": 11.2}),
        "cell_h": ParagraphStyle(
            "cell_h",
            **{**base, "fontName": "Helvetica-Bold", "fontSize": 8.8, "leading": 11.2, "textColor": colors.white},
        ),
        "cell_hr": ParagraphStyle(
            "cell_hr",
            alignment=TA_RIGHT,
            **{**base, "fontName": "Helvetica-Bold", "fontSize": 8.8, "leading": 11.2, "textColor": colors.white},
        ),
        "caption": ParagraphStyle(
            "caption",
            fontName="Helvetica-Oblique",
            fontSize=8.4,
            leading=11,
            textColor=colors.HexColor(SLATE),
            spaceAfter=8,
        ),
        "callout": ParagraphStyle("callout", leftIndent=0, **{**base, "fontSize": 9.6}),
    }


def _pdf_table(header, rows, align, styles, width, widths=None, bold_last=False):
    from reportlab.lib import colors
    from reportlab.platypus import Paragraph
    from reportlab.platypus import Table
    from reportlab.platypus import TableStyle

    n = len(header)
    if widths is None:
        lens = [max([len(header[c])] + [len(r[c]) for r in rows]) for c in range(n)]
        lens = [min(max(v, 11), 34) for v in lens]
        widths = [v / sum(lens) for v in lens]
    data = [[Paragraph(inline_rl(h), styles["cell_hr" if align[c] == "R" else "cell_h"]) for c, h in enumerate(header)]]
    for ri, r in enumerate(rows):
        row = []
        for c, cell in enumerate(r):
            txt = inline_rl(cell)
            if bold_last and ri == len(rows) - 1:
                txt = f"<b>{txt}</b>"
            row.append(Paragraph(txt, styles["cell_r" if align[c] == "R" else "cell"]))
        data.append(row)
    t = Table(data, colWidths=[w * width for w in widths], repeatRows=1, hAlign="LEFT")
    style = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(NAVY)),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor(GRID)),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
    ]
    for ri in range(1, len(data)):
        if ri % 2 == 0:
            style.append(("BACKGROUND", (0, ri), (-1, ri), colors.HexColor(LIGHT)))
    if bold_last:
        style.append(("LINEABOVE", (0, len(data) - 1), (-1, len(data) - 1), 1.2, colors.HexColor(NAVY)))
    t.setStyle(TableStyle(style))
    return t


def _pdf_chart(spec: dict, width: float):
    """A grouped vertical bar chart (reportlab.graphics), data labels on top of each bar."""
    from reportlab.graphics.charts.barcharts import VerticalBarChart
    from reportlab.graphics.charts.legends import Legend
    from reportlab.graphics.shapes import Drawing
    from reportlab.graphics.shapes import String
    from reportlab.lib import colors

    height = 210
    d = Drawing(width, height)
    d.add(
        String(0, height - 12, spec["title"], fontName="Helvetica-Bold", fontSize=10, fillColor=colors.HexColor(NAVY))
    )
    ch = VerticalBarChart()
    ch.x, ch.y, ch.width, ch.height = 44, 34, width - 70, height - 80
    ch.data = [s[1] for s in spec["series"]]
    ch.categoryAxis.categoryNames = spec["labels"]
    ch.categoryAxis.labels.fontName = "Helvetica"
    ch.categoryAxis.labels.fontSize = 8.5
    ch.valueAxis.valueMin = 0
    top = max(max(s[1]) for s in spec["series"])
    raw = top / 4
    mag = 10 ** math.floor(math.log10(raw))
    step = next(m * mag for m in (1, 2, 2.5, 5, 10) if raw <= m * mag)
    ch.valueAxis.valueMax = step * math.ceil(top * 1.1 / step)
    ch.valueAxis.valueStep = step
    ch.valueAxis.labels.fontName = "Helvetica"
    ch.valueAxis.labels.fontSize = 8
    ch.valueAxis.labelTextFormat = spec.get("axis_fmt", "%d")
    ch.valueAxis.gridStrokeColor = colors.HexColor(GRID)
    ch.valueAxis.visibleGrid = 1
    ch.groupSpacing = 12
    ch.barSpacing = 2
    ch.barLabelFormat = spec.get("label_fmt", "%d")
    ch.barLabels.fontName = "Helvetica"
    ch.barLabels.fontSize = 7.4
    ch.barLabels.nudge = 7
    palette = [colors.HexColor(GRID), colors.HexColor(NAVY), colors.HexColor(AMBER)]
    for i in range(len(spec["series"])):
        ch.bars[i].fillColor = palette[i % len(palette)]
        ch.bars[i].strokeColor = None
    d.add(ch)
    lg = Legend()
    lg.x, lg.y = width - 160, height - 14
    lg.alignment = "right"
    lg.fontName, lg.fontSize = "Helvetica", 8.5
    lg.columnMaximum = 1
    lg.deltax = 60
    lg.colorNamePairs = [(palette[i % len(palette)], s[0]) for i, s in enumerate(spec["series"])]
    d.add(lg)
    return d


def render_pdf(
    path: Path,
    meta: dict[str, str],
    blocks: list[Block],
    tables: dict[str, dict],
    charts: dict[str, dict],
) -> None:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.units import inch
    from reportlab.pdfgen import canvas as rl_canvas
    from reportlab.platypus import KeepTogether
    from reportlab.platypus import ListFlowable
    from reportlab.platypus import ListItem
    from reportlab.platypus import PageBreak
    from reportlab.platypus import Paragraph
    from reportlab.platypus import SimpleDocTemplate
    from reportlab.platypus import Spacer
    from reportlab.platypus import Table
    from reportlab.platypus import TableStyle

    st = _pdf_styles()
    margin = 0.85 * inch
    page_w, page_h = letter
    width = page_w - 2 * margin
    title = meta["title"]
    doc_line = " | ".join(
        x for x in (meta.get("doc_id"), f"Version {meta['version']}" if meta.get("version") else None) if x
    )

    class Numbered(rl_canvas.Canvas):
        def __init__(self, *args, **kwargs):
            super().__init__(*args, **kwargs)
            self._pages: list[dict] = []

        def showPage(self):  # noqa: N802 (reportlab API)
            self._pages.append(dict(self.__dict__))
            self._startPage()

        def save(self):
            total = len(self._pages)
            for state in self._pages:
                self.__dict__.update(state)
                self._decorate(total)
                super().showPage()
            super().save()

        def _decorate(self, total: int) -> None:
            self.saveState()
            if self._pageNumber > 1:
                self.setFont("Helvetica-Bold", 8.5)
                self.setFillColor(colors.HexColor(NAVY))
                self.drawString(margin, page_h - 0.55 * inch, "LUMEN RETAIL GROUP")
                self.setFont("Helvetica", 8.5)
                self.setFillColor(colors.HexColor(SLATE))
                self.drawRightString(page_w - margin, page_h - 0.55 * inch, title)
                self.setStrokeColor(colors.HexColor(GRID))
                self.line(margin, page_h - 0.62 * inch, page_w - margin, page_h - 0.62 * inch)
            self.setStrokeColor(colors.HexColor(GRID))
            self.line(margin, 0.62 * inch, page_w - margin, 0.62 * inch)
            self.setFont("Helvetica", 8)
            self.setFillColor(colors.HexColor(SLATE))
            self.drawString(margin, 0.45 * inch, f"{doc_line}   {meta.get('footer', 'Internal use only')}".strip())
            self.drawRightString(page_w - margin, 0.45 * inch, f"Page {self._pageNumber} of {total}")
            self.restoreState()

    def first_page(canv, doc):  # brand band
        canv.saveState()
        canv.setFillColor(colors.HexColor(NAVY))
        canv.rect(0, page_h - 0.52 * inch, page_w, 0.52 * inch, stroke=0, fill=1)
        canv.setFillColor(colors.HexColor(AMBER))
        canv.rect(0, page_h - 0.58 * inch, page_w, 0.06 * inch, stroke=0, fill=1)
        canv.setFillColor(colors.white)
        canv.setFont("Helvetica-Bold", 12)
        canv.drawString(margin, page_h - 0.34 * inch, "LUMEN RETAIL GROUP")
        canv.setFont("Helvetica", 9)
        canv.drawRightString(page_w - margin, page_h - 0.34 * inch, meta.get("department", ""))
        canv.restoreState()

    doc = SimpleDocTemplate(
        str(path),
        pagesize=letter,
        leftMargin=margin,
        rightMargin=margin,
        topMargin=0.95 * inch,
        bottomMargin=0.85 * inch,
        title=title,
        author="Lumen Retail Group",
        subject=meta.get("subject", title),
        creator="Lumen Retail Group document services",
        invariant=1,
    )
    story: list = [Spacer(1, 6), Paragraph(esc(title), st["title"])]
    if meta.get("subtitle"):
        story.append(Paragraph(esc(meta["subtitle"]), st["subtitle"]))
    info_rows = [
        (k, meta[f])
        for k, f in (
            ("Document", "doc_id"),
            ("Version", "version"),
            ("Effective", "effective"),
            ("Owner", "owner"),
            ("Audience", "classification"),
        )
        if meta.get(f)
    ]
    if info_rows:
        cells = []
        for k, v in info_rows:
            cells.append([Paragraph(f"<b>{esc(k)}</b>", st["cell"]), Paragraph(esc(v), st["cell"])])
        half = (len(cells) + 1) // 2
        left, right = cells[:half], cells[half:]
        while len(right) < len(left):
            right.append(["", ""])
        data = [a + b for a, b in zip(left, right, strict=True)]
        t = Table(data, colWidths=[0.14 * width, 0.36 * width, 0.14 * width, 0.36 * width], hAlign="LEFT")
        t.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor(LIGHT)),
                    ("LINEBELOW", (0, 0), (-1, -1), 0.4, colors.HexColor(GRID)),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("TOPPADDING", (0, 0), (-1, -1), 4),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ]
            )
        )
        story += [t, Spacer(1, 8)]

    def take_lead(flow: list) -> list:
        """Pop the heading and intro paragraph just before a table so they stay on the table's page."""
        lead: list = []
        if flow and isinstance(flow[-1], Paragraph) and flow[-1].style.name == "body":
            lead.insert(0, flow.pop())
        if flow and isinstance(flow[-1], Paragraph) and flow[-1].style.name in ("h1", "h2", "h3"):
            lead.insert(0, flow.pop())
        return lead

    for b in blocks:
        if b.kind == "h1":
            story.append(Paragraph(inline_rl(b.text), st["h1"]))
        elif b.kind == "h2":
            story.append(Paragraph(inline_rl(b.text), st["h2"]))
        elif b.kind == "h3":
            story.append(Paragraph(inline_rl(b.text), st["h3"]))
        elif b.kind == "p":
            story.append(Paragraph(inline_rl(b.text), st["body"]))
        elif b.kind in ("ul", "ol"):
            items = [ListItem(Paragraph(inline_rl(t), st["bullet"]), leftIndent=16) for t in b.items]
            if b.kind == "ul":
                story.append(ListFlowable(items, bulletType="bullet", start="•", leftIndent=16, bulletFontSize=9))
            else:
                story.append(
                    ListFlowable(
                        items,
                        bulletType="1",
                        bulletFormat="%s.",
                        leftIndent=18,
                        bulletFontName="Helvetica",
                        bulletFontSize=9.4,
                    )
                )
            story.append(Spacer(1, 4))
        elif b.kind == "table":
            tbl = _pdf_table(b.header, b.rows, b.align, st, width)
            if len(b.rows) <= 12:
                story.append(KeepTogether([*take_lead(story), tbl, Spacer(1, 8)]))
            else:
                story += [tbl, Spacer(1, 8)]
        elif b.kind == "callout":
            t = Table([[Paragraph(inline_rl(b.text), st["callout"])]], colWidths=[width], hAlign="LEFT")
            t.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#FFF6E5")),
                        ("LINEBEFORE", (0, 0), (0, -1), 3, colors.HexColor(AMBER)),
                        ("TOPPADDING", (0, 0), (-1, -1), 6),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                        ("LEFTPADDING", (0, 0), (-1, -1), 10),
                    ]
                )
            )
            story += [t, Spacer(1, 8)]
        elif b.kind == "pagebreak":
            story.append(PageBreak())
        elif b.kind == "directive":
            if b.directive == "table":
                spec = tables[b.key]
                parts = []
                if spec.get("caption"):
                    parts.append(Paragraph(f"<b>{esc(spec['caption'])}</b>", st["body"]))
                parts.append(
                    _pdf_table(
                        spec["header"],
                        spec["rows"],
                        spec["align"],
                        st,
                        width,
                        spec.get("widths"),
                        spec.get("bold_last", False),
                    )
                )
                if spec.get("note"):
                    parts.append(Spacer(1, 3))
                    parts.append(Paragraph(esc(spec["note"]), st["caption"]))
                else:
                    parts.append(Spacer(1, 8))
                story.append(KeepTogether([*take_lead(story), *parts]))
            elif b.directive == "chart":
                story += [_pdf_chart(charts[b.key], width), Spacer(1, 6)]
            else:
                raise ValueError(f"directive {b.directive} is not supported in PDF")
    doc.build(story, onFirstPage=first_page, canvasmaker=Numbered)


# --------------------------------------------------------------------------- scanned memo


def render_memo_pdf(meta: dict[str, str], blocks: list[Block], tables: dict[str, dict]) -> bytes:
    """A one-page typed memo with a rubber stamp, as PDF bytes (it is then rasterized and degraded)."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import BaseDocTemplate
    from reportlab.platypus import Frame
    from reportlab.platypus import PageTemplate
    from reportlab.platypus import Paragraph
    from reportlab.platypus import Spacer
    from reportlab.platypus import Table
    from reportlab.platypus import TableStyle

    page_w, page_h = letter
    margin = 0.95 * inch
    width = page_w - 2 * margin
    ink = colors.HexColor("#111111")
    body = ParagraphStyle("mb", fontName="Times-Roman", fontSize=11.6, leading=15.2, textColor=ink, spaceAfter=7)
    li = ParagraphStyle("li", parent=body, leftIndent=0, spaceAfter=6)
    hd = ParagraphStyle(
        "mh", fontName="Times-Bold", fontSize=11.6, leading=15, textColor=ink, spaceBefore=4, spaceAfter=3
    )
    cell = ParagraphStyle("mc", fontName="Times-Roman", fontSize=10.6, leading=13, textColor=ink)
    cellb = ParagraphStyle("mcb", parent=cell, fontName="Times-Bold")

    def stamp(canv, doc):
        canv.saveState()
        canv.setFont("Times-Bold", 17)
        canv.drawCentredString(page_w / 2, page_h - 0.95 * inch, "LUMEN RETAIL GROUP")
        canv.setFont("Times-Roman", 10.5)
        canv.drawCentredString(page_w / 2, page_h - 1.17 * inch, "Asset Protection  |  Loss Prevention Office")
        canv.setLineWidth(1.6)
        canv.line(margin, page_h - 1.3 * inch, page_w - margin, page_h - 1.3 * inch)
        canv.setLineWidth(0.5)
        canv.line(margin, page_h - 1.34 * inch, page_w - margin, page_h - 1.34 * inch)
        # rubber stamp
        canv.translate(page_w - 1.55 * inch, page_h - 0.98 * inch)
        canv.rotate(7)
        canv.setStrokeColor(colors.HexColor("#444444"))
        canv.setFillColor(colors.HexColor("#444444"))
        canv.setLineWidth(2.2)
        canv.roundRect(-0.9 * inch, -0.38 * inch, 1.8 * inch, 0.76 * inch, 6, stroke=1, fill=0)
        canv.setFont("Helvetica-Bold", 12)
        canv.drawCentredString(0, 0.12 * inch, meta.get("stamp_top", "POSTED"))
        canv.setFont("Helvetica-Bold", 9.5)
        canv.drawCentredString(0, -0.12 * inch, meta.get("stamp_bottom", ""))
        canv.restoreState()

    out = io.BytesIO()
    doc = BaseDocTemplate(out, pagesize=letter, title=meta["title"], author="Lumen Retail Group", invariant=1)
    frame = Frame(
        margin, 0.8 * inch, width, page_h - 2.25 * inch, leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0
    )
    doc.addPageTemplates([PageTemplate(id="memo", frames=[frame], onPage=stamp)])
    head = [
        ("TO:", meta["memo_to"]),
        ("FROM:", meta["memo_from"]),
        ("DATE:", meta["memo_date"]),
        ("RE:", meta["memo_re"]),
        ("MEMO NO.:", meta["memo_id"]),
    ]
    ht = Table(
        [[Paragraph(f"<b>{k}</b>", cell), Paragraph(esc(v), cell)] for k, v in head],
        colWidths=[1.15 * inch, width - 1.15 * inch],
    )
    ht.setStyle(
        TableStyle(
            [
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("TOPPADDING", (0, 0), (-1, -1), 1.5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 1.5),
                ("LINEBELOW", (0, -1), (-1, -1), 0.8, ink),
            ]
        )
    )
    story: list = [
        Paragraph(
            "<b>MEMORANDUM</b>", ParagraphStyle("mt", fontName="Times-Bold", fontSize=15, leading=20, spaceAfter=6)
        ),
        ht,
        Spacer(1, 10),
    ]
    for b in blocks:
        if b.kind in ("h1", "h2", "h3"):
            story.append(Paragraph(inline_rl(b.text), hd))
        elif b.kind == "p":
            story.append(Paragraph(inline_rl(b.text), body))
        elif b.kind in ("ul", "ol"):
            for n, t in enumerate(b.items, 1):
                mark = f"{n}." if b.kind == "ol" else "-"
                story.append(
                    Paragraph(
                        f"{mark}&nbsp;&nbsp;{inline_rl(t)}",
                        ParagraphStyle("x", parent=li, leftIndent=18, firstLineIndent=-18),
                    )
                )
        elif b.kind == "directive" and b.directive == "table":
            spec = tables[b.key]
            data = [[Paragraph(esc(h), cellb) for h in spec["header"]]]
            for r in spec["rows"]:
                data.append([Paragraph(esc(c), cell) for c in r])
            widths = spec.get("widths") or [1 / len(spec["header"])] * len(spec["header"])
            t = Table(data, colWidths=[w * width for w in widths], hAlign="LEFT")
            t.setStyle(
                TableStyle(
                    [
                        ("GRID", (0, 0), (-1, -1), 0.6, ink),
                        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                        ("TOPPADDING", (0, 0), (-1, -1), 3),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                    ]
                )
            )
            story += [t, Spacer(1, 9)]
    doc.build(story)
    return out.getvalue()


def rasterize_pdf(pdf_bytes: bytes, dpi: int = 160) -> Image.Image:
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(pdf_bytes)
    page = pdf[0]
    bitmap = page.render(scale=dpi / 72.0)
    img = bitmap.to_pil().convert("RGB")
    page.close()
    pdf.close()
    return img


def degrade_scan(img: Image.Image, seed: int, angle: float = 1.4) -> Image.Image:
    """Make a clean page look like a flatbed scan: grayscale, slight rotation, paper tone, shadow, noise, dust."""
    rng = np.random.default_rng(seed)
    g = img.convert("L").filter(ImageFilter.GaussianBlur(0.65))
    g = g.rotate(angle, resample=Image.Resampling.BICUBIC, expand=True, fillcolor=214)
    a = np.asarray(g, dtype=np.float32)
    h, w = a.shape
    a = 26.0 + a * (218.0 / 255.0)  # off-white paper, softened blacks
    xx = np.linspace(0.0, 1.0, w, dtype=np.float32)[None, :]
    yy = np.linspace(0.0, 1.0, h, dtype=np.float32)[:, None]
    shade = 1.0 - 0.085 * np.exp(-xx * 11.0) - 0.05 * (yy**2) * xx - 0.03 * np.exp(-(1.0 - xx) * 14.0)
    a *= shade
    a += rng.normal(0.0, 5.5, a.shape).astype(np.float32)
    for _ in range(140):  # dust
        cy, cx = int(rng.integers(0, h)), int(rng.integers(0, w))
        r = int(rng.integers(1, 3))
        a[max(cy - r, 0) : cy + r, max(cx - r, 0) : cx + r] -= float(rng.uniform(25, 70))
    for _ in range(3):  # faint vertical scanner streaks
        cx = int(rng.integers(int(0.1 * w), int(0.9 * w)))
        a[:, cx : cx + 1] -= float(rng.uniform(4, 9))
    return Image.fromarray(np.clip(a, 0, 255).astype(np.uint8), mode="L")


# --------------------------------------------------------------------------- DOCX


def render_docx(path: Path, meta: dict[str, str], blocks: list[Block], tables: dict[str, dict]) -> None:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches
    from docx.shared import Pt
    from docx.shared import RGBColor

    def rgb(hexstr: str) -> RGBColor:
        return RGBColor.from_string(hexstr.lstrip("#"))

    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Inches(1.0)
    sec.top_margin = Inches(0.9)
    sec.bottom_margin = Inches(0.9)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    for name, size, color in (("Heading 1", 15, NAVY), ("Heading 2", 12.5, TEAL), ("Heading 3", 11, NAVY)):
        s = doc.styles[name]
        s.font.name = "Calibri"
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.color.rgb = rgb(color)
    doc.styles["Title"].font.name = "Calibri"
    doc.styles["Title"].font.size = Pt(24)
    doc.styles["Title"].font.color.rgb = rgb(NAVY)

    def shade(cell, hexcolor: str) -> None:
        tcpr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), hexcolor.lstrip("#"))
        tcpr.append(shd)

    def add_runs(par, text: str, color: str | None = None, bold_all: bool = False, size: float | None = None) -> None:
        for t, b, i in inline_runs(text):
            run = par.add_run(t)
            run.bold = b or bold_all
            run.italic = i
            if color:
                run.font.color.rgb = rgb(color)
            if size:
                run.font.size = Pt(size)

    def add_table(header, rows, align) -> None:
        t = doc.add_table(rows=1, cols=len(header))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.LEFT
        for c, h in enumerate(header):
            cell = t.rows[0].cells[c]
            cell.text = ""
            add_runs(cell.paragraphs[0], h, color="#FFFFFF", bold_all=True, size=9.5)
            shade(cell, NAVY)
            if align[c] == "R":
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for ri, r in enumerate(rows):
            cells = t.add_row().cells
            for c, v in enumerate(r):
                cells[c].text = ""
                add_runs(cells[c].paragraphs[0], v, size=9.5)
                if align[c] == "R":
                    cells[c].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
                if ri % 2 == 1:
                    shade(cells[c], LIGHT)
        doc.add_paragraph()

    hdr = sec.header.paragraphs[0]
    hdr.text = f"LUMEN RETAIL GROUP  |  {meta['title']}"
    hdr.runs[0].font.size = Pt(8.5)
    hdr.runs[0].font.color.rgb = rgb(SLATE)
    ftr = sec.footer.paragraphs[0]
    ftr.text = f"{meta.get('doc_id', '')}  {meta.get('footer', 'Internal use only')}   |   Page "
    ftr.runs[0].font.size = Pt(8.5)
    ftr.runs[0].font.color.rgb = rgb(SLATE)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run = OxmlElement("w:r")
    txt = OxmlElement("w:t")
    txt.text = "1"
    run.append(txt)
    fld.append(run)
    ftr._p.append(fld)

    doc.add_paragraph(meta["title"], style="Title")
    if meta.get("subtitle"):
        p = doc.add_paragraph()
        add_runs(p, meta["subtitle"], color=SLATE, size=11.5)
    info = [
        (k, meta[f])
        for k, f in (
            ("Document", "doc_id"),
            ("Version", "version"),
            ("Effective", "effective"),
            ("Owner", "owner"),
            ("Audience", "classification"),
        )
        if meta.get(f)
    ]
    if info:
        t = doc.add_table(rows=0, cols=2)
        t.style = "Table Grid"
        for k, v in info:
            cells = t.add_row().cells
            cells[0].text = ""
            cells[1].text = ""
            add_runs(cells[0].paragraphs[0], k, bold_all=True, size=9.5)
            add_runs(cells[1].paragraphs[0], v, size=9.5)
            shade(cells[0], LIGHT)
            cells[0].width = Inches(1.2)
            cells[1].width = Inches(5.3)
        doc.add_paragraph()

    for b in blocks:
        if b.kind == "h1":
            doc.add_heading(b.text, level=1)
        elif b.kind == "h2":
            doc.add_heading(b.text, level=2)
        elif b.kind == "h3":
            doc.add_heading(b.text, level=3)
        elif b.kind == "p":
            add_runs(doc.add_paragraph(), b.text)
        elif b.kind in ("ul", "ol"):
            for t in b.items:
                p = doc.add_paragraph(style="List Bullet" if b.kind == "ul" else "List Number")
                add_runs(p, t)
        elif b.kind == "table":
            add_table(b.header, b.rows, b.align)
        elif b.kind == "callout":
            t = doc.add_table(rows=1, cols=1)
            t.style = "Table Grid"
            cell = t.rows[0].cells[0]
            cell.text = ""
            add_runs(cell.paragraphs[0], b.text, size=10)
            shade(cell, "#FFF6E5")
            doc.add_paragraph()
        elif b.kind == "pagebreak":
            doc.add_page_break()
        elif b.kind == "directive" and b.directive == "table":
            spec = tables[b.key]
            if spec.get("caption"):
                add_runs(doc.add_paragraph(), spec["caption"], bold_all=True)
            add_table(spec["header"], spec["rows"], spec["align"])
            if spec.get("note"):
                add_runs(doc.add_paragraph(), spec["note"], color=SLATE, size=9)
    cp = doc.core_properties
    cp.author = "Lumen Retail Group - Procurement"
    cp.last_modified_by = "Lumen Retail Group - Procurement"
    cp.title = meta["title"]
    cp.subject = meta.get("subject", meta["title"])
    cp.comments = "Synthetic document for a software demonstration."
    cp.keywords = "synthetic, retail, supplier terms"
    cp.created = FIXED_TIME
    cp.modified = FIXED_TIME
    cp.revision = 1
    doc.save(str(path))
    normalize_zip(path)


# --------------------------------------------------------------------------- PPTX


def render_pptx(
    path: Path,
    meta: dict[str, str],
    blocks: list[Block],
    tables: dict[str, dict],
    charts: dict[str, dict],
    kpis: dict[str, list[tuple[str, str, str]]],
) -> None:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Emu
    from pptx.util import Inches
    from pptx.util import Pt

    def rgb(hexstr: str) -> RGBColor:
        return RGBColor.from_string(hexstr.lstrip("#"))

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    sw = 13.333

    def set_text(tf, text: str, size: float, color: str = "#1F2430", bold: bool = False, align=None) -> None:
        tf.clear()
        first = True
        for line in text.split("\n"):
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            if align is not None:
                p.alignment = align
            for t, b, i in inline_runs(line):
                r = p.add_run()
                r.text = t
                r.font.size = Pt(size)
                r.font.bold = b or bold
                r.font.italic = i
                r.font.color.rgb = rgb(color)
                r.font.name = "Calibri"

    def textbox(slide, x, y, w, h, text, size=14, color="#1F2430", bold=False, align=None, anchor=None):
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tb.text_frame.word_wrap = True
        if anchor is not None:
            tb.text_frame.vertical_anchor = anchor
        set_text(tb.text_frame, text, size, color, bold, align)
        return tb

    def rect(slide, x, y, w, h, fill, shape=MSO_SHAPE.RECTANGLE):
        s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
        s.line.fill.background()
        s.shadow.inherit = False
        return s

    def chrome(slide, title: str, number: int) -> None:
        rect(slide, 0, 0, sw, 0.18, NAVY)
        rect(slide, 0, 0.18, sw, 0.05, AMBER)
        t = slide.shapes.title
        t.left, t.top, t.width, t.height = Inches(0.6), Inches(0.42), Inches(sw - 1.2), Inches(0.9)
        set_text(t.text_frame, title, 30, NAVY, True, PP_ALIGN.LEFT)
        t.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        textbox(slide, 0.6, 7.05, 8, 0.3, "Lumen Retail Group  |  Merchandising and Planning  |  Internal", 10, SLATE)
        textbox(slide, sw - 1.6, 7.05, 1.0, 0.3, str(number), 10, SLATE, align=PP_ALIGN.RIGHT)

    def add_table(slide, spec, x, y, w, row_h=0.38, font=12) -> None:
        n_rows, n_cols = len(spec["rows"]) + 1, len(spec["header"])
        gs = slide.shapes.add_table(n_rows, n_cols, Inches(x), Inches(y), Inches(w), Inches(row_h * n_rows))
        tbl = gs.table
        widths = spec.get("widths") or [1 / n_cols] * n_cols
        for c in range(n_cols):
            tbl.columns[c].width = Emu(int(Inches(w) * widths[c]))
        for r in range(n_rows):
            tbl.rows[r].height = Inches(row_h)
            for c in range(n_cols):
                cell = tbl.cell(r, c)
                text = spec["header"][c] if r == 0 else spec["rows"][r - 1][c]
                cell.margin_left = cell.margin_right = Inches(0.08)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                al = PP_ALIGN.RIGHT if spec["align"][c] == "R" else PP_ALIGN.LEFT
                bold = r == 0 or (spec.get("bold_last") and r == n_rows - 1)
                set_text(cell.text_frame, text, font, "#FFFFFF" if r == 0 else "#1F2430", bold, al)
                cell.fill.solid()
                cell.fill.fore_color.rgb = rgb(NAVY if r == 0 else (LIGHT if r % 2 == 0 else "#FFFFFF"))

    def add_bars(slide, spec, x, y, w, h) -> None:
        labels, values = spec["labels"], spec["series"][-1][1]
        top = max(values)
        row_h = h / len(labels)
        textbox(slide, x, y - 0.42, w, 0.35, spec["title"], 14, NAVY, True)
        for i, (lab, v) in enumerate(zip(labels, values, strict=True)):
            yy = y + i * row_h
            textbox(slide, x, yy, 2.1, row_h, lab, 13, "#1F2430", anchor=MSO_ANCHOR.MIDDLE)
            bar_w = max(0.05, (w - 3.5) * v / top)
            rect(slide, x + 2.15, yy + row_h * 0.18, bar_w, row_h * 0.64, AMBER if v == top else TEAL)
            textbox(
                slide, x + 2.2 + bar_w, yy, 1.3, row_h, spec["label_fmt"](v), 13, NAVY, True, anchor=MSO_ANCHOR.MIDDLE
            )

    def add_kpis(slide, cards, y) -> None:
        n = len(cards)
        gap = 0.25
        cw = (sw - 1.2 - gap * (n - 1)) / n
        for i, (big, label, sub) in enumerate(cards):
            x = 0.6 + i * (cw + gap)
            rect(slide, x, y, cw, 1.55, LIGHT, MSO_SHAPE.ROUNDED_RECTANGLE)
            rect(slide, x, y, 0.08, 1.55, AMBER)
            textbox(slide, x + 0.2, y + 0.08, cw - 0.3, 0.65, big, 28, NAVY, True)
            textbox(slide, x + 0.2, y + 0.72, cw - 0.3, 0.35, label, 13, "#1F2430", True)
            textbox(slide, x + 0.2, y + 1.06, cw - 0.3, 0.4, sub, 11, SLATE)

    # title slide
    s0 = prs.slides.add_slide(prs.slide_layouts[0])
    rect(s0, 0, 0, sw, 7.5, NAVY)
    rect(s0, 0, 4.9, sw, 0.07, AMBER)
    t = s0.shapes.title
    t.left, t.top, t.width, t.height = Inches(0.9), Inches(2.0), Inches(sw - 1.8), Inches(1.5)
    set_text(t.text_frame, meta["title"], 44, "#FFFFFF", True, PP_ALIGN.LEFT)
    sub = s0.placeholders[1]
    sub.left, sub.top, sub.width, sub.height = Inches(0.9), Inches(3.55), Inches(sw - 1.8), Inches(1.2)
    set_text(sub.text_frame, meta.get("subtitle", ""), 20, "#D8DEE9", False, PP_ALIGN.LEFT)
    textbox(s0, 0.9, 5.2, 9, 0.9, meta.get("presenter", ""), 16, "#FFFFFF")
    textbox(s0, 0.9, 6.7, 9, 0.4, "Synthetic data for a software demonstration. Internal use only.", 11, "#AAB4C5")

    # content slides: each `##` heading starts a slide
    slides: list[tuple[str, list[Block]]] = []
    for b in blocks:
        if b.kind == "h2":
            slides.append((b.text, []))
        elif slides:
            slides[-1][1].append(b)
    for n, (title, body) in enumerate(slides, start=2):
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        chrome(slide, title, n)
        notes = " ".join(b.text for b in body if b.kind == "callout")
        content = [b for b in body if b.kind != "callout"]
        table_spec = next((tables[b.key] for b in content if b.kind == "directive" and b.directive == "table"), None)
        chart_spec = next((charts[b.key] for b in content if b.kind == "directive" and b.directive == "chart"), None)
        kpi_cards = next((kpis[b.key] for b in content if b.kind == "directive" and b.directive == "kpis"), None)
        bullets: list[str] = []
        for b in content:
            if b.kind in ("ul", "ol"):
                bullets += b.items
            elif b.kind == "p":
                bullets.append(b.text)
        y0 = 1.55
        if kpi_cards:
            add_kpis(slide, kpi_cards, y0)
            y0 += 1.85
        side = table_spec is not None or chart_spec is not None
        text_w = 5.4 if side else sw - 1.2
        if bullets:
            tb = slide.shapes.add_textbox(Inches(0.6), Inches(y0), Inches(text_w), Inches(7.0 - y0 - 0.2))
            tb.text_frame.word_wrap = True
            first = True
            for item in bullets:
                p = tb.text_frame.paragraphs[0] if first else tb.text_frame.add_paragraph()
                first = False
                p.space_after = Pt(9)
                for t_, b_, i_ in inline_runs("•  " + item):
                    r = p.add_run()
                    r.text = t_
                    r.font.size = Pt(17 if not side else 15.5)
                    r.font.bold = b_
                    r.font.italic = i_
                    r.font.name = "Calibri"
                    r.font.color.rgb = rgb("#1F2430")
        if table_spec:
            nrows = len(table_spec["rows"]) + 1
            row_h = min(0.46, (7.0 - y0 - 0.3) / nrows)
            add_table(
                slide,
                table_spec,
                6.3,
                y0,
                sw - 6.9,
                row_h,
                font=12 if nrows < 10 and len(table_spec["header"]) < 6 else 11,
            )
        if chart_spec:
            add_bars(slide, chart_spec, 6.3, y0 + 0.5, sw - 6.9, min(4.6, 0.82 * len(chart_spec["labels"])))
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
    cp = prs.core_properties
    cp.author = "Lumen Retail Group - Merchandising and Planning"
    cp.last_modified_by = "Lumen Retail Group - Merchandising and Planning"
    cp.title = meta["title"]
    cp.subject = "Quarterly merchandising review"
    cp.comments = "Synthetic document for a software demonstration."
    cp.keywords = "synthetic, retail, merchandising"
    cp.created = FIXED_TIME
    cp.modified = FIXED_TIME
    cp.revision = 1
    prs.save(str(path))
    normalize_zip(path)
