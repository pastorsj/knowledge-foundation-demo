#!/usr/bin/env python3
# SPDX-FileCopyrightText: Copyright (c) 2026, NVIDIA CORPORATION & AFFILIATES. All rights reserved.
# SPDX-License-Identifier: Apache-2.0
# ruff: noqa: E501, PLR0917, PLW2901
# /// script
# requires-python = ">=3.12"
# dependencies = [
#     "numpy",
#     "pandas",
#     "openpyxl",
#     "reportlab",
#     "python-docx",
#     "python-pptx",
#     "pillow",
#     "pypdfium2",
# ]
# ///
"""Build the Financial Services (retail banking) industry pack: Harborview Community Bank.

    uv run data/packs/financial-services/generator/build.py

Everything is synthetic and seeded, and the same inputs give the same bytes. The script writes, under files/:

    banking/*.csv                  branches, customers, accounts, monthly_balances, loans, loan_payments,
                                   card_disputes
    banking/credit_risk_report.xlsx   two sheets: delinquency (by product, month) and provisions (by segment, quarter)
    policies/*                     consumer credit policy (PDF), KYC/AML procedure (PDF), complaint-handling
                                   procedure (DOCX), a scanned loan modification request (PNG)
    disclosures/*                  overdraft and fee schedule (PDF), Q3 2026 risk committee deck (PPTX)

and README.md (the data card and answer key, from content/README.template.md). The prose lives in content/*.md;
the numbers in the risk deck, the credit risk report and the answer key are computed from the generated tables.

Sections: world simulation (customers, accounts, loans, disputes), metrics, writers (CSV, XLSX), document
renderers (PDF, DOCX, PPTX, scanned PNG), README.
"""

from __future__ import annotations

import calendar
import datetime as dt
import io
import re
import shutil
import zipfile
import zlib
from datetime import date
from datetime import timedelta
from pathlib import Path

import numpy as np
import pandas as pd

HERE = Path(__file__).resolve().parent
PACK = HERE.parent
FILES = PACK / "files"
CONTENT = HERE / "content"

SEED = 20260930
AS_OF = date(2026, 9, 30)
WINDOW_START = date(2025, 4, 1)  # first month of the published history
SIM_START = date(2024, 10, 1)  # six months of burn-in so the window opens in a realistic state
N_SIM = 24  # simulated months: 2024-10 .. 2026-09
OUT_FIRST = 6  # index of the first published month (2025-04)
N_CUSTOMERS = 1060
FIXED_TIME = dt.datetime(2026, 9, 30, 12, 0, 0)  # document properties and zip entries

BANK = "Harborview Community Bank"


def month_end(y: int, m: int) -> date:
    return date(y, m, calendar.monthrange(y, m)[1])


MONTH_ENDS: list[date] = [month_end(2024 + (9 + i) // 12, (9 + i) % 12 + 1) for i in range(N_SIM)]
assert MONTH_ENDS[0] == date(2024, 10, 31) and MONTH_ENDS[-1] == AS_OF


def month_index(d: date) -> int:
    return (d.year - 2024) * 12 + d.month - 10


def rng_for(name: str) -> np.random.Generator:
    return np.random.default_rng([SEED, zlib.crc32(name.encode())])


def sigmoid(x):
    return 1.0 / (1.0 + np.exp(-x))


# --------------------------------------------------------------------------------------------------------------
# Reference data: policy parameters, branches, segments

# Underwriting limits in the consumer credit policy v5.1 (effective 2026-07-01). v5.0 (2025-07-01 .. 2026-06-30)
# allowed the "prior" DTI limits. The policy PDF's table is rendered from this list.
BANDS = [
    {
        "name": "Exceptional",
        "lo": 780,
        "hi": 850,
        "max_dti": 0.45,
        "prior_dti": 0.48,
        "ltv_auto": 1.20,
        "ltv_mtg": 0.97,
    },
    {"name": "Very Good", "lo": 720, "hi": 779, "max_dti": 0.43, "prior_dti": 0.45, "ltv_auto": 1.15, "ltv_mtg": 0.95},
    {"name": "Good", "lo": 670, "hi": 719, "max_dti": 0.40, "prior_dti": 0.43, "ltv_auto": 1.10, "ltv_mtg": 0.90},
    {"name": "Fair", "lo": 620, "hi": 669, "max_dti": 0.36, "prior_dti": 0.40, "ltv_auto": 1.00, "ltv_mtg": 0.85},
    {"name": "Poor", "lo": 300, "hi": 619, "max_dti": 0.30, "prior_dti": 0.36, "ltv_auto": 0.90, "ltv_mtg": 0.80},
]
BAND_NAMES = [b["name"] for b in BANDS]
POLICY_V51_DATE = date(2026, 7, 1)


def band_index(score: int | np.ndarray):
    score = np.asarray(score)
    idx = np.full(score.shape, 4)
    for i, b in enumerate(BANDS):
        idx = np.where(score >= b["lo"], np.minimum(idx, i), idx)
    return idx


# id, name, city, region, type, opened, size weight, staff
BRANCHES = [
    ("BR01", "Harborview Main Street", "Harborview", "Harbor Coast", "full_service", 1962, 2.2, 31),
    ("BR02", "Marina District", "Harborview", "Harbor Coast", "full_service", 1988, 1.6, 19),
    ("BR03", "Lighthouse Point", "Port Alder", "Harbor Coast", "full_service", 1995, 1.0, 14),
    ("BR04", "Seaport Plaza", "Port Alder", "Harbor Coast", "limited_service", 2011, 0.6, 6),
    ("BR05", "Union Square", "Cedar Falls", "Metro", "full_service", 1974, 1.7, 24),
    ("BR06", "Riverside", "Cedar Falls", "Metro", "full_service", 1999, 1.3, 17),
    ("BR07", "Eastgate", "Cedar Falls", "Metro", "full_service", 2004, 1.2, 15),
    ("BR08", "Westfield Mall", "Cedar Falls", "Metro", "limited_service", 2013, 0.7, 6),
    ("BR09", "Northgate", "Cedar Falls", "Metro", "full_service", 2008, 1.0, 13),
    ("BR10", "Millbrook", "Millbrook", "Valley", "full_service", 1968, 1.2, 16),
    ("BR11", "Oakridge", "Oakridge", "Valley", "full_service", 1991, 1.0, 13),
    ("BR12", "Fairview", "Fairview", "Valley", "full_service", 2001, 0.9, 12),
    ("BR13", "Stony Creek", "Stony Creek", "Valley", "limited_service", 2016, 0.5, 5),
    ("BR14", "Lakeside", "Lakeside", "Valley", "full_service", 2006, 0.8, 11),
    ("BR15", "Pinecrest", "Pinecrest", "Highlands", "full_service", 1979, 0.9, 12),
    ("BR16", "Summit Ridge", "Summit Ridge", "Highlands", "full_service", 2009, 0.7, 10),
    ("BR17", "Hollow Springs", "Hollow Springs", "Highlands", "limited_service", 2018, 0.4, 5),
    ("BR18", "Granite Falls", "Granite Falls", "Highlands", "full_service", 1985, 0.8, 11),
]
REGION_STRESS = {"Harbor Coast": 0.9, "Metro": 1.0, "Valley": 1.55, "Highlands": 1.25}

SEGMENTS = [
    # name, share, income median and sigma, score mean/sd, tenure gamma(shape, scale) and cap in years,
    # online banking rate, savings rate, credit card rate
    dict(
        name="Mass Market",
        p=0.44,
        inc=54000,
        sig=0.30,
        sc=668,
        sd=58,
        ts=2.2,
        tsc=4.5,
        tcap=30,
        on=0.72,
        sav=0.38,
        card=0.22,
    ),
    dict(
        name="Mass Affluent",
        p=0.20,
        inc=98000,
        sig=0.28,
        sc=735,
        sd=42,
        ts=2.4,
        tsc=5.0,
        tcap=30,
        on=0.85,
        sav=0.65,
        card=0.42,
    ),
    dict(
        name="Young Adult",
        p=0.16,
        inc=41000,
        sig=0.30,
        sc=645,
        sd=62,
        ts=1.6,
        tsc=1.6,
        tcap=8,
        on=0.93,
        sav=0.30,
        card=0.22,
    ),
    dict(
        name="Retiree",
        p=0.13,
        inc=47000,
        sig=0.35,
        sc=742,
        sd=48,
        ts=3.0,
        tsc=6.0,
        tcap=35,
        on=0.48,
        sav=0.62,
        card=0.30,
    ),
    dict(
        name="Private Client",
        p=0.07,
        inc=235000,
        sig=0.35,
        sc=778,
        sd=32,
        ts=2.5,
        tsc=5.0,
        tcap=30,
        on=0.80,
        sav=0.85,
        card=0.55,
    ),
]

PRODUCTS = ["personal", "auto", "mortgage"]


# --------------------------------------------------------------------------------------------------------------
# Simulation: customers


def make_branches() -> pd.DataFrame:
    return pd.DataFrame(
        [
            dict(
                branch_id=b[0],
                branch_name=b[1],
                city=b[2],
                region=b[3],
                branch_type=b[4],
                opened_year=b[5],
                staff_count=b[7],
            )
            for b in BRANCHES
        ]
    )


def make_customers() -> pd.DataFrame:
    rng = rng_for("customers")
    n = N_CUSTOMERS
    seg = rng.choice(len(SEGMENTS), size=n, p=[s["p"] for s in SEGMENTS])
    income = np.zeros(n)
    score = np.zeros(n)
    tenure = np.zeros(n)
    online = np.zeros(n, dtype=bool)
    for i, s in enumerate(SEGMENTS):
        m = seg == i
        k = int(m.sum())
        income[m] = s["inc"] * np.exp(rng.normal(0, s["sig"], k))
        score[m] = rng.normal(s["sc"], s["sd"], k)
        tenure[m] = np.minimum(rng.gamma(s["ts"], s["tsc"], k), s["tcap"])
        online[m] = rng.random(k) < s["on"]
    score = np.clip(np.round(score), 520, 840).astype(int)
    income = np.round(income / 500) * 500
    w = np.array([b[6] for b in BRANCHES])
    branch = rng.choice(len(BRANCHES), size=n, p=w / w.sum())
    since = [AS_OF - timedelta(days=int(t * 365.25)) for t in tenure]
    df = pd.DataFrame(
        dict(
            home_branch_idx=branch,
            segment=[SEGMENTS[i]["name"] for i in seg],
            seg_idx=seg,
            customer_since=since,
            credit_score=score,
            annual_income=income.astype(int),
            online_banking_enrolled=online.astype(int),
        )
    )
    df = df.sort_values(["customer_since", "seg_idx", "credit_score"], kind="stable").reset_index(drop=True)
    df.insert(0, "customer_id", [f"C{i + 1:05d}" for i in range(n)])
    df["home_branch_id"] = [BRANCHES[i][0] for i in df.home_branch_idx]
    df["region"] = [BRANCHES[i][3] for i in df.home_branch_idx]
    df["credit_score_band"] = [BAND_NAMES[i] for i in band_index(df.credit_score.values)]
    df["tenure_years"] = [round((AS_OF - d).days / 365.25, 1) for d in df.customer_since]
    return df


def make_stress(cust: pd.DataFrame) -> np.ndarray:
    """Monthly financial-stress level per customer, 0..1, over the simulated months.

    A baseline from the credit score, plus stress episodes (job loss, an illness, a rate shock) that ramp up over a
    few months and fade. Utilization, overdrafts, balances and missed payments all follow it, which is the signal
    the prediction templates find.
    """
    rng = rng_for("stress")
    n = len(cust)
    score = cust.credit_score.values
    base = np.clip((650 - score) / 500, 0, 0.25)
    stress = np.repeat(base[:, None], N_SIM, axis=1)
    young = (cust.segment.values == "Young Adult").astype(float)
    region = np.array([REGION_STRESS[r] for r in cust.region.values])
    p_ep = sigmoid(-2.45 + 1.0 * (700 - score) / 80 + 0.3 * young + np.log(region))
    has = rng.random(n) < p_ep
    onset = np.minimum(
        (rng.beta(1.5, 1.0, n) * (N_SIM - 1)).astype(int), N_SIM - 2
    )  # episodes get more common over time
    dur = rng.integers(3, 9, n)
    for c in np.flatnonzero(has):
        peak = min(1.0, 0.30 + 0.22 * (dur[c] - 1))
        for t in range(onset[c], N_SIM):
            k = t - onset[c]
            s = min(1.0, 0.30 + 0.22 * k) if k < dur[c] else max(0.0, peak - 0.25 * (k - dur[c] + 1))
            stress[c, t] = max(stress[c, t], s)
    return stress


# --------------------------------------------------------------------------------------------------------------
# Simulation: accounts, closures and monthly balances


def make_accounts(cust: pd.DataFrame) -> pd.DataFrame:
    rng = rng_for("accounts")
    rows = []
    for c, r in cust.iterrows():
        seg = SEGMENTS[r.seg_idx]
        score = r.credit_score
        affluent = r.segment in ("Mass Affluent", "Private Client")
        tenure_days = max((AS_OF - r.customer_since).days, 30)
        has_chk = rng.random() < 0.9
        has_sav = rng.random() < min(0.95, seg["sav"] * (1.15 if score >= 720 else 0.85))
        card_p = seg["card"] * (1.2 if score >= 700 else 0.9 if score >= 640 else 0.55)
        has_card = rng.random() < card_p
        if not (has_chk or has_sav or has_card):
            has_chk = True

        def open_date(frac: float, min_off: int, since=r.customer_since, span=tenure_days) -> date:
            off = min_off + int(rng.exponential(frac * span))
            return min(since + timedelta(days=off), AS_OF - timedelta(days=3))

        if has_chk:
            if r.segment == "Young Adult" and rng.random() < 0.7:
                prod = "Student Checking"
            elif r.segment == "Private Client" and rng.random() < 0.75:
                prod = "Premier Checking"
            elif affluent and rng.random() < 0.35:
                prod = "Premier Checking"
            else:
                prod = "Everyday Checking"
            rows.append((c, "checking", prod, open_date(0.0, int(rng.integers(0, 15))), np.nan))
        if has_sav:
            prod = "High-Yield Savings" if (affluent and rng.random() < 0.6) or rng.random() < 0.15 else "Basic Savings"
            rows.append((c, "savings", prod, open_date(0.25, 10), np.nan))
        if has_card:
            secured = score < 640 and rng.random() < 0.8
            prod = "Harbor Secured Card" if secured else "Harbor Rewards Card"
            limit_mult = [0.30, 0.24, 0.17, 0.10, 0.05][int(band_index(score))]
            limit = float(np.clip(round(r.annual_income * limit_mult / 500) * 500, 500, 25000))
            if secured:
                limit = float(np.clip(round(limit / 4 / 250) * 250, 500, 2500))
            rows.append((c, "credit_card", prod, open_date(0.35, 20), limit))
    acc = pd.DataFrame(rows, columns=["cust_idx", "account_type", "product_name", "opened_date", "credit_limit"])
    # the opening branch is the home branch for most, another branch for some
    branch = cust.home_branch_idx.values[acc.cust_idx.values].copy()
    other = rng.random(len(acc)) < 0.1
    branch[other] = rng.integers(0, len(BRANCHES), int(other.sum()))
    acc["branch_idx"] = branch
    acc = acc.sort_values(["opened_date", "cust_idx", "account_type"], kind="stable").reset_index(drop=True)
    acc.insert(0, "account_id", [f"A{i + 1:06d}" for i in range(len(acc))])
    acc["open_raw"] = [month_index(d) for d in acc.opened_date]
    acc["open_idx"] = acc.open_raw.clip(lower=0)
    return acc


def simulate_closures(cust: pd.DataFrame, acc: pd.DataFrame, stress: np.ndarray) -> pd.DataFrame:
    """Customer attrition: a customer decides to leave in some month and closes accounts that month or the next."""
    rng = rng_for("closures")
    n = len(cust)
    ci = acc.cust_idx.values
    typ = acc.account_type.values
    oi = acc.open_idx.values
    n_products = np.bincount(ci, minlength=n)
    online = cust.online_banking_enrolled.values
    tenure = cust.tenure_years.values
    close_idx = np.full(len(acc), 10**6)
    left = np.zeros(n, dtype=bool)
    for t in range(OUT_FIRST + 1, N_SIM):
        m_online = np.where(online == 1, 1.0, 1.9)
        m_tenure = np.where(tenure < 2, 2.1, np.where(tenure < 8, 1.0, 0.65))
        m_prod = np.where(n_products <= 1, 1.7, 1.0)
        h = CLOSURE_HAZARD * m_online * m_tenure * m_prod * np.exp(0.7 * stress[:, t - 1])
        has_open = np.zeros(n, dtype=bool)
        dep = (typ != "credit_card") & (oi <= t - 1) & (close_idx > t)
        has_open[ci[dep]] = True
        go = (~left) & has_open & (rng.random(n) < h)
        left |= go
        for a in np.flatnonzero(go[ci] & (oi <= t - 1) & (close_idx > t)):
            later = rng.random() < 0.15 and t + 1 < N_SIM
            close_idx[a] = t + 1 if later else t
        # a savings or card account closed on its own
        solo = (typ != "checking") & (oi <= t - 1) & (close_idx > t) & (rng.random(len(acc)) < 0.0018)
        close_idx[solo] = t
    acc = acc.copy()
    acc["close_idx"] = close_idx
    acc["silent"] = rng.random(len(acc)) < 0.30  # leaves without any warning in the balances
    days = []
    for a in range(len(acc)):
        t = close_idx[a]
        if t >= N_SIM:
            days.append(None)
        else:
            y, m = MONTH_ENDS[t].year, MONTH_ENDS[t].month
            days.append(date(y, m, int(rng.integers(1, calendar.monthrange(y, m)[1] + 1))))
    acc["closed_date"] = days
    return acc


def simulate_balances(cust: pd.DataFrame, acc: pd.DataFrame, stress: np.ndarray):
    """Monthly aggregates per account; returns the balance table and the card utilization path per customer."""
    rng = rng_for("balances")
    na = len(acc)
    ci = acc.cust_idx.values
    typ = acc.account_type.values
    oi = acc.open_idx.values
    raw = acc.open_raw.values
    cl = acc.close_idx.values
    silent = acc.silent.values
    seg = cust.seg_idx.values[ci]
    income_m = cust.annual_income.values[ci] / 12 * 0.8
    band = band_index(cust.credit_score.values[ci])
    limit = acc.credit_limit.fillna(0).values
    is_chk, is_sav, is_card = typ == "checking", typ == "savings", typ == "credit_card"

    seg_level = np.array([0.0, 0.25, -0.35, 0.1, 0.45])[seg]
    chk_level = income_m * np.exp(rng.normal(-0.35, 0.7, na) + seg_level)
    sav_level = income_m * np.exp(rng.normal(0.9, 0.9, na) + seg_level)
    u_mean = np.array([0.07, 0.14, 0.27, 0.42, 0.58])[band]
    u0 = np.clip(rng.normal(u_mean, 0.08, na), 0.01, 0.9)
    noise = np.zeros(na)
    prev = np.zeros(na)
    util_by_cust = np.full((len(cust), N_SIM), np.nan)
    rows = []
    for t in range(N_SIM):
        s_a = stress[ci, t]
        active = (oi <= t) & (t < cl)
        new = active & (raw == t)
        seen_first = active & (oi == t) & ~new
        noise = 0.6 * noise + rng.normal(0, 0.16, na)
        mtc = cl - t  # months to closure
        pre = (~silent) & (mtc >= 1) & (mtc <= 4)
        pre_bal = np.where(pre, 1 - 0.2 * (5 - np.minimum(mtc, 4)), 1.0)
        pre_in = np.where(pre, 1 - 0.16 * (5 - np.minimum(mtc, 4)), 1.0)

        end_chk = chk_level * np.exp(noise) * (1 - 0.55 * s_a) * pre_bal
        in_chk = income_m * rng.uniform(0.85, 1.1, na) * (1 - 0.35 * s_a) * pre_in
        lvl = sav_level * 1.004**t
        end_sav = lvl * np.exp(0.5 * noise) * (1 - 0.30 * s_a) * pre_bal
        in_sav = lvl * rng.uniform(0.02, 0.08, na) * (1 - 0.4 * s_a) * pre_in
        util_t = np.clip(u0 + 0.45 * s_a + 0.08 * noise + 0.003 * t, 0.0, 0.99) * pre_bal
        end_card = util_t * limit
        spend = limit * rng.uniform(0.07, 0.19, na) * (1 - 0.2 * s_a) * pre_in
        end = np.where(is_chk, end_chk, np.where(is_sav, end_sav, end_card))
        p0 = np.where(new, 0.0, np.where(seen_first, end, prev))
        # money is conserved: end = start + inflow - outflow (for a card the payments are the inflow)
        pay_min = np.where(p0 > 0, 0.03 * p0 + 25, 0.0)
        purchases = np.maximum(spend, end - p0 + pay_min)
        payments = p0 + purchases - end
        inflow_dep = np.where(is_chk, in_chk, in_sav)
        inflow_dep = np.where(new, end * 1.15, inflow_dep)
        outflow_dep = p0 + inflow_dep - end
        neg = outflow_dep < 0
        inflow_dep = np.where(neg, inflow_dep - outflow_dep, inflow_dep)
        outflow_dep = np.where(neg, 0.0, outflow_dep)
        inflow = np.where(is_card, payments, inflow_dep)
        outflow = np.where(is_card, purchases, outflow_dep)
        util = np.where(is_card, end / np.maximum(limit, 1), np.nan)
        # overdrafts happen on checking accounts
        ratio = end_chk / np.maximum(income_m, 1)
        lam = 0.02 + 0.55 * np.exp(-4 * ratio) + 1.8 * s_a + 0.002 * t + np.where(pre, 0.25, 0.0)
        od = np.where(is_chk, rng.poisson(lam), 0)
        neg_bal = is_chk & (od > 0) & (rng.random(na) < 0.35)
        end_rep = np.where(neg_bal, -rng.uniform(5, 180, na), end)
        txn_lam = np.where(is_chk, 22 + 10 * np.minimum(income_m / 6000, 2), np.where(is_sav, 3.2, 15.0))
        txn = rng.poisson(txn_lam)
        avg_bal = (p0 + end_rep) / 2 * rng.uniform(0.93, 1.07, na)
        avg_bal = np.where(neg_bal, np.minimum(avg_bal, end_rep * 0.5), avg_bal)
        sel = active & is_card
        util_by_cust[ci[sel], t] = util[sel]
        if t >= OUT_FIRST:
            for a in np.flatnonzero(active):
                rows.append(
                    (
                        a,
                        MONTH_ENDS[t],
                        round(float(end_rep[a]), 2),
                        round(float(avg_bal[a]), 2),
                        round(float(inflow[a]), 2),
                        round(float(outflow[a]), 2),
                        int(txn[a]),
                        int(od[a]),
                        round(float(util[a]), 4) if is_card[a] else None,
                    )
                )
        prev = np.where(active, end, 0.0)
    bal = pd.DataFrame(
        rows,
        columns=[
            "acc_idx",
            "month_end",
            "ending_balance",
            "average_daily_balance",
            "total_inflows",
            "total_outflows",
            "transaction_count",
            "overdraft_count",
            "utilization_ratio",
        ],
    )
    return bal, util_by_cust


# --------------------------------------------------------------------------------------------------------------
# Simulation: loans and their payments

CLOSURE_HAZARD = 0.0050  # monthly probability that a typical customer leaves
# share of customers with a mortgage, auto and personal loan, by segment
LOAN_P = {
    "Mass Market": (0.12, 0.31, 0.26),
    "Mass Affluent": (0.26, 0.26, 0.17),
    "Young Adult": (0.02, 0.29, 0.36),
    "Retiree": (0.14, 0.14, 0.17),
    "Private Client": (0.36, 0.17, 0.06),
}
TERMS = {
    "personal": ([24, 36, 48, 60], [0.2, 0.35, 0.25, 0.2]),
    "auto": ([48, 60, 72], [0.25, 0.45, 0.30]),
    "mortgage": ([180, 240, 360], [0.15, 0.10, 0.75]),
}
RATE_PREMIUM = {
    "personal": [0.0, 1.6, 3.6, 7.0, 11.0],
    "auto": [0.0, 0.8, 2.0, 4.6, 8.2],
    "mortgage": [0.0, 0.15, 0.35, 0.75, 1.3],
}
RATE_BASE = {"personal": 8.5, "auto": 5.2}
MORTGAGE_BASE = {2013: 4.0, 2014: 4.1, 2015: 4.0, 2016: 3.8, 2017: 4.1, 2018: 4.6, 2019: 4.0, 2020: 3.2, 2021: 3.0}
MORTGAGE_BASE |= {2022: 5.3, 2023: 6.8, 2024: 6.6, 2025: 6.3, 2026: 6.1}
# logit of a missed installment, by product, then the effects of credit score, stress, utilization and time
MISS_BASE = {"personal": -4.6, "auto": -4.7, "mortgage": -5.6}
PREPAY = {"personal": 0.009, "auto": 0.006, "mortgage": 0.0035}  # monthly early payoff of a current loan
CHARGE_OFF_ARREARS = {"personal": 4, "auto": 4, "mortgage": 6}  # missed installments before charge-off (120/180 days)
RECOVERY_RATE = 0.15


def due_on(orig: date, k: int) -> date:
    idx = orig.year * 12 + orig.month - 1 + k
    return date(idx // 12, idx % 12 + 1, min(orig.day, 28))


def amortized_payment(principal: float, rate_pct: float, term: int) -> float:
    r = rate_pct / 1200
    return round(principal * r / (1 - (1 + r) ** -term), 2)


def amort_balance(principal: float, rate_pct: float, term: int, paid: int) -> float:
    if paid >= term:
        return 0.0
    r = rate_pct / 1200
    return principal * ((1 + r) ** term - (1 + r) ** paid) / ((1 + r) ** term - 1)


def rand_date(rng, lo: date, hi: date) -> date:
    return lo + timedelta(days=int(rng.integers(0, (hi - lo).days + 1)))


def originate_loans(cust: pd.DataFrame, rng) -> list[dict]:
    recs = []
    for c, r in cust.iterrows():
        low = r.credit_score < 640
        pm, pa, pp = LOAN_P[r.segment]
        for prod, p in (
            ("mortgage", pm * (0.4 if low else 1.0)),
            ("auto", pa),
            ("personal", pp * (1.5 if low else 1.0)),
        ):
            if rng.random() >= p:
                continue
            in_window = rng.random() < {"personal": 0.60, "auto": 0.55, "mortgage": 0.12}[prod]
            if in_window:
                lo, hi = WINDOW_START, date(2026, 9, 15)
            else:
                lo = {"personal": date(2023, 1, 1), "auto": date(2021, 10, 1), "mortgage": date(2013, 1, 1)}[prod]
                hi = date(2025, 3, 31)
            lo = max(lo, r.customer_since + timedelta(days=30))
            if lo > hi:
                continue
            orig = rand_date(rng, lo, hi)
            terms, tp = TERMS[prod]
            term = int(rng.choice(terms, p=tp))
            score_o = int(np.clip(r.credit_score + round(rng.normal(0, 18)), 500, 850))
            band = int(band_index(score_o))
            if prod == "personal":
                principal = float(np.clip(np.exp(rng.normal(np.log(11000), 0.6)), 2500, 50000))
                principal = min(principal, max(2500, r.annual_income * 0.7))
                principal = round(principal / 100) * 100
            elif prod == "auto":
                principal = round(float(np.clip(np.exp(rng.normal(np.log(28000), 0.45)), 8000, 75000)) / 100) * 100
            else:
                principal = round(float(np.clip(r.annual_income * rng.uniform(1.8, 4.2), 90000, 750000)) / 1000) * 1000
            base = MORTGAGE_BASE[orig.year] if prod == "mortgage" else RATE_BASE[prod]
            rate = round(max(2.5, base + RATE_PREMIUM[prod][band] + rng.normal(0, 0.25)), 2)
            b = BANDS[band]
            cap = b["prior_dti"] if orig < POLICY_V51_DATE else b["max_dti"]
            if rng.random() < 0.025:
                dti = cap + rng.uniform(0.01, 0.05)  # an approved exception
            else:
                dti = rng.normal(cap - 0.06, 0.07)
                for _ in range(30):
                    if 0.10 <= dti <= cap:
                        break
                    dti = rng.normal(cap - 0.06, 0.07)
                dti = float(np.clip(dti, 0.10, cap))
            ltv = None
            if prod != "personal":
                lcap = b["ltv_auto"] if prod == "auto" else b["ltv_mtg"]
                ltv = round(float(np.clip(rng.normal(lcap - (0.2 if prod == "auto" else 0.12), 0.12), 0.35, lcap)), 2)
            branch = int(r.home_branch_idx) if rng.random() < 0.85 else int(rng.integers(0, len(BRANCHES)))
            recs.append(
                dict(
                    cust_idx=c,
                    branch_idx=branch,
                    product=prod,
                    principal=float(principal),
                    interest_rate=rate,
                    term_months=term,
                    origination_date=orig,
                    monthly_payment=amortized_payment(principal, rate, term),
                    credit_score_at_origination=score_o,
                    credit_band=BAND_NAMES[band],
                    dti_at_origination=round(float(dti), 2),
                    ltv_at_origination=ltv,
                )
            )
    return recs


def simulate_loans(cust: pd.DataFrame, stress: np.ndarray, util: np.ndarray):
    """Originate loans, then play each loan's installments month by month with a roll-rate model.

    A current loan misses its next installment with a probability that rises with a weaker credit score, the
    customer's stress, card utilization and the passage of time; an installment that is missed can be cured (all the
    arrears paid on a later day) or roll to the next due date; the loan is charged off after four missed
    installments (six for a mortgage). Burn-in months (2024-10 to 2025-03) are simulated, but only installments
    still open at the window's start are published.
    """
    rng = rng_for("loans")
    recs = originate_loans(cust, rng)
    loans, pays = [], []
    for li, rec in enumerate(recs):
        c, prod, orig, term = rec["cust_idx"], rec["product"], rec["origination_date"], rec["term_months"]
        score = int(cust.credit_score.iat[c])
        k = max(1, (SIM_START.year * 12 + SIM_START.month - 1) - (orig.year * 12 + orig.month - 1))
        while k > 1 and due_on(orig, k - 1) >= SIM_START:
            k -= 1
        while due_on(orig, k) < SIM_START:
            k += 1
        n_pre = k - 1
        if k > term:
            continue  # matured before the simulation began
        arrears: list[dict] = []
        rows: list[dict] = []
        status, closed, trouble, stale = "active", None, False, 0
        while True:
            d = due_on(orig, k)
            in_term = k <= term
            if d > AS_OF or (not in_term and not arrears):
                break
            mi = month_index(d)
            sc = float(stress[c, mi])
            u = util[c, mi]
            u = 0.3 if np.isnan(u) else float(u)
            row = None
            if arrears:
                if len(arrears) + stale >= CHARGE_OFF_ARREARS[prod]:
                    status, closed = "charged_off", d
                    break
                p_cure = 0.04 + 0.52 * 0.68 ** (len(arrears) + stale - 1) * (1 - 0.35 * sc)
                if rng.random() < p_cure:
                    cure = d + timedelta(days=int(rng.integers(0, 28)))
                    for a in arrears:
                        a["paid_date"] = cure
                    arrears, stale = [], 0
                    if in_term:
                        row = dict(k=k, due_date=d, paid_date=cure)
                elif in_term:
                    row = dict(k=k, due_date=d, paid_date=None)
                    arrears.append(row)
                else:
                    stale += 1  # past the final installment, the arrears keep aging
            else:
                z = (
                    MISS_BASE[prod]
                    + 0.8 * (700 - score) / 80
                    + 2.6 * sc
                    + 1.4 * (u - 0.3)
                    + 0.018 * mi
                    + (0.7 if trouble else 0.0)
                )
                z_late = -3.3 + 0.5 * (700 - score) / 80 + 1.8 * sc + (0.6 if trouble else 0.0)
                if rng.random() < sigmoid(z):
                    row = dict(k=k, due_date=d, paid_date=None)
                    arrears.append(row)
                    trouble = True
                elif rng.random() < sigmoid(z_late):
                    row = dict(k=k, due_date=d, paid_date=d + timedelta(days=int(rng.integers(1, 26))))
                    trouble = True
                else:
                    row = dict(k=k, due_date=d, paid_date=d - timedelta(days=int(rng.integers(0, 4))))
            if row is not None:
                rows.append(row)
            settled = row is not None and row["paid_date"] is not None and not arrears
            if in_term and k == term and settled:
                status, closed = "paid_off", row["paid_date"]
                break
            if in_term and k < term and settled and rng.random() < PREPAY[prod]:
                status, closed = "paid_off", row["paid_date"] + timedelta(days=int(rng.integers(2, 21)))
                break
            k += 1
        for row in rows:
            if row["paid_date"] is not None and row["paid_date"] > AS_OF:
                row["paid_date"] = None  # not yet received at the reporting date
        if closed is not None and closed > AS_OF:
            status, closed = "active", None
        rec = dict(rec, status=status, closed_date=closed, n_pre=n_pre, tmp=li)
        rec["maturity_date"] = due_on(orig, term)
        loans.append(rec)
        for row in rows:
            row["tmp"] = li
            row["scheduled_amount"] = rec["monthly_payment"]
            if row["paid_date"] is not None:
                row["days_past_due"] = max(0, (row["paid_date"] - row["due_date"]).days)
            else:
                end = closed if status == "charged_off" else AS_OF
                row["days_past_due"] = max(0, (end - row["due_date"]).days)
            row["paid_amount"] = rec["monthly_payment"] if row["paid_date"] is not None else 0.0
        pays.extend(rows)
    ln = pd.DataFrame(loans)
    ln = ln[ln.closed_date.isna() | (ln.closed_date >= WINDOW_START)]
    ln = ln.sort_values(["origination_date", "cust_idx", "product"], kind="stable").reset_index(drop=True)
    ln.insert(0, "loan_id", [f"L{i + 1:05d}" for i in range(len(ln))])
    pm = pd.DataFrame(pays)
    pm = pm.merge(ln[["tmp", "loan_id"]], on="tmp", how="inner").drop(columns="tmp")
    ln = ln.drop(columns="tmp")
    return ln, pm


# --------------------------------------------------------------------------------------------------------------
# Simulation: card disputes


def simulate_disputes(cust: pd.DataFrame, acc: pd.DataFrame, stress: np.ndarray) -> pd.DataFrame:
    rng = rng_for("disputes")
    n = len(cust)
    has_card = np.zeros(n, dtype=bool)
    has_card[acc.cust_idx.values[acc.account_type.values == "credit_card"]] = True
    young = (cust.segment.values == "Young Adult").astype(float)
    online = cust.online_banking_enrolled.values
    prior = np.zeros(n)
    by_cust = {c: g for c, g in acc.groupby("cust_idx")}
    reasons = [
        "unauthorized_transaction",
        "merchant_dispute",
        "duplicate_charge",
        "service_not_received",
        "incorrect_amount",
    ]
    rp = [0.34, 0.26, 0.14, 0.15, 0.11]
    rows = []
    for t in range(OUT_FIRST, N_SIM):
        lam = (
            DISPUTE_BASE
            * np.where(has_card, 2.4, 1.0)
            * (1 + 0.5 * young)
            * (1 + 0.6 * np.minimum(prior, 3))
            * (1 + 0.8 * stress[:, t])
            * np.where(online == 1, 1.15, 1.0)
        )
        counts = rng.poisson(lam)
        for c in np.flatnonzero(counts):
            g = by_cust.get(c)
            if g is None:
                continue
            m_end = MONTH_ENDS[t]
            m_start = date(m_end.year, m_end.month, 1)
            for _ in range(int(counts[c])):
                ok = g[(g.opened_date <= m_end) & (g.close_idx > t)]
                if ok.empty:
                    continue
                cards = ok[ok.account_type == "credit_card"]
                chk = ok[ok.account_type == "checking"]
                if not cards.empty and (chk.empty or rng.random() < 0.65):
                    a = cards.iloc[int(rng.integers(0, len(cards)))]
                    channel = "credit_card"
                elif not chk.empty:
                    a = chk.iloc[int(rng.integers(0, len(chk)))]
                    channel = "debit_card"
                else:
                    continue
                d = m_start + timedelta(days=int(rng.integers(0, m_end.day)))
                d = max(d, a.opened_date + timedelta(days=1))
                d = min(d, m_end)
                amount = round(float(np.clip(np.exp(rng.normal(np.log(90), 0.9)), 8, 4500)), 2)
                if d > date(2026, 9, 8):
                    status = "open"
                else:
                    status = "customer_credited" if rng.random() < 0.64 else "denied"
                rows.append(
                    dict(
                        account_id=a.account_id,
                        cust_idx=c,
                        dispute_date=d,
                        amount=amount,
                        reason=reasons[int(rng.choice(5, p=rp))],
                        channel=channel,
                        status=status,
                    )
                )
                prior[c] += 1
    df = pd.DataFrame(rows).sort_values(["dispute_date", "account_id"], kind="stable").reset_index(drop=True)
    df.insert(0, "dispute_id", [f"D{i + 1:05d}" for i in range(len(df))])
    return df


DISPUTE_BASE = 0.0095


# --------------------------------------------------------------------------------------------------------------
# The world


class World:
    """Every table as a DataFrame in the published column order, plus the hidden working columns."""

    def __init__(self) -> None:
        self.branches = make_branches()
        cust = make_customers()
        stress = make_stress(cust)
        acc = make_accounts(cust)
        acc = simulate_closures(cust, acc, stress)
        bal, util = simulate_balances(cust, acc, stress)
        loans, pays = simulate_loans(cust, stress, util)
        disp = simulate_disputes(cust, acc, stress)
        self.cust_full, self.acc_full, self.loans_full, self.pays_full = cust, acc, loans, pays
        self.stress = stress
        self.disputes_full = disp
        # published tables
        self.customers = cust[
            [
                "customer_id",
                "home_branch_id",
                "segment",
                "region",
                "customer_since",
                "tenure_years",
                "credit_score",
                "credit_score_band",
                "annual_income",
                "online_banking_enrolled",
            ]
        ].copy()
        a = acc.copy()
        a["customer_id"] = cust.customer_id.values[a.cust_idx.values]
        a["branch_id"] = [BRANCHES[i][0] for i in a.branch_idx]
        a["status"] = np.where(a.closed_date.isna(), "open", "closed")
        self.accounts = a[
            [
                "account_id",
                "customer_id",
                "branch_id",
                "account_type",
                "product_name",
                "opened_date",
                "closed_date",
                "status",
                "credit_limit",
            ]
        ].copy()
        bal = bal.copy()
        bal["account_id"] = acc.account_id.values[bal.acc_idx.values]
        bal.insert(0, "balance_id", range(1, len(bal) + 1))
        self.monthly_balances = bal[
            [
                "balance_id",
                "account_id",
                "month_end",
                "ending_balance",
                "average_daily_balance",
                "total_inflows",
                "total_outflows",
                "transaction_count",
                "overdraft_count",
                "utilization_ratio",
            ]
        ].copy()
        ln = loans.copy()
        ln["customer_id"] = cust.customer_id.values[ln.cust_idx.values]
        ln["branch_id"] = [BRANCHES[i][0] for i in ln.branch_idx]
        self.loans = ln[
            [
                "loan_id",
                "customer_id",
                "branch_id",
                "product",
                "principal",
                "interest_rate",
                "term_months",
                "monthly_payment",
                "origination_date",
                "maturity_date",
                "credit_score_at_origination",
                "credit_band",
                "dti_at_origination",
                "ltv_at_origination",
                "status",
                "closed_date",
            ]
        ].copy()
        p = pays.copy()
        p = p[(p.due_date >= WINDOW_START) | p.paid_date.isna() | (p.paid_date >= WINDOW_START)]
        p = p.sort_values(["due_date", "loan_id"], kind="stable").reset_index(drop=True)
        p.insert(0, "payment_id", [f"P{i + 1:06d}" for i in range(len(p))])
        p["payment_status"] = np.where(
            p.paid_date.isna(), "unpaid", np.where(p.days_past_due > 0, "paid_late", "paid_on_time")
        )
        p["installment_number"] = p.k
        self.loan_payments = p[
            [
                "payment_id",
                "loan_id",
                "installment_number",
                "due_date",
                "scheduled_amount",
                "paid_amount",
                "paid_date",
                "days_past_due",
                "payment_status",
            ]
        ].copy()
        d = disp.copy()
        d["customer_id"] = cust.customer_id.values[d.cust_idx.values]
        self.card_disputes = d[
            ["dispute_id", "account_id", "customer_id", "dispute_date", "amount", "reason", "channel", "status"]
        ].copy()


# --------------------------------------------------------------------------------------------------------------
# Metrics: delinquency, balances, provisions (computed from the tables, never typed in)

# Allowance for loan losses: a reserve rate on the outstanding balance, by delinquency bucket, product and credit band.
RESERVE_CURRENT = {"personal": 0.032, "auto": 0.016, "mortgage": 0.0035}
BAND_RESERVE_MULT = [0.4, 0.6, 1.0, 1.8, 3.0]
BUCKET_RATE = [(90, 0.70), (60, 0.35), (30, 0.12)]  # 90+, 60-89, 30-59 days past due
LOSS_SEVERITY = {"personal": 1.0, "auto": 0.7, "mortgage": 0.4}


class Metrics:
    """Point-in-time views of the loan book, built from the full loan and payment frames."""

    def __init__(self, w: World) -> None:
        self.w = w
        L = w.loans_full.copy()
        L["segment"] = w.cust_full.segment.values[L.cust_idx.values]
        L["band_idx"] = [BAND_NAMES.index(b) for b in L.credit_band]
        L["orig_ts"] = pd.to_datetime(L.origination_date)
        L["closed_ts"] = pd.to_datetime(L.closed_date)
        self.L = L.set_index("loan_id", drop=False).rename_axis(None)
        P = w.pays_full.copy()
        P["due_ts"] = pd.to_datetime(P.due_date)
        P["paid_ts"] = pd.to_datetime(P.paid_date)
        self.P = P

    def book(self, e: date) -> pd.DataFrame:
        """Active loans at e with their worst days past due, balance and delinquency bucket."""
        ts = pd.Timestamp(e)
        L, P = self.L, self.P
        act = L[(L.orig_ts <= ts) & (L.closed_ts.isna() | (L.closed_ts > ts))].copy()
        due = P[P.due_ts <= ts]
        open_ = due[due.paid_ts.isna() | (due.paid_ts > ts)]
        worst = (ts - open_.due_ts).dt.days.groupby(open_.loan_id).max()
        act["dpd"] = act.loan_id.map(worst).fillna(0).astype(int)
        paid = due[due.paid_ts.notna() & (due.paid_ts <= ts)].groupby("loan_id").size()
        act["k_paid"] = act.n_pre + act.loan_id.map(paid).fillna(0).astype(int)
        act["balance"] = [
            amort_balance(r.principal, r.interest_rate, r.term_months, r.k_paid) for r in act.itertuples()
        ]
        act["delinquent_30"] = act.dpd >= 30
        act["delinquent_90"] = act.dpd >= 90
        return act

    def delinquency(self, e: date) -> pd.DataFrame:
        b = self.book(e)
        g = b.groupby("product")
        out = pd.DataFrame(
            dict(
                active_loans=g.size(),
                loans_30plus=g.delinquent_30.sum(),
                loans_90plus=g.delinquent_90.sum(),
                balance=g.balance.sum(),
            )
        ).reindex(PRODUCTS)
        out["rate_30"] = out.loans_30plus / out.active_loans
        out["rate_90"] = out.loans_90plus / out.active_loans
        return out

    def allowance(self, b: pd.DataFrame) -> pd.Series:
        rate = []
        for r in b.itertuples():
            bucket = next((x for d, x in BUCKET_RATE if r.dpd >= d), None)
            if bucket is not None:
                rate.append(bucket * LOSS_SEVERITY[r.product])
            else:
                mult = 2.0 if r.dpd > 0 else 1.0
                rate.append(RESERVE_CURRENT[r.product] * BAND_RESERVE_MULT[r.band_idx] * mult)
        return pd.Series(np.array(rate) * b.balance.values, index=b.index)

    def charge_offs(self, lo: date, hi: date) -> pd.DataFrame:
        """Loans charged off in (lo, hi] with the balance owed at charge-off."""
        L = self.L
        co = L[
            (L.status == "charged_off") & (L.closed_ts > pd.Timestamp(lo)) & (L.closed_ts <= pd.Timestamp(hi))
        ].copy()
        ts = co.closed_ts
        P = self.P[self.P.loan_id.isin(co.loan_id)]
        paid = P[P.paid_ts.notna()].merge(co[["loan_id", "closed_ts"]], on="loan_id")
        paid = paid[paid.paid_ts <= paid.closed_ts].groupby("loan_id").size()
        co["k_paid"] = co.n_pre + co.loan_id.map(paid).fillna(0).astype(int)
        co["balance"] = [amort_balance(r.principal, r.interest_rate, r.term_months, r.k_paid) for r in co.itertuples()]
        del ts
        return co

    def delinquency_sheet(self) -> pd.DataFrame:
        rows = []
        for e in MONTH_ENDS[OUT_FIRST:]:
            d = self.delinquency(e)
            for prod in PRODUCTS:
                r = d.loc[prod]
                rows.append(
                    (
                        e,
                        prod,
                        int(r.active_loans),
                        int(r.loans_30plus),
                        int(r.loans_90plus),
                        round(float(r.rate_30), 4),
                        round(float(r.rate_90), 4),
                    )
                )
        return pd.DataFrame(
            rows,
            columns=[
                "month_end",
                "product",
                "active_loans",
                "loans_30plus",
                "loans_90plus",
                "delinquency_rate_30plus",
                "delinquency_rate_90plus",
            ],
        )

    def provisions_sheet(self) -> pd.DataFrame:
        qends = [
            date(2025, 3, 31),
            date(2025, 6, 30),
            date(2025, 9, 30),
            date(2025, 12, 31),
            date(2026, 3, 31),
            date(2026, 6, 30),
            date(2026, 9, 30),
        ]
        books = {}
        for e in qends:
            b = self.book(e)
            b["allowance"] = self.allowance(b)
            books[e] = b
        segs = [s["name"] for s in SEGMENTS]
        rows = []
        for i in range(1, len(qends)):
            e, prev = qends[i], qends[i - 1]
            b, bp = books[e], books[prev]
            co = self.charge_offs(prev, e)
            for seg in segs:
                bs, bps, cs = b[b.segment == seg], bp[bp.segment == seg], co[co.segment == seg]
                gross = round(float(cs.balance.sum()), 0)
                recov = round(gross * RECOVERY_RATE, 0)
                net = gross - recov
                allowance = round(float(bs.allowance.sum()), 0)
                expense = round(allowance - round(float(bps.allowance.sum()), 0) + net, 0)
                outstanding = round(float(bs.balance.sum()), 0)
                rows.append(
                    (
                        e,
                        seg,
                        len(bs),
                        int(outstanding),
                        int(allowance),
                        int(expense),
                        int(gross),
                        int(recov),
                        int(net),
                        round(allowance / outstanding, 4) if outstanding else 0.0,
                    )
                )
        return pd.DataFrame(
            rows,
            columns=[
                "quarter_end",
                "segment",
                "active_loans",
                "outstanding_balance",
                "allowance",
                "provision_expense",
                "gross_charge_offs",
                "recoveries",
                "net_charge_offs",
                "coverage_ratio",
            ],
        )


def pct(x: float, nd: int = 1) -> str:
    return f"{x * 100:.{nd}f}%"


def money(x: float) -> str:
    return f"${x:,.0f}"


def millions(x: float) -> str:
    return f"${x / 1e6:,.1f} million"


# --------------------------------------------------------------------------------------------------------------
# Writers: CSV, XLSX and byte-stable zip containers


def normalize_zip_bytes(data: bytes) -> bytes:
    """Rewrite an OOXML container with a fixed timestamp and no host metadata, so the same content gives the same bytes."""
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        items = [(i.filename, zf.read(i.filename)) for i in zf.infolist()]
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as out:
        for name, blob in items:
            if name.endswith((".xlsx", ".docx")) and name.startswith(("ppt/embeddings/", "word/embeddings/")):
                blob = normalize_zip_bytes(blob)
            if name == "docProps/core.xml":  # openpyxl stamps the save time into dcterms:modified
                stamp = FIXED_TIME.strftime("%Y-%m-%dT%H:%M:%SZ").encode()
                blob = re.sub(
                    rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)", rb"\g<1>" + stamp + rb"\g<2>", blob
                )
            zi = zipfile.ZipInfo(name, date_time=(2026, 9, 30, 12, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o644 << 16
            zi.create_system = 3
            out.writestr(zi, blob)
    return buf.getvalue()


def normalize_zip(path: Path) -> None:
    path.write_bytes(normalize_zip_bytes(path.read_bytes()))


def clean_numbers(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    for c in df.columns:
        if pd.api.types.is_float_dtype(df[c]):
            df[c] = df[c] + 0.0  # turns -0.0 into 0.0
    return df


def write_csv(df: pd.DataFrame, name: str) -> Path:
    path = FILES / "banking" / f"{name}.csv"
    path.parent.mkdir(parents=True, exist_ok=True)
    clean_numbers(df).to_csv(path, index=False, lineterminator="\n")
    return path


def write_xlsx(path: Path, sheets: dict[str, pd.DataFrame], widths: dict[str, int] | None = None) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment
    from openpyxl.styles import Font
    from openpyxl.styles import PatternFill

    wb = Workbook()
    wb.remove(wb.active)
    for name, df in sheets.items():
        ws = wb.create_sheet(name)
        ws.append(list(df.columns))
        for row in df.itertuples(index=False):
            ws.append([v.item() if hasattr(v, "item") else v for v in row])
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="12355B")
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for i, col in enumerate(df.columns, start=1):
            letter = ws.cell(row=1, column=i).column_letter
            ws.column_dimensions[letter].width = max(12, min(30, len(col) + 3))
            if col.endswith(("_date", "_end")):
                for r in range(2, len(df) + 2):
                    ws.cell(row=r, column=i).number_format = "yyyy-mm-dd"
            elif "rate" in col or col == "coverage_ratio":
                for r in range(2, len(df) + 2):
                    ws.cell(row=r, column=i).number_format = "0.00%"
            elif col in (
                "outstanding_balance",
                "allowance",
                "provision_expense",
                "gross_charge_offs",
                "recoveries",
                "net_charge_offs",
            ):
                for r in range(2, len(df) + 2):
                    ws.cell(row=r, column=i).number_format = "#,##0"
        ws.freeze_panes = "A2"
    p = wb.properties
    p.creator = f"{BANK} (synthetic)"
    p.lastModifiedBy = f"{BANK} (synthetic)"
    p.title = "Credit Risk Report, Q3 2026"
    p.created = FIXED_TIME
    p.modified = FIXED_TIME
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    normalize_zip(path)


# --------------------------------------------------------------------------------------------------------------
# Content: a small Markdown subset (front matter, ## and ### headings, paragraphs, bullets, numbered lists, pipe
# tables, [[table:name]] and [[bars:name]] directives, > notes) with {{placeholders}} filled from computed facts


def parse_front(text: str) -> tuple[dict[str, str], str]:
    meta: dict[str, str] = {}
    if text.startswith("---\n"):
        head, _, body = text[4:].partition("\n---\n")
        for line in head.splitlines():
            k, _, v = line.partition(":")
            meta[k.strip()] = v.strip()
        return meta, body
    return meta, text


def fill(text: str, ctx: dict) -> str:
    def sub(m: re.Match) -> str:
        key = m.group(1).strip()
        if key not in ctx:
            raise KeyError(f"placeholder {{{{{key}}}}} has no value")
        return str(ctx[key])

    return re.sub(r"\{\{([^}]+)\}\}", sub, text)


def load_md(name: str, ctx: dict | None = None) -> tuple[dict[str, str], list[tuple]]:
    meta, body = parse_front((CONTENT / name).read_text(encoding="utf-8"))
    if ctx is not None:
        body = fill(body, ctx)
        meta = {k: fill(v, ctx) for k, v in meta.items()}
    blocks: list[tuple] = []
    lines = body.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1
        elif line.startswith("### "):
            blocks.append(("h3", line[4:].strip()))
            i += 1
        elif line.startswith("## "):
            blocks.append(("h2", line[3:].strip()))
            i += 1
        elif line.startswith("[[") and line.rstrip().endswith("]]"):
            kind, _, ref = line.strip()[2:-2].partition(":")
            blocks.append((kind, ref))
            i += 1
        elif line.startswith(">"):
            notes = []
            while i < len(lines) and lines[i].startswith(">"):
                notes.append(lines[i][1:].strip())
                i += 1
            blocks.append(("note", " ".join(notes)))
        elif line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(set(c) <= set("-: ") for c in cells):
                    rows.append(cells)
                i += 1
            blocks.append(("rows", rows))
        elif re.match(r"^(- |\d+\. )", line):
            ordered = bool(re.match(r"^\d+\. ", line))
            items: list[str] = []
            while i < len(lines) and (re.match(r"^(- |\d+\. )", lines[i]) or lines[i].startswith("  ")):
                if lines[i].startswith("  "):
                    items[-1] += " " + lines[i].strip()
                else:
                    items.append(re.sub(r"^(- |\d+\. )", "", lines[i]).strip())
                i += 1
            blocks.append(("ol" if ordered else "ul", items))
        else:
            para = []
            while i < len(lines) and lines[i].strip() and not re.match(r"^(#|\||>|- |\d+\. |\[\[)", lines[i]):
                para.append(lines[i].strip())
                i += 1
            text = " ".join(para)
            kind = "caption" if text.startswith("**Table") and text.endswith("**") else "p"
            blocks.append((kind, text))
    return meta, blocks


def runs(text: str) -> list[tuple[str, bool, bool]]:
    """Split **bold** and *italic* spans."""
    out: list[tuple[str, bool, bool]] = []
    for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if not part:
            continue
        if part.startswith("**"):
            out.append((part[2:-2], True, False))
        elif part.startswith("*") and len(part) > 2:
            out.append((part[1:-1], False, True))
        else:
            out.append((part, False, False))
    return out


def markup(text: str) -> str:
    from xml.sax.saxutils import escape

    return "".join(("<b>%s</b>" if b else "<i>%s</i>" if i else "%s") % escape(t) for t, b, i in runs(text))


def score_range(b: dict) -> str:
    if b["hi"] >= 850:
        return f"{b['lo']} and above"
    if b["lo"] <= 300:
        return f"{b['hi']} and below"
    return f"{b['lo']} to {b['hi']}"


def band_tables() -> dict[str, tuple[list[str], list[list[str]]]]:
    bands = (
        ["Credit band", "Credit score", "Maximum DTI", "Maximum auto LTV", "Maximum mortgage LTV", "Pricing tier"],
        [
            [b["name"], score_range(b), pct(b["max_dti"], 0), pct(b["ltv_auto"], 0), pct(b["ltv_mtg"], 0), "ABCDE"[i]]
            for i, b in enumerate(BANDS)
        ],
    )
    sup = (
        ["Credit band", "Maximum DTI, version 5.0", "Maximum DTI, version 5.1 (current)", "Change"],
        [
            [
                b["name"],
                pct(b["prior_dti"], 0),
                pct(b["max_dti"], 0),
                f"-{round((b['prior_dti'] - b['max_dti']) * 100)} points",
            ]
            for b in BANDS
        ],
    )
    return {"bands": bands, "superseded": sup}


# --------------------------------------------------------------------------------------------------------------
# Facts: every number a document or the answer key quotes, computed from the tables

PRODUCT_LABEL = {"personal": "Personal loans", "auto": "Auto loans", "mortgage": "Mortgages"}
LOW_BANDS = ("Fair", "Poor")


def pts(delta: float) -> str:
    return "0.0 pts" if abs(delta) < 0.0005 else f"{delta * 100:+.1f} pts"


def money_signed(x: float) -> str:
    return f"-${abs(x):,.0f}" if x < 0 else f"${x:,.0f}"


def money_text(x: float) -> str:
    return f"a net release of ${abs(x):,.0f}" if x < 0 else f"${x:,.0f}"


def loan_label(m: Metrics, anchor: date, horizon: int = 90) -> pd.DataFrame:
    """The loan-default template's label: any installment due in the next 90 days ends up 30 or more days past due."""
    book = m.book(anchor)
    lo, hi = pd.Timestamp(anchor), pd.Timestamp(anchor + timedelta(days=horizon))
    win = m.P[(m.P.due_ts >= lo) & (m.P.due_ts < hi)]
    worst = win.groupby("loan_id").days_past_due.max()
    book["y"] = book.loan_id.map(worst).fillna(0) >= 30
    return book


def compute_facts(w: World, m: Metrics, delinq: pd.DataFrame, prov: pd.DataFrame) -> dict:
    F: dict = {}
    now, q2, py = AS_OF, date(2026, 6, 30), date(2025, 9, 30)
    book = {e: m.book(e) for e in (now, q2, py)}
    dl = {e: m.delinquency(e) for e in (now, q2, py)}
    b = book[now]
    F["active_loans"] = f"{len(b):,}"
    F["balance_total"] = millions(b.balance.sum())
    for prod in PRODUCTS:
        F[f"share_{prod}"] = pct((b["product"] == prod).mean(), 0)
    orig = w.loans[(w.loans.origination_date >= date(2026, 7, 1)) & (w.loans.origination_date <= now)]
    F["originations_q3"] = len(orig)
    for prod in PRODUCTS:
        F[f"originations_q3_{prod}"] = int((orig["product"] == prod).sum())
    for tag, e in (("now", now), ("q2", q2), ("py", py)):
        bb = book[e]
        F[f"rate30_all_{tag}"] = pct(bb.delinquent_30.mean())
        F[f"rate90_all_{tag}"] = pct(bb.delinquent_90.mean())
        for prod in PRODUCTS:
            F[f"rate30_{prod}_{tag}"] = pct(dl[e].rate_30[prod])
            F[f"rate90_{prod}_{tag}"] = pct(dl[e].rate_90[prod])
    top, low = dl[now].rate_30.idxmax(), dl[now].rate_30.idxmin()
    F["top30_product_lower"], F["top30_rate"] = PRODUCT_LABEL[top].lower(), pct(dl[now].rate_30[top])
    F["low30_product_lower"], F["low30_rate"] = PRODUCT_LABEL[low].lower(), pct(dl[now].rate_30[low])
    d_all = book[now].delinquent_30.mean() - book[py].delinquent_30.mean()
    F["delinq_direction"] = "up from" if d_all > 0.0005 else "down from" if d_all < -0.0005 else "unchanged from"
    F["delinq_loans_now"] = int(b.delinquent_30.sum())
    F["loans90_now"] = int(b.delinquent_90.sum())

    tables: dict[str, tuple[list[str], list[list[str]]]] = dict(band_tables())
    rows = []
    for prod in PRODUCTS:
        sel = b[b["product"] == prod]
        rows.append(
            [PRODUCT_LABEL[prod], f"{len(sel):,}", money(sel.balance.sum()), f"{sel.interest_rate.mean():.2f}%"]
        )
    rows.append(["All consumer loans", f"{len(b):,}", money(b.balance.sum()), f"{b.interest_rate.mean():.2f}%"])
    tables["portfolio"] = (["Product", "Active loans", "Outstanding principal", "Average interest rate"], rows)

    rows = []
    for prod in PRODUCTS:
        r = [PRODUCT_LABEL[prod]] + [pct(dl[e].rate_30[prod]) for e in (py, q2, now)]
        rows.append(r + [pts(dl[now].rate_30[prod] - dl[q2].rate_30[prod])])
    rows.append(
        ["All consumer loans"]
        + [pct(book[e].delinquent_30.mean()) for e in (py, q2, now)]
        + [pts(book[now].delinquent_30.mean() - book[q2].delinquent_30.mean())]
    )
    tables["delinquency30"] = (["Product", "30 Sep 2025", "30 Jun 2026", "30 Sep 2026", "Change since 30 Jun"], rows)
    F["bars_delinquency30"] = [(PRODUCT_LABEL[p], float(dl[now].rate_30[p])) for p in PRODUCTS] + [
        ("All consumer loans", float(b.delinquent_30.mean()))
    ]

    rows = []
    for prod in PRODUCTS:
        sel = b[b["product"] == prod]
        rows.append(
            [PRODUCT_LABEL[prod], f"{len(sel):,}", f"{int(sel.delinquent_90.sum())}", pct(dl[now].rate_90[prod])]
        )
    rows.append(["All consumer loans", f"{len(b):,}", f"{int(b.delinquent_90.sum())}", pct(b.delinquent_90.mean())])
    tables["delinquency90"] = (["Product", "Active loans", "Loans 90+ days past due", "90+ rate"], rows)

    q3 = prov[prov.quarter_end == now]
    q2p = prov[prov.quarter_end == q2]
    rows = []
    for r in q3.itertuples():
        rows.append(
            [
                r.segment,
                money(r.outstanding_balance),
                money(r.allowance),
                pct(r.coverage_ratio, 2),
                money(r.net_charge_offs),
                money_signed(r.provision_expense),
            ]
        )
    rows.append(
        [
            "All segments",
            money(q3.outstanding_balance.sum()),
            money(q3.allowance.sum()),
            pct(q3.allowance.sum() / q3.outstanding_balance.sum(), 2),
            money(q3.net_charge_offs.sum()),
            money_signed(q3.provision_expense.sum()),
        ]
    )
    tables["provisions"] = (
        ["Segment", "Outstanding", "Allowance", "Coverage", "Net charge-offs, Q3", "Provision expense, Q3"],
        rows,
    )
    co = m.charge_offs(q2, now)
    F["co_q3_count"] = len(co)
    F["co_q3_net"] = money(q3.net_charge_offs.sum())
    F["allowance_now"] = money(q3.allowance.sum())
    F["coverage_now"] = pct(q3.allowance.sum() / q3.outstanding_balance.sum(), 2)
    F["provision_q3"] = money_text(q3.provision_expense.sum())
    F["provision_q2"] = money_text(q2p.provision_expense.sum())

    mb = w.monthly_balances
    acct = w.accounts.set_index("account_id")
    mb = mb.assign(account_type=mb.account_id.map(acct.account_type))

    def month_slice(e: date) -> pd.DataFrame:
        return mb[mb.month_end == e]

    cards_now, cards_py = month_slice(now), month_slice(py)
    cards_now = cards_now[cards_now.account_type == "credit_card"]
    cards_py = cards_py[cards_py.account_type == "credit_card"]
    F["util_now"] = pct(cards_now.utilization_ratio.mean())
    F["util_py"] = pct(cards_py.utilization_ratio.mean())
    chk_now = month_slice(now)
    chk_now = chk_now[chk_now.account_type == "checking"]
    chk_py = month_slice(py)
    chk_py = chk_py[chk_py.account_type == "checking"]
    F["od_share_now"] = pct((chk_now.overdraft_count >= 1).mean())
    F["od_share_py"] = pct((chk_py.overdraft_count >= 1).mean())
    hot = cards_now[cards_now.utilization_ratio > 0.7]
    F["stressed_cards"] = len(hot)
    hot_cust = set(acct.customer_id.reindex(hot.account_id))
    borrowers = set(b.cust_idx.map(lambda i: w.cust_full.customer_id.iat[i]))
    F["stressed_cards_borrowers"] = len(hot_cust & borrowers)

    cohort = book[q2][book[q2].delinquent_30]
    still = 0
    for r in cohort.itertuples():
        row = m.L.loc[r.loan_id]
        if row.status == "charged_off" and row.closed_date <= now and row.closed_date > q2:
            still += 1
        elif r.loan_id in set(b.loan_id[b.delinquent_30]):
            still += 1
    F["roll_rate"] = pct(still / max(len(cohort), 1), 0)

    reg = b.assign(region=w.cust_full.region.values[b.cust_idx.values]).groupby("region")
    rr = pd.DataFrame(dict(n=reg.size(), bad=reg.delinquent_30.sum())).assign(rate=lambda d: d.bad / d.n)
    rr = rr.sort_values("rate", ascending=False)
    tables["regions"] = (
        ["Region", "Active loans", "Loans 30+ days past due", "Delinquency rate"],
        [[r, f"{int(x.n)}", f"{int(x.bad)}", pct(x.rate)] for r, x in rr.iterrows()],
    )
    F["worst_region"], F["best_region"] = rr.index[0], rr.index[-1]
    F["worst_region_rate"], F["best_region_rate"] = pct(rr.rate.iloc[0]), pct(rr.rate.iloc[-1])
    bad = b[b.delinquent_30]
    F["share_low_band"] = pct(bad.credit_band.isin(LOW_BANDS).mean(), 0)
    F["share_low_band_book"] = pct(b.credit_band.isin(LOW_BANDS).mean(), 0)
    F["_tables"] = tables
    F["_book"] = book
    F["_dl"] = dl
    return F


# --------------------------------------------------------------------------------------------------------------
# Renderers: PDF (reportlab platypus), DOCX (python-docx), PPTX (python-pptx)

NAVY = "#12355B"
TEAL = "#1F7A8C"
ZEBRA = "#F1F5F9"
GRID = "#B8C4D2"
SYNTHETIC_NOTICE = f"Synthetic document for a software demonstration. {BANK} is fictional."


def resolve_table(ref: str, tables: dict) -> tuple[list[str], list[list[str]]]:
    if ref not in tables:
        raise KeyError(f"table {ref!r} is not defined")
    return tables[ref]


def col_weights(header: list[str], rows: list[list[str]]) -> list[float]:
    w = []
    for j, h in enumerate(header):
        longest = max([len(h) * 0.7] + [len(r[j]) for r in rows])
        w.append(max(8.0, min(46.0, longest)) ** 0.85)
    total = sum(w)
    return [x / total for x in w]


def render_pdf(md_name: str, out: Path, ctx: dict | None, tables: dict) -> None:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import BaseDocTemplate
    from reportlab.platypus import Frame
    from reportlab.platypus import KeepTogether
    from reportlab.platypus import ListFlowable
    from reportlab.platypus import ListItem
    from reportlab.platypus import PageTemplate
    from reportlab.platypus import Paragraph
    from reportlab.platypus import Spacer
    from reportlab.platypus import Table
    from reportlab.platypus import TableStyle

    meta, blocks = load_md(md_name, ctx)
    navy, teal = colors.HexColor(NAVY), colors.HexColor(TEAL)
    body = ParagraphStyle("body", fontName="Helvetica", fontSize=9.6, leading=13.4, alignment=TA_LEFT, spaceAfter=6)
    h2 = ParagraphStyle(
        "h2",
        parent=body,
        fontName="Helvetica-Bold",
        fontSize=13,
        leading=16,
        textColor=navy,
        spaceBefore=12,
        spaceAfter=5,
        keepWithNext=1,
    )
    h3 = ParagraphStyle(
        "h3",
        parent=body,
        fontName="Helvetica-Bold",
        fontSize=10.8,
        leading=14,
        textColor=teal,
        spaceBefore=8,
        spaceAfter=3,
        keepWithNext=1,
    )
    cap = ParagraphStyle(
        "cap",
        parent=body,
        fontName="Helvetica-Bold",
        fontSize=9,
        leading=12,
        textColor=navy,
        spaceBefore=6,
        spaceAfter=3,
        keepWithNext=1,
    )
    cell = ParagraphStyle("cell", parent=body, fontSize=8.4, leading=10.6, spaceAfter=0)
    head = ParagraphStyle("head", parent=cell, fontName="Helvetica-Bold", textColor=colors.white)
    title = ParagraphStyle(
        "title", parent=body, fontName="Helvetica-Bold", fontSize=22, leading=26, textColor=navy, spaceAfter=2
    )
    sub = ParagraphStyle(
        "sub",
        parent=body,
        fontName="Helvetica-Oblique",
        fontSize=11,
        leading=14,
        textColor=colors.HexColor("#4A5A6B"),
        spaceAfter=10,
    )
    width = letter[0] - 2 * 0.9 * inch

    def make_table(header: list[str], rows: list[list[str]]):
        weights = col_weights(header, rows)
        data = [[Paragraph(markup(h), head) for h in header]] + [[Paragraph(markup(c), cell) for c in r] for r in rows]
        t = Table(data, colWidths=[x * width for x in weights], repeatRows=1)
        style = [
            ("BACKGROUND", (0, 0), (-1, 0), navy),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor(GRID)),
            ("LEFTPADDING", (0, 0), (-1, -1), 5),
            ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]
        for i in range(2, len(data), 2):
            style.append(("BACKGROUND", (0, i), (-1, i), colors.HexColor(ZEBRA)))
        t.setStyle(TableStyle(style))
        return t

    story: list = [Paragraph(markup(meta["title"]), title), Paragraph(markup(meta["subtitle"]), sub)]
    control = [
        ["Document", meta["doc_id"], "Version", meta["version"]],
        ["Effective", meta["effective"], "Owner", meta["owner"]],
        ["Approved", meta["approved"], "Classification", "Internal use"],
    ]
    ct = Table(
        [[Paragraph(markup(c), head if j % 2 == 0 else cell) for j, c in enumerate(r)] for r in control],
        colWidths=[0.95 * inch, 2.2 * inch, 1.0 * inch, 2.15 * inch],
    )
    ct.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (0, -1), teal),
                ("BACKGROUND", (2, 0), (2, -1), teal),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor(GRID)),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]
        )
    )
    story += [ct, Spacer(1, 8)]
    pending_caption = None
    for kind, val in blocks:
        if kind == "h2":
            story.append(Paragraph(markup(val), h2))
        elif kind == "h3":
            story.append(Paragraph(markup(val), h3))
        elif kind == "p":
            story.append(Paragraph(markup(val), body))
        elif kind == "caption":
            pending_caption = Paragraph(markup(val[2:-2]), cap)
        elif kind in ("table", "rows"):
            header, rows = (val[0], val[1:]) if kind == "rows" else resolve_table(val, tables)
            tbl = make_table(header, rows)
            if pending_caption is None:
                story.append(tbl)
            elif len(rows) < 12:
                story.append(KeepTogether([pending_caption, tbl]))
            else:
                story += [pending_caption, tbl]
            story.append(Spacer(1, 8))
            pending_caption = None
        elif kind in ("ul", "ol"):
            items = [ListItem(Paragraph(markup(x), body), leftIndent=14) for x in val]
            story.append(
                ListFlowable(
                    items,
                    bulletType="bullet" if kind == "ul" else "1",
                    start="•" if kind == "ul" else None,
                    leftIndent=14,
                    bulletFontSize=9,
                )
            )
        elif kind == "note":
            continue
        else:
            raise ValueError(f"unsupported block {kind}")

    def decorate(canvas, doc):
        canvas.saveState()
        canvas.setStrokeColor(navy)
        canvas.setLineWidth(1.2)
        canvas.line(0.9 * inch, letter[1] - 0.62 * inch, letter[0] - 0.9 * inch, letter[1] - 0.62 * inch)
        canvas.setFont("Helvetica-Bold", 8.5)
        canvas.setFillColor(navy)
        canvas.drawString(0.9 * inch, letter[1] - 0.52 * inch, BANK.upper())
        canvas.setFont("Helvetica", 8.5)
        canvas.drawRightString(
            letter[0] - 0.9 * inch, letter[1] - 0.52 * inch, f"{meta['doc_id']}  |  Version {meta['version']}"
        )
        canvas.setFillColor(colors.HexColor("#5B6B7C"))
        canvas.setFont("Helvetica", 7.6)
        canvas.drawString(0.9 * inch, 0.5 * inch, SYNTHETIC_NOTICE)
        canvas.drawRightString(letter[0] - 0.9 * inch, 0.5 * inch, f"Page {doc.page}")
        canvas.restoreState()

    doc = BaseDocTemplate(
        str(out),
        pagesize=letter,
        invariant=1,
        title=meta["title"],
        author=f"{BANK} (synthetic)",
        subject=meta["subtitle"],
        creator="financial-services pack generator",
        leftMargin=0.9 * inch,
        rightMargin=0.9 * inch,
        topMargin=0.9 * inch,
        bottomMargin=0.85 * inch,
    )
    frame = Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height, id="body")
    doc.addPageTemplates([PageTemplate(id="p", frames=[frame], onPage=decorate)])
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.build(story)


def render_docx(md_name: str, out: Path, ctx: dict | None, tables: dict) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt
    from docx.shared import RGBColor

    meta, blocks = load_md(md_name, ctx)
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    for name, size in (("Heading 1", 15), ("Heading 2", 12.5)):
        st = doc.styles[name]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(NAVY[1:])

    def add_runs(par, text: str, size: float | None = None) -> None:
        for t, bold, italic in runs(text):
            r = par.add_run(t)
            r.bold, r.italic = bold, italic
            if size:
                r.font.size = Pt(size)

    def shade(cell, color: str) -> None:
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), color)
        tcPr.append(shd)

    def add_table(header: list[str], rows: list[list[str]]) -> None:
        t = doc.add_table(rows=1, cols=len(header))
        t.style = "Table Grid"
        for j, h in enumerate(header):
            c = t.rows[0].cells[j]
            c.text = ""
            r = c.paragraphs[0].add_run(h)
            r.bold = True
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(255, 255, 255)
            shade(c, NAVY[1:])
        for i, row in enumerate(rows):
            cells = t.add_row().cells
            for j, v in enumerate(row):
                cells[j].text = ""
                add_runs(cells[j].paragraphs[0], v, 9.5)
                if i % 2 == 1:
                    shade(cells[j], ZEBRA[1:])
        doc.add_paragraph()

    doc.add_paragraph(meta["title"], style="Title")
    p = doc.add_paragraph()
    r = p.add_run(meta["subtitle"])
    r.italic = True
    add_table(
        ["Document", "Version", "Effective", "Owner", "Approved"],
        [[meta["doc_id"], meta["version"], meta["effective"], meta["owner"], meta["approved"]]],
    )
    for kind, val in blocks:
        if kind == "h2":
            doc.add_heading(val, level=1)
        elif kind == "h3":
            doc.add_heading(val, level=2)
        elif kind == "p":
            add_runs(doc.add_paragraph(), val)
        elif kind == "caption":
            par = doc.add_paragraph()
            r = par.add_run(val[2:-2])
            r.bold = True
        elif kind in ("table", "rows"):
            header, rows = (val[0], val[1:]) if kind == "rows" else resolve_table(val, tables)
            add_table(header, rows)
        elif kind in ("ul", "ol"):
            for item in val:
                add_runs(doc.add_paragraph(style="List Bullet" if kind == "ul" else "List Number"), item)
        elif kind == "note":
            continue
        else:
            raise ValueError(f"unsupported block {kind}")
    sec = doc.sections[0]
    hp = sec.header.paragraphs[0]
    hp.text = f"{BANK.upper()}  |  {meta['doc_id']}  |  Version {meta['version']}"
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    sec.footer.paragraphs[0].text = SYNTHETIC_NOTICE
    cp = doc.core_properties
    cp.author = f"{BANK} (synthetic)"
    cp.last_modified_by = f"{BANK} (synthetic)"
    cp.title = meta["title"]
    cp.subject = meta["subtitle"]
    cp.created = FIXED_TIME
    cp.modified = FIXED_TIME
    cp.revision = 1
    cp.comments = "Synthetic document generated for a software demonstration."
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    normalize_zip(out)


def render_pptx(md_name: str, out: Path, ctx: dict, tables: dict, bars: dict) -> None:
    from lxml import etree
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN
    from pptx.oxml.ns import qn
    from pptx.util import Emu
    from pptx.util import Inches
    from pptx.util import Pt

    meta, blocks = load_md(md_name, ctx)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    navy, teal = RGBColor.from_string(NAVY[1:]), RGBColor.from_string(TEAL[1:])
    grey = RGBColor(0x5B, 0x6B, 0x7C)

    def textbox(slide, x, y, w, h, text, size, bold=False, color=None, align=None):
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = text
        r.font.size, r.font.bold = Pt(size), bold
        if color is not None:
            r.font.color.rgb = color
        if align is not None:
            p.alignment = align
        return tb

    def footer(slide, n: int) -> None:
        textbox(
            slide,
            0.6,
            7.0,
            10.5,
            0.3,
            f"{BANK}  |  Credit Risk Committee, Q3 2026  |  Synthetic data for a software demonstration",
            10,
            color=grey,
        )
        textbox(slide, 12.0, 7.0, 0.8, 0.3, str(n), 10, color=grey, align=PP_ALIGN.RIGHT)

    # title slide
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = meta["title"]
    s.placeholders[1].text = meta["subtitle"]
    band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(2.2), prs.slide_width, Inches(0.12))
    band.fill.solid()
    band.fill.fore_color.rgb = teal
    band.line.fill.background()
    t = s.shapes.title
    t.left, t.top, t.width, t.height = Inches(0.8), Inches(2.5), Inches(11.7), Inches(1.4)
    t.text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT
    t.text_frame.paragraphs[0].runs[0].font.size = Pt(40)
    t.text_frame.paragraphs[0].runs[0].font.bold = True
    t.text_frame.paragraphs[0].runs[0].font.color.rgb = navy
    st = s.placeholders[1]
    st.left, st.top, st.width, st.height = Inches(0.8), Inches(4.0), Inches(11.7), Inches(1.0)
    st.text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT
    st.text_frame.paragraphs[0].runs[0].font.size = Pt(18)
    st.text_frame.paragraphs[0].runs[0].font.color.rgb = grey
    textbox(s, 0.8, 5.4, 11.7, 0.4, f"Prepared by {meta['author']}", 14, color=grey)
    s.notes_slide.notes_text_frame.text = "Synthetic data for a software demonstration."

    slides: list[tuple[str, list[tuple]]] = []
    for kind, val in blocks:
        if kind == "h2":
            slides.append((val, []))
        else:
            slides[-1][1].append((kind, val))

    def add_bullets(slide, x, y, w, h, items: list[str], size: float = 16) -> None:
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        for i, item in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", str(Emu(Inches(0.3))))
            pPr.set("indent", str(-Emu(Inches(0.3))))
            bu = etree.SubElement(pPr, qn("a:buChar"))
            bu.set("char", "•")
            p.space_after = Pt(7)
            for text, bold, italic in runs(item):
                r = p.add_run()
                r.text = text
                r.font.size, r.font.bold, r.font.italic = Pt(size), bold, italic

    def add_table(slide, header, rows, x, y, w) -> float:
        n = len(rows) + 1
        gf = slide.shapes.add_table(n, len(header), Inches(x), Inches(y), Inches(w), Inches(0.42 * n))
        tbl = gf.table
        weights = col_weights(header, rows)
        for j, wt in enumerate(weights):
            tbl.columns[j].width = Inches(w * wt)
        for i in range(n):
            tbl.rows[i].height = Inches(0.42)
            for j in range(len(header)):
                c = tbl.cell(i, j)
                c.margin_left = c.margin_right = Inches(0.08)
                c.fill.solid()
                text = header[j] if i == 0 else rows[i - 1][j]
                c.fill.fore_color.rgb = (
                    navy if i == 0 else (RGBColor(0xF1, 0xF5, 0xF9) if i % 2 == 0 else RGBColor(255, 255, 255))
                )
                tf = c.text_frame
                tf.word_wrap = True
                p = tf.paragraphs[0]
                r = p.add_run()
                r.text = text
                r.font.size = Pt(13)
                r.font.bold = i == 0 or rows[i - 1][0].startswith("All ")
                r.font.color.rgb = RGBColor(255, 255, 255) if i == 0 else RGBColor(0x1B, 0x26, 0x33)
        return y + 0.42 * n

    for n, (title, content) in enumerate(slides, start=2):
        s = prs.slides.add_slide(prs.slide_layouts[5])
        s.shapes.title.text = title
        t = s.shapes.title
        t.left, t.top, t.width, t.height = Inches(0.6), Inches(0.35), Inches(12.1), Inches(0.85)
        para = t.text_frame.paragraphs[0]
        para.alignment = PP_ALIGN.LEFT
        para.runs[0].font.size = Pt(28)
        para.runs[0].font.bold = True
        para.runs[0].font.color.rgb = navy
        rule = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(1.22), Inches(12.1), Inches(0.05))
        rule.fill.solid()
        rule.fill.fore_color.rgb = teal
        rule.line.fill.background()
        y = 1.5
        has_bars = any(k == "bars" for k, _ in content)
        notes: list[str] = []
        for kind, val in content:
            if kind == "table":
                header, rows = resolve_table(val, tables)
                width_in = 7.9 if has_bars else min(12.1, 2.5 * len(header))
                y = add_table(s, header, rows, 0.6, y, width_in) + 0.3
            elif kind == "bars":
                data = bars[val]
                top = 1.5
                textbox(s, 8.8, top - 0.05, 4.0, 0.3, "30+ days past due, 30 Sep 2026", 12, True, navy)
                peak = max(v for _, v in data) or 1
                for i, (label, v) in enumerate(data):
                    yy = top + 0.4 + i * 0.62
                    textbox(s, 8.8, yy, 4.0, 0.28, label, 12, color=grey)
                    bw = max(0.05, 3.1 * v / peak)
                    bar = s.shapes.add_shape(
                        MSO_SHAPE.RECTANGLE, Inches(8.85), Inches(yy + 0.3), Inches(bw), Inches(0.2)
                    )
                    bar.fill.solid()
                    bar.fill.fore_color.rgb = teal if not label.startswith("All") else navy
                    bar.line.fill.background()
                    textbox(s, 8.85 + bw + 0.05, yy + 0.22, 0.9, 0.3, pct(v), 12, True, navy)
            elif kind == "ul":
                add_bullets(s, 0.6, y, 7.9 if has_bars else 12.1, 7.0 - y, val, 16 if len(val) <= 4 else 15)
            elif kind == "note":
                notes.append(val)
            elif kind == "p":
                textbox(s, 0.6, y, 12.1, 0.8, val, 16)
        if notes:
            s.notes_slide.notes_text_frame.text = " ".join(notes)
        footer(s, n)

    cp = prs.core_properties
    cp.author = f"{BANK} (synthetic)"
    cp.last_modified_by = f"{BANK} (synthetic)"
    cp.title = meta["title"]
    cp.subject = meta["subtitle"]
    cp.created = FIXED_TIME
    cp.modified = FIXED_TIME
    cp.revision = 1
    cp.comments = "Synthetic document generated for a software demonstration."
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    normalize_zip(out)


# --------------------------------------------------------------------------------------------------------------
# The scanned loan modification request (PNG): a typeset form, filled in, signed, then degraded like a scan

FIRST_NAMES = ["Marisol", "Daniel", "Priya", "Terrence", "Lucia", "Anders", "Imani", "Wesley", "Noor", "Gideon"]
LAST_NAMES = [
    "Quintero",
    "Whitcombe",
    "Raghavan",
    "Oyelaran",
    "Brandt",
    "Kowalczyk",
    "Delacroix",
    "Fairbanks",
    "Haddad",
    "Lindqvist",
]
FORM_DATES = {"request": date(2026, 9, 14), "received": date(2026, 9, 16), "decision": date(2026, 9, 30)}
REQUEST_MONTHS = 6
REQUEST_START = date(2026, 10, 1)
REQUEST_REDUCTION = 0.40


def pick_borrower(w: World, m: Metrics) -> dict:
    """The first personal loan that is 30 to 89 days past due at the reporting date, with a made-up borrower name."""
    rng = rng_for("form")
    b = m.book(AS_OF)
    cand = b[(b["product"] == "personal") & (b.dpd >= 30) & (b.dpd < 90)].sort_values("loan_id")
    if cand.empty:
        raise RuntimeError("no personal loan is 30 to 89 days past due at the reporting date")
    r = cand.iloc[0]
    name = f"{FIRST_NAMES[int(rng.integers(0, len(FIRST_NAMES)))]} {'ABCDEFGHJKLMNPRSTW'[int(rng.integers(0, 18))]}. {LAST_NAMES[int(rng.integers(0, len(LAST_NAMES)))]}"
    new_payment = round(r.monthly_payment * (1 - REQUEST_REDUCTION))
    return dict(
        name=name,
        loan_id=r.loan_id,
        customer_id=w.cust_full.customer_id.iat[int(r.cust_idx)],
        principal=float(r.principal),
        rate=float(r.interest_rate),
        term=int(r.term_months),
        payment=float(r.monthly_payment),
        new_payment=int(new_payment),
        dpd=int(r.dpd),
        phone=f"(555) 010-{int(rng.integers(1000, 9999)):04d}",
    )


def render_form_pdf(f: dict) -> bytes:
    import textwrap

    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    meta, blocks = load_md("loan_modification_request.md")
    paras = {blocks[i][1]: blocks[i + 1][1] for i, bl in enumerate(blocks[:-1]) if bl[0] == "h2"}
    narrative, ack = paras["Hardship statement"], paras["Acknowledgment"]
    W, H = letter
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=letter, invariant=1)
    c.setTitle("Loan modification request (scanned)")
    ink = (0.05, 0.08, 0.30)
    dark = (0.12, 0.12, 0.14)

    def y(v: float) -> float:
        return H - v

    def bar(top: float, text: str) -> None:
        c.setFillColorRGB(0.22, 0.24, 0.28)
        c.rect(40, y(top + 15), 532, 15, stroke=0, fill=1)
        c.setFillColorRGB(1, 1, 1)
        c.setFont("Helvetica-Bold", 8.6)
        c.drawString(46, y(top + 10.5), text)

    def field(x: float, top: float, w: float, label: str, value: str, h: float = 30) -> None:
        c.setStrokeColorRGB(*dark)
        c.setLineWidth(0.6)
        c.rect(x, y(top + h), w, h, stroke=1, fill=0)
        c.setFillColorRGB(0.35, 0.35, 0.38)
        c.setFont("Helvetica", 6.6)
        c.drawString(x + 4, y(top + 8), label.upper())
        c.setFillColorRGB(*ink)
        c.setFont("Courier-Bold", 10.6)
        c.drawString(x + 5, y(top + 23), value)

    def checkbox(x: float, top: float, label: str, checked: bool, size: float = 9) -> float:
        c.setStrokeColorRGB(*dark)
        c.setLineWidth(0.8)
        c.rect(x, y(top + size), size, size, stroke=1, fill=0)
        if checked:
            c.setStrokeColorRGB(*ink)
            c.setLineWidth(1.5)
            c.line(x + 1.5, y(top + 1.8), x + size - 1.2, y(top + size - 1.2))
            c.line(x + 1.5, y(top + size - 1.2), x + size - 1.2, y(top + 1.8))
        c.setFillColorRGB(*dark)
        c.setFont("Helvetica", 8.4)
        c.drawString(x + size + 4, y(top + 7.4), label)
        return x + size + 10 + c.stringWidth(label, "Helvetica", 8.4)

    # header
    c.setFillColorRGB(*dark)
    c.setFont("Helvetica-Bold", 15)
    c.drawString(40, y(48), BANK.upper())
    c.setFont("Helvetica", 8.5)
    c.drawString(40, y(62), "Loan Servicing Department  |  P.O. Box 4417, Harborview")
    c.setFont("Helvetica-Bold", 11)
    c.drawRightString(572, y(46), f"Form {meta['form_id']}")
    c.setFont("Helvetica", 8.5)
    c.drawRightString(572, y(60), meta["form_revision"])
    c.setLineWidth(2)
    c.line(40, y(72), 572, y(72))
    c.setFont("Helvetica-Bold", 16)
    c.drawCentredString(W / 2, y(98), "REQUEST FOR LOAN MODIFICATION")
    c.setFont("Helvetica", 9)
    c.drawCentredString(
        W / 2, y(112), "Financial hardship  |  Complete every section, sign and attach documentation of the hardship"
    )

    bar(124, "1.  BORROWER AND LOAN INFORMATION")
    field(40, 142, 252, "Borrower name", f["name"])
    field(292, 142, 120, "Customer ID", f["customer_id"])
    field(412, 142, 160, "Loan number", f["loan_id"])
    field(40, 172, 252, "Loan type", "Personal loan (unsecured)")
    field(292, 172, 120, "Original principal", f"${f['principal']:,.2f}")
    field(412, 172, 160, "Current monthly payment", f"${f['payment']:,.2f}")
    field(40, 202, 252, "Telephone", f["phone"])
    field(292, 202, 120, "Date of request", FORM_DATES["request"].strftime("%m/%d/%Y"))
    field(412, 202, 160, "Best time to call", "Evenings after 6 PM")

    bar(242, "2.  REASON FOR HARDSHIP  (check all that apply)")
    x = 44.0
    for label, checked in (
        ("Job loss", False),
        ("Reduced income or work hours", True),
        ("Medical expense", False),
        ("Death of co-borrower", False),
        ("Disaster", False),
        ("Other", False),
    ):
        x = checkbox(x, 263, label, checked)
    c.setStrokeColorRGB(*dark)
    c.setLineWidth(0.6)
    c.rect(40, y(364), 532, 84, stroke=1, fill=0)
    c.setFillColorRGB(0.35, 0.35, 0.38)
    c.setFont("Helvetica", 6.6)
    c.drawString(44, y(288), "DESCRIBE YOUR SITUATION")
    c.setFillColorRGB(*ink)
    c.setFont("Courier", 8.7)
    for i, line in enumerate(textwrap.wrap(narrative, 94)[:7]):
        c.drawString(46, y(300 + i * 11.2), line)

    bar(376, "3.  MODIFICATION REQUESTED")
    new_payment = f"${f['new_payment']:,}"
    checkbox(44, 400, "Temporary payment reduction:  new monthly payment of", True)
    c.setFillColorRGB(*dark)
    c.setLineWidth(0.6)
    c.line(300, y(408), 352, y(408))
    c.setFont("Courier-Bold", 10.6)
    c.setFillColorRGB(*ink)
    c.drawString(303, y(406), new_payment)
    c.setFillColorRGB(*dark)
    c.setFont("Helvetica", 8.4)
    c.drawString(358, y(406.6), "for")
    c.line(373, y(408), 395, y(408))
    c.setFont("Courier-Bold", 10.6)
    c.setFillColorRGB(*ink)
    c.drawString(380, y(406), str(REQUEST_MONTHS))
    c.setFillColorRGB(*dark)
    c.setFont("Helvetica", 8.4)
    c.drawString(399, y(406.6), "months, beginning")
    c.line(471, y(408), 548, y(408))
    c.setFont("Courier-Bold", 10.6)
    c.setFillColorRGB(*ink)
    c.drawString(474, y(406), REQUEST_START.strftime("%m/%d/%Y"))
    checkbox(44, 420, "Term extension of ______ months (no more than 12)", False)
    checkbox(44, 440, "Waiver of late fees assessed in the last 60 days", True)

    bar(462, "4.  BORROWER CERTIFICATION")
    c.setFillColorRGB(*dark)
    c.setFont("Helvetica", 7.9)
    for i, line in enumerate(textwrap.wrap(ack, 118)):
        c.drawString(44, y(492 + i * 10.4), line)
    # signature
    sx, sy = 50.0, y(572)
    t = np.linspace(0, 1, 500)
    xs = sx + 190 * t + 7 * np.cos(24 * np.pi * t) * (0.35 + 0.65 * np.sin(np.pi * t))
    ys = sy + 13 * np.sin(24 * np.pi * t) * (0.3 + 0.7 * np.sin(np.pi * t)) + 7 * np.sin(3.2 * np.pi * t) + 6
    c.setStrokeColorRGB(*ink)
    c.setLineWidth(1.3)
    path = c.beginPath()
    path.moveTo(float(xs[0]), float(ys[0]))
    for px, py in zip(xs[1:], ys[1:], strict=True):
        path.lineTo(float(px), float(py))
    c.drawPath(path, stroke=1, fill=0)
    c.setLineWidth(1.0)
    c.line(sx + 70, sy - 3, sx + 245, sy + 7)  # the flourish under the name
    c.setStrokeColorRGB(*dark)
    c.setLineWidth(0.7)
    c.line(44, y(584), 300, y(584))
    c.line(330, y(584), 460, y(584))
    c.setFillColorRGB(0.35, 0.35, 0.38)
    c.setFont("Helvetica", 6.8)
    c.drawString(44, y(592), "BORROWER SIGNATURE")
    c.drawString(330, y(592), "DATE SIGNED")
    c.setFillColorRGB(*ink)
    c.setFont("Courier-Bold", 10.6)
    c.drawString(336, y(580), FORM_DATES["request"].strftime("%m/%d/%Y"))
    c.drawString(44, y(618), f["name"])
    c.setStrokeColorRGB(*dark)
    c.line(44, y(622), 300, y(622))
    c.setFillColorRGB(0.35, 0.35, 0.38)
    c.setFont("Helvetica", 6.8)
    c.drawString(44, y(630), "PRINTED NAME")

    # bank use box
    c.setStrokeColorRGB(*dark)
    c.setLineWidth(1)
    c.rect(40, y(740), 532, 88, stroke=1, fill=0)
    c.setFillColorRGB(*dark)
    c.setFont("Helvetica-Bold", 8)
    c.drawString(46, y(662), "FOR BANK USE ONLY")
    c.setFont("Helvetica", 8.4)
    c.drawString(46, y(682), "Date received:")
    c.drawString(46, y(700), "Case number:")
    c.drawString(46, y(718), "Decision due (10 business days):")
    c.setFillColorRGB(*ink)
    c.setFont("Courier-Bold", 10)
    c.drawString(130, y(682), FORM_DATES["received"].strftime("%m/%d/%Y"))
    c.drawString(130, y(700), "HM-2026-0412")
    c.drawString(200, y(718), FORM_DATES["decision"].strftime("%m/%d/%Y"))
    # received stamp
    c.saveState()
    c.translate(430, y(700))
    c.rotate(7)
    c.setStrokeColorRGB(0.25, 0.1, 0.1)
    c.setFillColorRGB(0.25, 0.1, 0.1)
    c.setLineWidth(2)
    c.roundRect(-70, -22, 140, 44, 4, stroke=1, fill=0)
    c.setFont("Helvetica-Bold", 14)
    c.drawCentredString(0, 4, "RECEIVED")
    c.setFont("Helvetica-Bold", 8)
    c.drawCentredString(0, -8, "SEP 16 2026  LOAN SERVICING")
    c.restoreState()
    c.setFillColorRGB(0.4, 0.4, 0.42)
    c.setFont("Helvetica", 6.8)
    c.drawString(40, y(764), f"Form {meta['form_id']} {meta['form_revision']}  |  Specimen: {SYNTHETIC_NOTICE}")
    c.showPage()
    c.save()
    return buf.getvalue()


def render_scan(f: dict, out: Path) -> None:
    """Typeset the form, render it, then make it look scanned: grayscale, a slight tilt, shading and noise."""
    import pypdfium2 as pdfium
    from PIL import Image
    from PIL import ImageFilter

    pdf = pdfium.PdfDocument(render_form_pdf(f))
    page = pdf[0].render(scale=150 / 72).to_pil().convert("L")
    rng = rng_for("scan")
    arr = np.asarray(page, dtype=np.float64)
    h, w = arr.shape
    yy, xx = np.mgrid[0:h, 0:w]
    paper = (
        238 - 10 * (xx / w) - 6 * (yy / h) + 7 * np.exp(-(((xx - 0.08 * w) / (0.05 * w)) ** 2))
    )  # shading, a shadow near the left edge
    arr = arr / 255.0 * paper
    img = Image.fromarray(np.clip(arr, 0, 255).astype(np.uint8)).filter(ImageFilter.GaussianBlur(0.8))
    img = img.rotate(-1.3, resample=Image.BICUBIC, fillcolor=205)
    arr = np.asarray(img, dtype=np.float64)
    arr += rng.normal(0, 4.5, arr.shape)
    specks = rng.random(arr.shape) < 0.00018
    arr[specks] -= rng.uniform(40, 110, int(specks.sum()))
    arr[:, int(0.93 * w) : int(0.93 * w) + 2] -= 7  # a faint scanner streak
    img = Image.fromarray(np.clip(arr, 0, 255).astype(np.uint8), mode="L")
    out.parent.mkdir(parents=True, exist_ok=True)
    img.save(out, format="PNG", optimize=True)
    Image.open(out).verify()


# --------------------------------------------------------------------------------------------------------------
# README: the data card and the answer key, computed from the tables


def md_table(header: list[str], rows: list[list]) -> str:
    out = ["| " + " | ".join(header) + " |", "|" + "|".join("---" for _ in header) + "|"]
    out += ["| " + " | ".join(str(c) for c in r) + " |" for r in rows]
    return "\n".join(out)


def signal_checks(w: World, m: Metrics) -> dict:
    """How far simple, visible signals move the prediction labels (the lift a model should find)."""
    S: dict = {}
    anchor = date(2026, 6, 30)
    book = loan_label(m, anchor)
    S["loan_anchor"] = anchor
    S["loan_n"] = len(book)
    S["loan_base"] = book.y.mean()
    S["loan_unpaid"] = book[book.dpd > 0].y.mean()
    S["loan_unpaid_n"] = int((book.dpd > 0).sum())
    S["loan_clean"] = book[book.dpd == 0].y.mean()
    S["loan_low"] = book[book.credit_band.isin(LOW_BANDS)].y.mean()
    S["loan_high"] = book[~book.credit_band.isin(LOW_BANDS)].y.mean()
    S["loan_positives"] = int(book.y.sum())
    mb0 = w.monthly_balances
    acct0 = w.accounts.set_index("account_id")
    cards = mb0[(mb0.month_end == anchor) & (mb0.account_id.map(acct0.account_type) == "credit_card")]
    util_by_cust = (
        cards.assign(customer_id=cards.account_id.map(acct0.customer_id)).groupby("customer_id").utilization_ratio.max()
    )
    book["util"] = [util_by_cust.get(w.cust_full.customer_id.iat[int(i)], np.nan) for i in book.cust_idx]
    S["loan_hi_util"] = book[book.util > 0.6].y.mean()
    S["loan_hi_util_n"] = int((book.util > 0.6).sum())
    S["loan_lo_util"] = book[book.util <= 0.6].y.mean()
    now = m.book(AS_OF)
    S["now_30_89"] = int(((now.dpd >= 30) & (now.dpd < 90)).sum())
    S["now_1_29"] = int(((now.dpd >= 1) & (now.dpd < 30)).sum())
    S["now_90"] = int((now.dpd >= 90).sum())

    mb = w.monthly_balances
    acct = w.accounts.set_index("account_id")
    cust = w.customers.set_index("customer_id")
    chk = mb[mb.account_id.map(acct.account_type) == "checking"]
    at = chk[chk.month_end == anchor].set_index("account_id")
    later = chk[chk.month_end.isin([date(2026, 7, 31), date(2026, 8, 31)])].account_id.unique()
    at["y"] = ~at.index.isin(later)
    earlier = chk[chk.month_end == date(2026, 3, 31)].set_index("account_id").ending_balance
    at["drop"] = (at.ending_balance < 0.5 * earlier.reindex(at.index)).fillna(False)
    at["online"] = at.index.map(lambda a: cust.online_banking_enrolled[acct.customer_id[a]])
    S["acct_n"] = len(at)
    S["acct_base"] = at.y.mean()
    S["acct_closed"] = int(at.y.sum())
    S["acct_drop"] = at[at["drop"]].y.mean()
    S["acct_drop_n"] = int(at["drop"].sum())
    S["acct_nodrop"] = at[~at["drop"]].y.mean()
    S["acct_offline"] = at[at.online == 0].y.mean()
    S["acct_online"] = at[at.online == 1].y.mean()
    n_prod = w.accounts.groupby("customer_id").size()
    cu = at.index.map(lambda a: acct.customer_id[a])
    at["single"] = [n_prod[c] == 1 for c in cu]
    at["short"] = [cust.tenure_years[c] < 2 for c in cu]
    S["acct_single"], S["acct_multi"] = at[at.single].y.mean(), at[~at.single].y.mean()
    S["acct_short"], S["acct_long"] = at[at.short].y.mean(), at[~at.short].y.mean()

    d = w.card_disputes
    a2 = date(2026, 8, 31)
    in_win = set(d[(d.dispute_date > a2) & (d.dispute_date <= AS_OF)].customer_id)
    prior = set(d[(d.dispute_date > date(2026, 2, 28)) & (d.dispute_date <= a2)].customer_id)
    card_holders = set(w.accounts[w.accounts.account_type == "credit_card"].customer_id)
    allc = list(w.customers.customer_id)
    S["disp_base"] = np.mean([c in in_win for c in allc])
    S["disp_prior"] = np.mean([c in in_win for c in allc if c in prior])
    S["disp_noprior"] = np.mean([c in in_win for c in allc if c not in prior])
    S["disp_card"] = np.mean([c in in_win for c in allc if c in card_holders])
    S["disp_nocard"] = np.mean([c in in_win for c in allc if c not in card_holders])
    return S


def answer_key(
    w: World, m: Metrics, F: dict, delinq: pd.DataFrame, prov: pd.DataFrame, borrower: dict, sig: dict
) -> dict:
    A: dict = {}
    good = next(b for b in BANDS if b["name"] == "Good")
    A["ans_underwriting"] = (
        f"A 700 score is in the **Good** band (670 to 719) of Table 3 in `consumer_credit_policy.pdf` (version 5.1). "
        f"The maximum DTI is **{pct(good['max_dti'], 0)}** and the maximum mortgage LTV is **{pct(good['ltv_mtg'], 0)}** "
        f"(maximum auto LTV {pct(good['ltv_auto'], 0)}). Version 5.0 allowed a DTI of {pct(good['prior_dti'], 0)}, so an answer of "
        f"{pct(good['prior_dti'], 0)} reads the superseded table in Table 5."
    )
    ln = w.loans[w.loans.loan_id == borrower["loan_id"]].iloc[0]
    now = F["_book"][AS_OF]
    dpd = int(now[now.loan_id == borrower["loan_id"]].dpd.iloc[0])
    A["ans_scan"] = (
        f"The scanned `loan_modification_request_scan.png` is a signed Form LS-114 from **{borrower['name']}** "
        f"(customer {borrower['customer_id']}) for **loan {borrower['loan_id']}**, a personal loan with an original principal of "
        f"${borrower['principal']:,.2f} and a current monthly payment of ${borrower['payment']:,.2f}. The borrower asks for a temporary payment "
        f"reduction to **${borrower['new_payment']:,} a month for {REQUEST_MONTHS} months beginning {REQUEST_START:%m/%d/%Y}** "
        f"({pct(REQUEST_REDUCTION, 0)} lower) and a waiver of late fees, because the employer cut work hours from 40 to 22 a week. Signed "
        f"{FORM_DATES['request']:%m/%d/%Y}, stamped received {FORM_DATES['received']:%m/%d/%Y}, case HM-2026-0412, decision due "
        f"{FORM_DATES['decision']:%m/%d/%Y}. The request is inside the policy limits (at most a 50 percent reduction for at most 6 months). "
        f"In the tables, {borrower['loan_id']} is a {ln['product']} loan of ${ln.principal:,.0f} at {ln.interest_rate}% over {ln.term_months} months "
        f"with a ${ln.monthly_payment:,.2f} payment, and at 30 September 2026 its oldest unpaid installment is {dpd} days past due."
    )
    A["ans_complaint"] = (
        "A complaint about a denied loan modification is a **Lending and loan servicing** complaint (Table 1 of "
        "`complaint_handling_procedure.docx`, owner Loan Servicing): acknowledge within **2 business days** and resolve within "
        "**15 business days**. Any complaint not resolved by the **10th business day** is escalated automatically to the Complaint Resolution "
        "Team, which may extend the resolution time by up to 15 business days if it tells the customer in writing before the deadline. A "
        "complaint open more than 30 business days is reported weekly to the Head of Customer Care."
    )
    mb = w.monthly_balances
    acct = w.accounts.set_index("account_id")
    dep = mb[(mb.month_end == AS_OF) & (mb.account_id.map(acct.account_type).isin(["checking", "savings"]))].copy()
    dep["branch_id"] = dep.account_id.map(acct.branch_id)
    top = dep.groupby("branch_id").ending_balance.sum().sort_values(ascending=False)
    names = w.branches.set_index("branch_id").branch_name
    rows = [[i + 1, b, names[b], money(v)] for i, (b, v) in enumerate(top.head(5).items())]
    A["ans_branches"] = (
        "Sum `ending_balance` of the checking and savings rows in `monthly_balances` for `month_end = 2026-09-30`, grouped by the account's "
        "`branch_id`:\n\n"
        + md_table(["Rank", "Branch", "Name", "Deposit balances"], rows)
        + f"\n\nAll 18 branches together hold {money(dep.ending_balance.sum())} across {len(dep):,} deposit accounts."
    )
    d = delinq
    sep26 = d[d.month_end == AS_OF].set_index("product")
    sep25 = d[d.month_end == date(2025, 9, 30)].set_index("product")
    rows = [
        [
            p,
            pct(sep26.delinquency_rate_30plus[p], 2),
            pct(sep25.delinquency_rate_30plus[p], 2),
            int(sep26.loans_30plus[p]),
            int(sep26.active_loans[p]),
        ]
        for p in PRODUCTS
    ]
    top_p = sep26.delinquency_rate_30plus.idxmax()
    A["ans_delinquency"] = (
        f"From `credit_risk_report_delinquency` (`month_end = 2026-09-30`): the highest 30-day delinquency rate is **{top_p}** at "
        f"**{pct(sep26.delinquency_rate_30plus[top_p], 2)}** (September 2025: {pct(sep25.delinquency_rate_30plus[top_p], 2)}).\n\n"
        + md_table(
            ["Product", "30+ rate, Sep 2026", "30+ rate, Sep 2025", "Loans 30+, Sep 2026", "Active loans, Sep 2026"],
            rows,
        )
    )
    dd = w.card_disputes[(w.card_disputes.dispute_date >= date(2026, 7, 1)) & (w.card_disputes.dispute_date <= AS_OF)]
    by = (
        dd.groupby("reason")
        .agg(n=("dispute_id", "count"), amount=("amount", "sum"))
        .sort_values(["n", "amount"], ascending=False)
    )
    rows = [[r, int(x.n), money(x.amount)] for r, x in by.iterrows()]
    A["ans_disputes"] = (
        f"Filter `card_disputes` to `dispute_date` from 2026-07-01 to 2026-09-30: **{len(dd)} disputes** with a total disputed amount of "
        f"**{money(dd.amount.sum())}** (${dd.amount.sum():,.2f}). The most common reason is **{by.index[0]}** ({int(by.n.iloc[0])} disputes).\n\n"
        + md_table(["Reason", "Disputes", "Amount"], rows)
    )
    cap = {b["name"]: b["max_dti"] for b in BANDS}
    p26 = w.loans[(w.loans["product"] == "personal") & (w.loans.origination_date >= date(2026, 1, 1))].copy()
    p26["cap"] = p26.credit_band.map(cap)
    viol = p26[p26.dti_at_origination > p26.cap].sort_values("loan_id")
    rows = [
        [
            r.loan_id,
            r.credit_band,
            pct(r.dti_at_origination, 0),
            pct(r.cap, 0),
            f"{r.origination_date}",
            money(r.principal),
        ]
        for r in viol.itertuples()
    ]
    A["ans_dti"] = (
        f"Maximum DTI by band comes from Table 3 of the credit policy (Exceptional 45%, Very Good 43%, Good 40%, Fair 36%, Poor 30%). "
        f"Of the {len(p26)} personal loans originated from 2026-01-01, **{len(viol)}** have `dti_at_origination` above the limit for their "
        f"`credit_band` (the band at origination; a DTI equal to the limit is allowed):\n\n"
        + md_table(["Loan", "Credit band", "DTI at origination", "Maximum DTI", "Originated", "Principal"], rows)
    )
    sheet = F["_dl"][AS_OF]
    rows = [
        [
            PRODUCT_LABEL[p],
            int(sheet.active_loans[p]),
            int(sheet.loans_30plus[p]),
            pct(sheet.rate_30[p]),
            F[f"rate30_{p}_now"],
        ]
        for p in PRODUCTS
    ]
    b = F["_book"][AS_OF]
    rows.append(
        ["All consumer loans", len(b), int(b.delinquent_30.sum()), pct(b.delinquent_30.mean()), F["rate30_all_now"]]
    )
    A["ans_deck"] = (
        "Yes. Slide 3 of `q3_2026_risk_committee.pptx` reports 30+ day delinquency of "
        f"**{F['rate30_personal_now']}** (personal), **{F['rate30_auto_now']}** (auto) and **{F['rate30_mortgage_now']}** (mortgage), "
        f"**{F['rate30_all_now']}** overall at 30 September 2026. The deck's definition is active loans (`status = 'active'`) with at least one "
        "installment in `loan_payments` that is unpaid (`paid_date IS NULL`) and has `days_past_due >= 30`, divided by active loans. The "
        "tables give the same figures:\n\n"
        + md_table(
            ["Product", "Active loans", "Loans 30+ days past due", "Rate from the tables", "Rate in the deck"], rows
        )
        + f"\n\nThe same rates are in `credit_risk_report_delinquency` for `month_end = 2026-09-30`. The deck also reports {F['loans90_now']} loans "
        f"90+ days past due ({F['rate90_all_now']})."
    )
    A["ans_loan_pred"] = (
        f"There is no single right list; Kumo returns a probability per loan for the `loan_default_90d` template. The base rate is "
        f"{pct(sig['loan_base'])} (at the earlier anchor {sig['loan_anchor']}, {sig['loan_positives']} of {sig['loan_n']} active loans had an "
        f"installment due in the next 90 days that ended 30+ days past due). Signals a good model finds: loans with an unpaid installment at the anchor "
        f"({sig['loan_unpaid_n']} loans) had a {pct(sig['loan_unpaid'])} rate against {pct(sig['loan_clean'])} for loans with none; Fair and "
        f"Poor credit bands had {pct(sig['loan_low'])} against {pct(sig['loan_high'])}; borrowers with a credit card above 60 percent utilization at the "
        f"anchor ({sig['loan_hi_util_n']} loans) had {pct(sig['loan_hi_util'])} against {pct(sig['loan_lo_util'])} for borrowers with a card below that. At 2026-09-30, {sig['now_30_89']} active loans are "
        f"30 to 89 days past due and {sig['now_90']} are 90+; they should rank near the top. Predictions are produced by Kumo and are not "
        "reproducible from the files alone."
    )
    A["ans_acct_pred"] = (
        f"Kumo returns a closure probability per checking account for `account_closure_90d` (an account is closed when no `monthly_balances` row "
        f"follows in the next 90 days). Base rate at the earlier anchor {sig['loan_anchor']}: {pct(sig['acct_base'])} "
        f"({sig['acct_closed']} of {sig['acct_n']} checking accounts open at that date had no balance row for July and August). Signals: accounts whose "
        f"balance fell by more than half in three months ({sig['acct_drop_n']} accounts) closed at {pct(sig['acct_drop'])} against "
        f"{pct(sig['acct_nodrop'])} for the rest; customers not enrolled in online banking closed at {pct(sig['acct_offline'])} against "
        f"{pct(sig['acct_online'])}; customers with a single product closed at {pct(sig['acct_single'])} against {pct(sig['acct_multi'])}, and "
        f"customers with under two years of tenure at {pct(sig['acct_short'])} against {pct(sig['acct_long'])}."
    )
    A["sig_disputes"] = (
        f"For `card_dispute_30d`, the base rate at the anchor 2026-08-31 is {pct(sig['disp_base'], 2)} of customers with a dispute in the next 30 days; "
        f"customers with a dispute in the prior six months had {pct(sig['disp_prior'], 2)} against {pct(sig['disp_noprior'], 2)}, and card holders "
        f"{pct(sig['disp_card'], 2)} against {pct(sig['disp_nocard'], 2)} for the rest."
    )
    return A


def render_readme(w: World, frames: dict, A: dict, F: dict, sig: dict) -> None:
    sizes = {p.name: p.stat().st_size for p in FILES.rglob("*") if p.is_file()}
    total = sum(sizes.values())
    ctx = dict(
        A,
        n_customers=f"{len(w.customers):,}",
        total_mb=f"{total / 1e6:.1f}",
        rows_table=md_table(
            ["Table", "Rows"],
            [[f"`{name}`", f"{len(df):,}"] for name, df in frames.items()]
            + [["`credit_risk_report_delinquency`", "54"], ["`credit_risk_report_provisions`", "30"]],
        ),
        total_rows=f"{sum(len(df) for df in frames.values()) + 84:,}",
        loans_active=F["active_loans"],
        borrowers=f"{w.loans.customer_id.nunique():,} ({pct(w.loans.customer_id.nunique() / len(w.customers), 0)})",
    )
    text = fill((CONTENT / "README.template.md").read_text(encoding="utf-8"), ctx)
    (PACK / "README.md").write_text(text, encoding="utf-8")


# --------------------------------------------------------------------------------------------------------------
# Main


def build_tables(w: World) -> dict[str, pd.DataFrame]:
    accounts = w.accounts.copy()
    accounts["credit_limit"] = accounts.credit_limit.astype("Int64")
    loans = w.loans.copy()
    loans["principal"] = loans.principal.astype(int)
    return {
        "branches": w.branches,
        "customers": w.customers,
        "accounts": accounts,
        "monthly_balances": w.monthly_balances,
        "loans": loans,
        "loan_payments": w.loan_payments,
        "card_disputes": w.card_disputes,
    }


def main() -> None:
    if FILES.exists():
        shutil.rmtree(FILES)
    w = World()
    m = Metrics(w)
    frames = build_tables(w)
    for name, df in frames.items():
        write_csv(df, name)
    delinq, prov = m.delinquency_sheet(), m.provisions_sheet()
    write_xlsx(FILES / "banking" / "credit_risk_report.xlsx", {"delinquency": delinq, "provisions": prov})
    F = compute_facts(w, m, delinq, prov)
    tables, bars = F["_tables"], {"delinquency30": F["bars_delinquency30"]}
    borrower = pick_borrower(w, m)
    render_pdf("consumer_credit_policy.md", FILES / "policies" / "consumer_credit_policy.pdf", F, tables)
    render_pdf("kyc_aml_procedure.md", FILES / "policies" / "kyc_aml_customer_due_diligence.pdf", F, tables)
    render_docx("complaint_handling_procedure.md", FILES / "policies" / "complaint_handling_procedure.docx", F, tables)
    render_scan(borrower, FILES / "policies" / "loan_modification_request_scan.png")
    render_pdf("overdraft_fee_schedule.md", FILES / "disclosures" / "overdraft_fee_schedule.pdf", F, tables)
    render_pptx("risk_committee_deck.md", FILES / "disclosures" / "q3_2026_risk_committee.pptx", F, tables, bars)
    sig = signal_checks(w, m)
    A = answer_key(w, m, F, delinq, prov, borrower, sig)
    render_readme(w, frames, A, F, sig)
    total = sum(p.stat().st_size for p in FILES.rglob("*") if p.is_file())
    print(f"wrote {sum(1 for p in FILES.rglob('*') if p.is_file())} files, {total / 1e6:.2f} MB under {FILES}")
    for name, df in frames.items():
        print(f"  {name}: {len(df):,} rows")


if __name__ == "__main__":
    main()
