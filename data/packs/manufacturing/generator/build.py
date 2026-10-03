#!/usr/bin/env -S uv run --script
# SPDX-FileCopyrightText: Copyright (c) 2026, NVIDIA CORPORATION & AFFILIATES. All rights reserved.
# SPDX-License-Identifier: Apache-2.0
#
# /// script
# requires-python = ">=3.12"
# dependencies = [
#   "numpy>=2.0",
#   "pandas>=2.2",
#   "openpyxl>=3.1",
#   "reportlab>=4.0",
#   "python-docx>=1.1",
#   "python-pptx>=1.0",
#   "pillow>=10.1",
#   "pypdfium2>=4.30",
#   "pyyaml>=6.0",
# ]
# ///
"""Generate the Atlas Precision Components pack: a fictional machining manufacturer.

    uv run data/packs/manufacturing/generator/build.py

Everything is written under ../files/ (and ../README.md) from one seeded simulation, so the tables, the reports
and the answer key always agree, and the same inputs always give the same bytes:

    files/tables/*.csv, production_plan.xlsx   plants, machines, sensor_readings, maintenance_events, work_orders,
                                               quality_defects and a two-sheet production plan
    files/procedures/                          preventive maintenance manual (PDF), lockout/tagout SOP (PDF), a
                                               scanned handwritten round sheet (PNG)
    files/reports/                             supplier quality agreement (DOCX), Q3 2026 plant review (PPTX), an 8D
                                               report (PDF)

The prose lives in content/*.md with {placeholders} that are filled from figures computed on the generated
tables. Everything here is synthetic: no real company, plant, machine, supplier or person.
"""

from __future__ import annotations

import argparse
import io
import math
import re
import sys
import zipfile
from dataclasses import dataclass
from dataclasses import field
from datetime import date
from datetime import datetime
from datetime import timedelta
from pathlib import Path

import numpy as np
import pandas as pd

HERE = Path(__file__).resolve().parent
PACK = HERE.parent
FILES = PACK / "files"
CONTENT = HERE / "content"

SEED = 20260930
START = date(2025, 7, 1)  # first day of history
AS_OF = date(2026, 9, 30)  # the data is current to here
SIM_END = date(2026, 10, 31)  # the simulation runs a month past as_of; that month is never written
PRE_DAYS = 400  # calendar days simulated before START, so each machine has a PM-2 history on day one

CAL0 = START - timedelta(days=PRE_DAYS)
S0 = PRE_DAYS  # day index of START (day indexes count from CAL0)
ASOF_I = S0 + (AS_OF - START).days
END_I = S0 + (SIM_END - START).days
NDAYS = END_I + 1

# Calendar -------------------------------------------------------------------------------------------------------

PLANTS = [
    {
        "plant_id": "DAY",
        "plant_name": "Dayton Plant",
        "city": "Dayton, OH",
        "country": "United States",
        "timezone": "America/New_York",
        "opened_year": 1998,
        "headcount": 312,
        "shifts_per_day": 2,
        "main_products": "Transmission housings, brackets and mounts, pump components",
        "climate": 0.0,
        "plan_factor": 0.93,
    },
    {
        "plant_id": "MTY",
        "plant_name": "Monterrey Plant",
        "city": "Monterrey, NL",
        "country": "Mexico",
        "timezone": "America/Monterrey",
        "opened_year": 2009,
        "headcount": 268,
        "shifts_per_day": 2,
        "main_products": "Valve bodies, pump components, brackets and mounts",
        "climate": 3.0,
        "plan_factor": 0.985,
    },
    {
        "plant_id": "BRN",
        "plant_name": "Brno Plant",
        "city": "Brno",
        "country": "Czechia",
        "timezone": "Europe/Prague",
        "opened_year": 2014,
        "headcount": 184,
        "shifts_per_day": 2,
        "main_products": "Valve bodies, transmission housings, pump components",
        "climate": -2.0,
        "plan_factor": 0.90,
    },
]
PLANT_IDS = [p["plant_id"] for p in PLANTS]
PLANT_NAME = {p["plant_id"]: p["plant_name"] for p in PLANTS}

SHUTDOWN = (date(2025, 12, 24), date(2026, 1, 2))  # every plant
HOLIDAYS = {
    "DAY": [
        date(2025, 7, 4),
        date(2025, 9, 1),
        date(2025, 11, 27),
        date(2025, 11, 28),
        date(2026, 5, 25),
        date(2026, 7, 3),
        date(2026, 9, 7),
    ],
    "MTY": [
        date(2025, 9, 16),
        date(2025, 11, 17),
        date(2026, 2, 2),
        date(2026, 3, 16),
        date(2026, 5, 1),
        date(2026, 9, 16),
    ],
    "BRN": [
        date(2025, 10, 28),
        date(2025, 11, 17),
        date(2026, 4, 6),
        date(2026, 5, 1),
        date(2026, 5, 8),
        date(2026, 7, 6),
        date(2026, 9, 28),
    ],
}


def day_date(i: int) -> date:
    return CAL0 + timedelta(days=int(i))


def day_index(d: date) -> int:
    return (d - CAL0).days


def build_open_masks() -> dict[str, np.ndarray]:
    masks = {}
    for pid in PLANT_IDS:
        closed = set(HOLIDAYS[pid])
        mask = np.zeros(NDAYS, dtype=bool)
        for i in range(NDAYS):
            d = day_date(i)
            mask[i] = d.weekday() < 5 and d not in closed and not (SHUTDOWN[0] <= d <= SHUTDOWN[1] and d >= START)
        masks[pid] = mask
    return masks


# Machines -------------------------------------------------------------------------------------------------------

ROSTER = {
    "DAY": {"CNC Mill": 8, "CNC Lathe": 5, "Hydraulic Press": 4, "Robot Cell": 3, "Surface Grinder": 1},
    "MTY": {"CNC Mill": 6, "CNC Lathe": 5, "Hydraulic Press": 3, "Robot Cell": 3, "Surface Grinder": 2},
    "BRN": {"CNC Mill": 6, "CNC Lathe": 4, "Hydraulic Press": 3, "Robot Cell": 2, "Surface Grinder": 1},
}

# pm2: PM-2 interval in calendar days (the maintenance manual's Table 1); v0: baseline vibration (mm/s RMS);
# t0: baseline max temperature (C); a0: baseline alarms a day; hz: daily failure hazard; rate: units per 16 h.
TYPES = {
    "CNC Mill": {
        "code": "MIL",
        "pm2": 90,
        "v0": 1.6,
        "t0": 47.0,
        "a0": 0.18,
        "hz": 0.0057,
        "rate": 70,
        "models": [("Tauber Machine Works", "VMC-850"), ("Tauber Machine Works", "VMC-1100")],
        "codes": ["SPN-BRG", "COOL-PMP", "ATC-JAM", "SRV-DRV"],
        "cw": [0.34, 0.22, 0.24, 0.20],
        "pm_hours": (6.0, 9.0),
    },
    "CNC Lathe": {
        "code": "LTH",
        "pm2": 120,
        "v0": 1.9,
        "t0": 51.0,
        "a0": 0.20,
        "hz": 0.0066,
        "rate": 130,
        "models": [("Vantor Systems", "TC-300"), ("Vantor Systems", "TC-450")],
        "codes": ["SPN-BRG", "CHK-HYD", "COOL-PMP", "SRV-DRV"],
        "cw": [0.35, 0.20, 0.20, 0.25],
        "pm_hours": (7.0, 10.0),
    },
    "Hydraulic Press": {
        "code": "PRS",
        "pm2": 60,
        "v0": 2.7,
        "t0": 57.0,
        "a0": 0.22,
        "hz": 0.0083,
        "rate": 420,
        "models": [("Helmsford Hydraulics", "HP-250T"), ("Helmsford Hydraulics", "HP-400T")],
        "codes": ["HYD-LEAK", "PMP-WEAR", "VLV-STK"],
        "cw": [0.40, 0.30, 0.30],
        "pm_hours": (4.0, 6.0),
    },
    "Robot Cell": {
        "code": "RBT",
        "pm2": 180,
        "v0": 1.2,
        "t0": 43.0,
        "a0": 0.15,
        "hz": 0.0046,
        "rate": 260,
        "models": [("Brenmark Robotics", "RW-6 weld cell"), ("Brenmark Robotics", "RH-12 handling cell")],
        "codes": ["RDC-WEAR", "CBL-FLT", "SRV-DRV"],
        "cw": [0.40, 0.35, 0.25],
        "pm_hours": (8.0, 12.0),
    },
    "Surface Grinder": {
        "code": "GRD",
        "pm2": 90,
        "v0": 1.4,
        "t0": 45.0,
        "a0": 0.14,
        "hz": 0.0055,
        "rate": 160,
        "models": [("Ostrander Grinding", "SG-1200")],
        "codes": ["WHL-IMB", "SPN-BRG", "COOL-PMP"],
        "cw": [0.40, 0.30, 0.30],
        "pm_hours": (5.0, 8.0),
    },
}
TYPE_ORDER = list(TYPES)

# Failure codes: description, repair downtime range (hours), and how the failure shows in the sensors as weights
# on (vibration, temperature, alarms) during the weeks before it.
FAILURE_CODES = {
    "SPN-BRG": ("Spindle bearing wear", (30.0, 60.0), (1.0, 0.6, 0.4), "spindle cartridge replaced"),
    "COOL-PMP": ("Coolant pump failure", (4.0, 10.0), (0.2, 1.0, 0.5), "coolant pump replaced"),
    "ATC-JAM": ("Automatic tool changer jam", (3.0, 9.0), (0.2, 0.1, 1.0), "tool changer cleared and re-timed"),
    "SRV-DRV": ("Servo drive fault", (6.0, 20.0), (0.3, 0.5, 1.0), "servo drive replaced and axis re-homed"),
    "CHK-HYD": ("Chuck hydraulic fault", (5.0, 14.0), (0.3, 0.7, 0.8), "chuck cylinder resealed"),
    "HYD-LEAK": ("Hydraulic leak", (6.0, 18.0), (0.2, 1.0, 0.4), "hydraulic line and seals replaced"),
    "PMP-WEAR": ("Hydraulic pump wear", (24.0, 48.0), (1.0, 0.7, 0.3), "hydraulic pump replaced"),
    "VLV-STK": ("Valve sticking", (4.0, 12.0), (0.2, 0.8, 0.8), "proportional valve cleaned and recalibrated"),
    "RDC-WEAR": ("Gear reducer wear", (24.0, 60.0), (1.0, 0.5, 0.4), "joint reducer replaced"),
    "CBL-FLT": ("Cable harness fault", (5.0, 14.0), (0.1, 0.1, 1.0), "dress pack cable replaced"),
    "WHL-IMB": ("Grinding wheel imbalance", (8.0, 20.0), (1.0, 0.2, 0.3), "wheel replaced and balanced"),
}

# Failures and scripted PM history that the reports refer to.
FORCED_FAILURES = {
    # The 8D report's machine: a spindle bearing that wore for four weeks while its PM-2 ran late.
    "MTY-LTH-04": [{"day": date(2026, 8, 24), "code": "SPN-BRG", "ramp": 28, "amp": 1.25, "downtime": 52.0}],
    # The scanned round sheet's press: a ram seal weeping on 14 Sep, leaking out three days later.
    "DAY-PRS-03": [{"day": date(2026, 9, 17), "code": "HYD-LEAK", "ramp": 21, "amp": 1.2, "downtime": 14.0}],
}
FORCED_PM = {"MTY-LTH-04": {"dsp_now": 21, "gaps": [148]}}  # last PM-2 on 9 Sep 2026, the one before on 14 Apr 2026
N_OVERDUE = 7  # machines whose last PM-2 is past the manual's interval at as_of (chosen with the seed)

# Parts: number, description, family, customer, unit cost (USD), rate factor, baseline scrap rate.
PARTS = [
    ("AP-3301", "Hydraulic valve body 3/4 in", "Valve bodies", "Brightline Hydraulics", 38.0, 1.00, 0.0115),
    ("AP-3312", "Proportional valve body PV-12", "Valve bodies", "Brightline Hydraulics", 64.0, 0.85, 0.0130),
    ("AP-3340", "Four-port manifold block", "Valve bodies", "Corvane Aerospace", 52.0, 0.90, 0.0120),
    ("AP-4410", "Transmission housing T4", "Transmission housings", "Nordhaven Motors", 118.0, 0.70, 0.0105),
    ("AP-4425", "Transmission end cover", "Transmission housings", "Kestrel Powertrain", 142.0, 0.75, 0.0115),
    ("AP-5102", "Motor mount bracket", "Brackets & mounts", "Tamsin Agricultural Equipment", 14.0, 1.25, 0.0080),
    ("AP-5118", "Pump support bracket", "Brackets & mounts", "Ostergaard Rail Systems", 19.0, 1.15, 0.0085),
    ("AP-5130", "Sensor mounting plate", "Brackets & mounts", "Kestrel Powertrain", 26.0, 1.10, 0.0085),
    ("AP-6205", "Gear pump body", "Pump components", "Brightline Hydraulics", 45.0, 0.90, 0.0100),
    ("AP-6220", "Pump cover plate", "Pump components", "Tamsin Agricultural Equipment", 71.0, 0.95, 0.0100),
]
FAMILIES = ["Valve bodies", "Transmission housings", "Brackets & mounts", "Pump components"]
FAMILY_MIX = {  # plant -> weights over FAMILIES
    "DAY": [0.00, 0.45, 0.35, 0.20],
    "MTY": [0.55, 0.00, 0.20, 0.25],
    "BRN": [0.35, 0.25, 0.20, 0.20],
}
BORE_PART = "AP-3312"
QUARTERS = [
    ("2025-Q3", "2025-07-01", "2025-10-01"),
    ("2025-Q4", "2025-10-01", "2026-01-01"),
    ("2026-Q1", "2026-01-01", "2026-04-01"),
    ("2026-Q2", "2026-04-01", "2026-07-01"),
    ("2026-Q3", "2026-07-01", "2026-10-01"),
]
SCRAP_TARGET_PCT = {"DAY": 1.4, "MTY": 1.2, "BRN": 1.6}
DOWNTIME_TARGET_HOURS = {"DAY": 210.0, "MTY": 160.0, "BRN": 140.0}
DEFECT_SCALE = 0.40  # calibrates the fleet to about 500 defect records over the history

DEFECT_TYPES = {
    "CNC Mill": (
        ["dimension_oob", "surface_finish", "burr", "thread_defect", "material_flaw"],
        [0.34, 0.26, 0.2, 0.12, 0.08],
    ),
    "CNC Lathe": (
        ["dimension_oob", "surface_finish", "burr", "thread_defect", "material_flaw"],
        [0.38, 0.24, 0.16, 0.14, 0.08],
    ),
    "Hydraulic Press": (
        ["crack", "dimension_oob", "burr", "flatness_oob", "material_flaw"],
        [0.28, 0.24, 0.2, 0.16, 0.12],
    ),
    "Robot Cell": (["weld_defect", "dimension_oob", "surface_finish", "material_flaw"], [0.55, 0.15, 0.2, 0.10]),
    "Surface Grinder": (["surface_finish", "flatness_oob", "dimension_oob", "material_flaw"], [0.4, 0.3, 0.22, 0.08]),
}
DEFECT_BASE = {
    "CNC Mill": 0.50,
    "CNC Lathe": 0.55,
    "Hydraulic Press": 0.32,
    "Robot Cell": 0.45,
    "Surface Grinder": 0.38,
}


@dataclass
class Machine:
    machine_id: str
    plant_id: str
    machine_type: str
    manufacturer: str
    model: str
    commissioned: date
    cell: str
    criticality: str
    weak: float
    v0: float
    t0: float
    overdue: bool = False
    events: list = field(default_factory=list)

    @property
    def tp(self) -> dict:
        return TYPES[self.machine_type]


def build_machines(rng: np.random.Generator) -> list[Machine]:
    new_machines = {"DAY-MIL-08": date(2025, 11, 10), "MTY-RBT-03": date(2026, 2, 16), "BRN-LTH-04": date(2026, 4, 20)}
    machines = []
    for pid in PLANT_IDS:
        counter = 0
        for mtype in TYPE_ORDER:
            tp = TYPES[mtype]
            for n in range(1, ROSTER[pid][mtype] + 1):
                mid = f"{pid}-{tp['code']}-{n:02d}"
                maker, model = tp["models"][int(rng.integers(0, len(tp["models"])))]
                if mid in new_machines:
                    comm = new_machines[mid]
                else:
                    comm = date(2011, 1, 1) + timedelta(days=int(rng.integers(0, 365 * 13)))
                crit = ["high", "medium", "low"][int(rng.choice(3, p=[0.30, 0.45, 0.25]))]
                weak = float(np.clip(rng.lognormal(0.0, 0.6), 0.3, 4.0))
                v0 = float(tp["v0"] * rng.lognormal(0.0, 0.12))
                t0 = float(tp["t0"] + rng.normal(0.0, 2.5))
                cell = f"Cell {'ABCDEFGH'[counter // 3]}"
                counter += 1
                machines.append(Machine(mid, pid, mtype, maker, model, comm, cell, crit, weak, v0, t0))
    return machines


# Simulation -----------------------------------------------------------------------------------------------------


def shift_open(i: int, mask: np.ndarray, step: int) -> int:
    while 0 <= i < NDAYS and not mask[i]:
        i += step
    return i


def make_pm(m: Machine, rng: np.random.Generator, mask: np.ndarray) -> list[int]:
    """PM-2 day indexes for one machine, built backwards from its last PM-2 before as_of."""
    interval = m.tp["pm2"]
    comm_i = day_index(m.commissioned)
    forced = FORCED_PM.get(m.machine_id)
    if forced:
        dsp_now = forced["dsp_now"]
    elif m.overdue:
        dsp_now = interval + int(rng.integers(8, 46))
    else:
        dsp_now = int(rng.integers(3, interval - 11))
    floor_i = comm_i + 20
    last = max(ASOF_I - dsp_now, floor_i)
    last = shift_open(last, mask, -1)
    dates = [last]
    gaps = list(forced["gaps"]) if forced else []
    d = last
    while d >= S0:
        if gaps:
            gap = gaps.pop(0)
        elif rng.random() < 0.05:
            gap = interval + int(rng.integers(5, 26))  # a late service
        else:
            gap = interval - int(rng.integers(2, 10))  # planners book a little early
        d = shift_open(d - gap, mask, -1)
        if d < floor_i:
            break
        dates.append(d)
    dates.reverse()
    if m.overdue:
        nxt = ASOF_I + int(rng.integers(3, 21))
    else:
        nxt = dates[-1] + interval - int(rng.integers(2, 10))
    while nxt <= END_I:
        dates.append(shift_open(nxt, mask, 1))
        nxt = dates[-1] + interval - int(rng.integers(2, 10))
    return dates


def ramp_profile(length: int, kind: str) -> np.ndarray:
    x = (np.arange(length) + 1.0) / length
    if kind == "failure":
        return x**1.6
    return 1.0 - np.abs(2.0 * x - 1.0)  # a near miss rises and falls back


def simulate_machine(m: Machine, k: int, open_mask: np.ndarray, climate: float) -> dict:
    rng_pm = np.random.default_rng([SEED, k, 1])
    rng_f = np.random.default_rng([SEED, k, 2])
    rng_s = np.random.default_rng([SEED, k, 3])
    tp = m.tp
    interval = tp["pm2"]
    comm_i = day_index(m.commissioned)
    first_i = max(S0, comm_i)

    pm = make_pm(m, rng_pm, open_mask)
    pm_hours = {d: round(float(rng_pm.uniform(*tp["pm_hours"])), 1) for d in pm if S0 <= d <= ASOF_I}
    anchors = np.array(sorted([*pm, comm_i]))

    open_idx = np.nonzero(open_mask)[0]
    open_idx = open_idx[open_idx >= first_i]
    pos = np.searchsorted(anchors, open_idx, side="right") - 1
    dsp = open_idx - anchors[np.maximum(pos, 0)]
    stress = np.minimum(dsp / interval, 2.0)
    age = (open_idx - comm_i) / 365.25
    agef = np.clip(1.0 + 0.03 * (age - 6.0), 0.8, 1.5)
    hazard = tp["hz"] * m.weak * agef * np.exp(2.0 * (stress - 0.5))
    u = rng_f.random(len(open_idx))

    forced = {}
    for f in FORCED_FAILURES.get(m.machine_id, []):
        forced[day_index(f["day"])] = f
    blocked = np.zeros(len(open_idx), dtype=bool)
    for fi in forced:
        blocked |= (open_idx >= fi - 75) & (open_idx <= fi + 75)

    failures = []  # day index, code, ramp length, amplitude, downtime hours, hour of day
    down = np.zeros(NDAYS, dtype=bool)
    last_fail = -100
    p = 0
    while p < len(open_idx):
        i = int(open_idx[p])
        spec = forced.get(i)
        hit = spec is not None
        if not hit and not blocked[p] and i - last_fail >= 14 and u[p] < hazard[p]:
            hit = True
        if not hit:
            p += 1
            continue
        if spec:
            code, ramp, amp, hours = spec["code"], spec["ramp"], spec["amp"], spec["downtime"]
        else:
            code = tp["codes"][int(rng_f.choice(len(tp["codes"]), p=tp["cw"]))]
            ramp = int(rng_f.integers(18, 36))
            amp = float(rng_f.uniform(0.9, 1.4))
            lo, hi = FAILURE_CODES[code][1]
            hours = round(float(rng_f.uniform(lo, hi)), 1)
        hour = int(rng_f.integers(6, 22))
        minute = int(rng_f.integers(0, 60))
        failures.append((i, code, ramp, amp, hours, hour, minute))
        n_down = max(0, math.ceil(hours / 16.0) - 1)
        for q in range(1, n_down + 1):
            if p + q < len(open_idx):
                down[open_idx[p + q]] = True
        last_fail = i
        p += 1 + n_down

    # Sensors, for every calendar day, then masked to the days the machine ran.
    load = np.clip(rng_s.normal(0.72, 0.12, NDAYS), 0.30, 1.00)
    vib_noise = rng_s.lognormal(0.0, 0.055, NDAYS)
    temp_noise = rng_s.normal(0.0, 1.1, NDAYS)
    alarm_u = rng_s.random(NDAYS)
    vib_add = np.zeros(NDAYS)
    tmp_add = np.zeros(NDAYS)
    alm_add = np.zeros(NDAYS)

    def add_ramp(end_i: int, length: int, amp: float, weights: tuple, kind: str) -> None:
        lo = end_i - length
        days = np.arange(lo, end_i)
        keep = days >= 0
        prof = ramp_profile(length, kind)[keep] * amp
        days = days[keep]
        vib_add[days] += m.v0 * 1.3 * prof * weights[0]
        tmp_add[days] += 11.0 * prof * weights[1]
        alm_add[days] += 3.5 * prof * weights[2]

    for i, code, ramp, amp, _hours, _h, _mi in failures:
        add_ramp(i, ramp, amp, FAILURE_CODES[code][2], "failure")
    n_near = int(rng_s.poisson(0.9))
    near_misses = []
    for _ in range(n_near):
        centre = int(rng_s.choice(open_idx))
        length = int(rng_s.integers(10, 21))
        peak = float(rng_s.uniform(0.35, 0.6))
        code = tp["codes"][int(rng_s.integers(0, len(tp["codes"])))]
        add_ramp(centre + length // 2, length, peak, FAILURE_CODES[code][2], "near")
        near_misses.append((centre, code))

    service = np.array(sorted({*pm, *[f[0] for f in failures], comm_i}))
    spos = np.searchsorted(service, np.arange(NDAYS), side="right") - 1
    since = np.arange(NDAYS) - service[np.maximum(spos, 0)]
    stress_w = np.minimum(since / interval, 1.6)

    season = 2.2 * np.sin(2 * np.pi * (np.arange(NDAYS) - (day_index(date(2026, 4, 20)))) / 365.25)
    vib = m.v0 * (1.0 + 0.30 * (load - 0.72)) * (1.0 + 0.10 * stress_w) * vib_noise + vib_add
    temp = m.t0 + 9.0 * (load - 0.72) + season + climate * 0.5 + 1.5 * stress_w + temp_noise + tmp_add
    lam = tp["a0"] * (1.0 + 0.8 * (load - 0.72)) + 0.6 * np.maximum(0.0, vib / m.v0 - 1.35) + alm_add
    alarms = np.zeros(NDAYS, dtype=int)
    # Inverse-CDF Poisson from one uniform a day keeps the stream independent of the masks above.
    for i in range(NDAYS):
        alarms[i] = _poisson_inv(lam[i], alarm_u[i])
    hours = np.round(16.0 * load, 1)

    fail_day = {f[0]: f for f in failures}
    rows = []
    for day_i in open_idx:
        i = int(day_i)
        if i > ASOF_I or down[i]:
            continue
        h = float(hours[i])
        a = int(alarms[i])
        if i in pm_hours:
            h = max(1.0, round(h - pm_hours[i], 1))
        if i in fail_day:
            h = round(float(rng_s.uniform(1.0, 7.0)), 1)
            a += 3
        rows.append((i, round(float(vib[i]), 2), round(float(temp[i]), 1), h, a))

    fut_alarms = int(sum(alarms[int(i)] for i in open_idx if ASOF_I < i <= ASOF_I + 30 and not down[int(i)]))

    events = []
    for d in pm:
        if S0 <= d <= ASOF_I:
            events.append(
                {
                    "i": d,
                    "event_type": "planned",
                    "task_code": "PM-2",
                    "failure_code": None,
                    "downtime_hours": pm_hours[d],
                    "hour": int(rng_pm.integers(6, 14)),
                    "minute": int(rng_pm.integers(0, 60)),
                    "cost": round(
                        float(
                            rng_pm.uniform(0.6, 1.4)
                            * {
                                "CNC Mill": 1450,
                                "CNC Lathe": 1650,
                                "Hydraulic Press": 980,
                                "Robot Cell": 2100,
                                "Surface Grinder": 1200,
                            }[m.machine_type]
                        ),
                        2,
                    ),
                }
            )
    for i, code, _ramp, _amp, hours_down, hour, minute in failures:
        if i <= ASOF_I:
            events.append(
                {
                    "i": i,
                    "event_type": "unplanned",
                    "task_code": "CM",
                    "failure_code": code,
                    "downtime_hours": hours_down,
                    "hour": hour,
                    "minute": minute,
                    "cost": round(float(rng_f.uniform(0.5, 1.5) * hours_down * 85.0 + rng_f.uniform(150, 900)), 2),
                }
            )
    return {
        "rows": rows,
        "events": events,
        "failures": failures,
        "pm": [d for d in pm if S0 <= d <= ASOF_I],
        "pm_all": pm,
        "near_misses": near_misses,
        "fut_alarms": fut_alarms,
    }


def _poisson_inv(lam: float, u: float) -> int:
    k = 0
    p = math.exp(-lam)
    c = p
    while u > c and k < 40:
        k += 1
        p *= lam / k
        c += p
    return k


# Tables ---------------------------------------------------------------------------------------------------------


def build_tables() -> dict:
    rng = np.random.default_rng([SEED, 0])
    machines = build_machines(rng)
    open_masks = build_open_masks()

    # Pick the machines whose last PM-2 is past the manual's interval at as_of.
    cand = [m for m in machines if m.machine_id not in FORCED_PM and day_index(m.commissioned) < S0]
    while True:  # a spread across plants and types, drawn with the seed
        pick = sorted(int(x) for x in rng.choice(len(cand), size=N_OVERDUE, replace=False))
        by_plant = [sum(cand[q].plant_id == p for q in pick) for p in PLANT_IDS]
        by_type = [sum(cand[q].machine_type == t for q in pick) for t in TYPE_ORDER]
        if max(by_plant) <= 3 and min(by_plant) >= 2 and max(by_type) <= 2:
            break
    for q in pick:
        cand[q].overdue = True

    sims = {}
    plant_climate = {p["plant_id"]: p["climate"] for p in PLANTS}
    for k, m in enumerate(machines):
        sims[m.machine_id] = simulate_machine(m, k + 1, open_masks[m.plant_id], plant_climate[m.plant_id])

    plants_df = pd.DataFrame([{k: v for k, v in p.items() if k not in ("climate", "plan_factor")} for p in PLANTS])
    machines_df = pd.DataFrame(
        [
            {
                "machine_id": m.machine_id,
                "plant_id": m.plant_id,
                "machine_type": m.machine_type,
                "manufacturer": m.manufacturer,
                "model": m.model,
                "cell": m.cell,
                "criticality": m.criticality,
                "commissioned_date": m.commissioned.isoformat(),
            }
            for m in machines
        ]
    )

    sensor_rows = []
    for m in machines:
        for i, vib, temp, hrs, alm in sims[m.machine_id]["rows"]:
            sensor_rows.append((day_date(i).isoformat(), m.machine_id, vib, temp, hrs, alm))
    sensor_df = pd.DataFrame(
        sensor_rows,
        columns=[
            "reading_date",
            "machine_id",
            "avg_vibration_mm_s",
            "max_temperature_c",
            "spindle_hours",
            "alarm_count",
        ],
    )
    sensor_df = sensor_df.sort_values(["reading_date", "machine_id"], kind="mergesort").reset_index(drop=True)
    sensor_df.insert(0, "reading_id", np.arange(1, len(sensor_df) + 1))

    ev_rows = []
    for m in machines:
        for e in sims[m.machine_id]["events"]:
            ts = datetime.combine(day_date(e["i"]), datetime.min.time()) + timedelta(
                hours=e["hour"], minutes=e["minute"]
            )
            if e["event_type"] == "planned":
                note = f"PM-2 full service completed on the {m.machine_type.lower()}"
            else:
                note = f"{FAILURE_CODES[e['failure_code']][0]}: {FAILURE_CODES[e['failure_code']][3]}"
            ev_rows.append(
                (
                    ts.strftime("%Y-%m-%d %H:%M:%S"),
                    m.machine_id,
                    e["event_type"],
                    e["task_code"],
                    e["failure_code"],
                    e["downtime_hours"],
                    e["cost"],
                    note,
                )
            )
    events_df = pd.DataFrame(
        ev_rows,
        columns=[
            "started_at",
            "machine_id",
            "event_type",
            "task_code",
            "failure_code",
            "downtime_hours",
            "parts_and_labor_cost_usd",
            "notes",
        ],
    )
    events_df = events_df.sort_values(["started_at", "machine_id"], kind="mergesort").reset_index(drop=True)
    events_df.insert(0, "event_id", [f"ME-{n:05d}" for n in range(1, len(events_df) + 1)])

    wo_df, defects_df, daily = build_work_orders(machines, sims)
    plan_df, targets_df = build_plan(machines, daily, open_masks)

    hidden = []
    for m in machines:
        for i, code, ramp, _amp, hours_down, _h, _mi in sims[m.machine_id]["failures"]:
            if ASOF_I < i <= ASOF_I + 30:
                hidden.append((m.machine_id, day_date(i).isoformat(), code, ramp, hours_down))
    return {
        "machines": machines,
        "sims": sims,
        "plants": plants_df,
        "machines_df": machines_df,
        "sensor": sensor_df,
        "events": events_df,
        "work_orders": wo_df,
        "defects": defects_df,
        "plan": plan_df,
        "targets": targets_df,
        "hidden_failures": sorted(hidden, key=lambda r: (r[1], r[0])),
        "future_alarms": {m.machine_id: sims[m.machine_id]["fut_alarms"] for m in machines},
        "open_masks": open_masks,
    }


def build_work_orders(machines: list[Machine], sims: dict) -> tuple[pd.DataFrame, pd.DataFrame, dict]:
    rng = np.random.default_rng([SEED, 9001])
    part_by_family = {f: [p for p in PARTS if p[2] == f] for f in FAMILIES}
    parts = {p[0]: p for p in PARTS}
    wo_rows = []
    for m in machines:
        rows = sims[m.machine_id]["rows"]
        if not rows:
            continue
        mix = np.array(FAMILY_MIX[m.plant_id])
        pos = 0
        while pos < len(rows):
            n = int(rng.integers(14, 27))
            chunk = rows[pos : pos + n]
            pos += n
            first_day, last_day = day_date(chunk[0][0]), day_date(chunk[-1][0])
            fam = FAMILIES[int(rng.choice(len(FAMILIES), p=mix))]
            options = part_by_family[fam]
            part = options[int(rng.integers(0, len(options)))]
            if m.machine_id == "MTY-LTH-04" and last_day >= date(2026, 7, 20) and first_day <= date(2026, 8, 24):
                part = parts[BORE_PART]
            wo_rows.append({"machine": m, "chunk": chunk, "planned_days": n, "part": part})

    # Number the work orders in start order.
    wo_rows.sort(key=lambda w: (w["chunk"][0][0], w["machine"].machine_id))
    records = []
    daily = {}  # (plant, day index) -> good units that day
    defect_rows = []
    for number, w in enumerate(wo_rows, start=1):
        m = w["machine"]
        part = w["part"]
        chunk = w["chunk"]
        pn, desc, fam, cust, unit_cost, rate_f, scrap_p = part
        units_by_day = []
        for i, _v, _t, hrs, _a in chunk:
            units_by_day.append(m.tp["rate"] * rate_f * hrs / 16.0 * float(rng.normal(1.0, 0.03)))
        qty_started = int(round(sum(units_by_day)))
        qty_started = max(qty_started, 20)
        # Defects follow the machine's condition: each day's weight grows with its vibration and temperature.
        day_w = np.array(
            [
                u * math.exp(1.6 * (c[1] / m.v0 - 1.0)) * math.exp(0.04 * max(c[2] - m.t0, -4.0))
                for c, u in zip(chunk, units_by_day, strict=True)
            ]
        )
        start_i, end_i = chunk[0][0], chunk[-1][0]
        completed = len(chunk) == w["planned_days"]
        start_d, end_d = day_date(start_i), day_date(end_i)
        wo_id = f"WO-{number:05d}"
        due_d = start_d + timedelta(days=int(math.ceil(w["planned_days"] * 7 / 5)) + int(rng.integers(0, 4)))

        lam = DEFECT_BASE[m.machine_type] * float(day_w.sum()) / 1000.0 * DEFECT_SCALE
        n_def = int(rng.poisson(min(lam, 14.0)))
        scrap_defect_qty = 0
        defects = []
        types, tw = DEFECT_TYPES[m.machine_type]
        for _ in range(n_def):
            dtype = types[int(rng.choice(len(types), p=np.array(tw) / sum(tw)))]
            sev = ["minor", "major", "critical"][int(rng.choice(3, p=[0.58, 0.32, 0.10]))]
            if sev == "minor":
                qty = int(rng.integers(1, 11))
            elif sev == "major":
                qty = int(rng.integers(8, 46))
            else:
                qty = int(rng.integers(20, 91))
            qty = max(1, min(qty, int(qty_started * 0.4)))
            if dtype == "material_flaw":
                disp = "return_to_supplier" if rng.random() < 0.9 else "scrap"
                stage = "incoming" if rng.random() < 0.7 else "in_process"
            else:
                probs = {"minor": [0.25, 0.50, 0.25], "major": [0.55, 0.45, 0.0], "critical": [0.85, 0.15, 0.0]}[sev]
                disp = ["scrap", "rework", "use_as_is"][int(rng.choice(3, p=probs))]
                stage = "in_process" if rng.random() < 0.65 else "final_inspection"
            day_i = chunk[int(rng.choice(len(chunk), p=day_w / day_w.sum()))][0]
            ts = datetime.combine(day_date(day_i), datetime.min.time()) + timedelta(
                hours=int(rng.integers(7, 22)), minutes=int(rng.integers(0, 60))
            )
            cost = {"scrap": 1.0, "rework": 0.3}.get(disp, 0.0) * qty * unit_cost
            if disp == "scrap":
                scrap_defect_qty += qty
            defects.append([ts, m.machine_id, wo_id, dtype, sev, stage, qty, disp, round(cost, 2)])
        if m.machine_id == "MTY-LTH-04" and pn == BORE_PART:
            # The 8D case: bores drifting out of tolerance while the spindle bearing wears (3 to 21 August).
            for c in chunk:
                if date(2026, 8, 3) <= day_date(c[0]) <= date(2026, 8, 21) and rng.random() < 0.75:
                    sev = "major" if rng.random() < 0.7 else "minor"
                    qty = int(rng.integers(12, 46)) if sev == "major" else int(rng.integers(3, 11))
                    disp = "scrap" if rng.random() < 0.4 else "rework"
                    ts = datetime.combine(day_date(c[0]), datetime.min.time()) + timedelta(
                        hours=int(rng.integers(7, 22)), minutes=int(rng.integers(0, 60))
                    )
                    cost = {"scrap": 1.0, "rework": 0.3}[disp] * qty * unit_cost
                    if disp == "scrap":
                        scrap_defect_qty += qty
                    defects.append(
                        [ts, m.machine_id, wo_id, "dimension_oob", sev, "in_process", qty, disp, round(cost, 2)]
                    )
        base_scrap = int(rng.binomial(qty_started, scrap_p))
        qty_scrap = min(base_scrap + scrap_defect_qty, int(qty_started * 0.6))
        qty_good = qty_started - qty_scrap
        defect_rows.extend(defects)
        good_frac = qty_good / qty_started
        for (i, *_), units in zip(chunk, units_by_day, strict=True):
            daily[(m.plant_id, i)] = daily.get((m.plant_id, i), 0.0) + units * good_frac
        records.append(
            {
                "work_order_id": wo_id,
                "plant_id": m.plant_id,
                "machine_id": m.machine_id,
                "part_number": pn,
                "product_family": fam,
                "customer": cust,
                "start_date": start_d.isoformat(),
                "due_date": due_d.isoformat(),
                "completed_date": end_d.isoformat() if completed else None,
                "status": "completed" if completed else "in_progress",
                "qty_started": qty_started,
                "qty_good": qty_good,
                "qty_scrap": qty_scrap,
            }
        )

    # The one escape to a customer: a lot shipped from the cluster's first work order and returned on 14 Aug.
    escape_wo = next(
        r
        for r in records
        if r["machine_id"] == "MTY-LTH-04"
        and r["part_number"] == BORE_PART
        and r["start_date"] <= "2026-08-03" <= (r["completed_date"] or "2026-09-30")
    )
    defect_rows.append(
        [
            datetime(2026, 8, 14, 9, 40),
            "MTY-LTH-04",
            escape_wo["work_order_id"],
            "dimension_oob",
            "critical",
            "customer_return",
            48,
            "scrap",
            round(48 * 64.0, 2),
        ]
    )
    wo_df = pd.DataFrame(records)
    d_df = pd.DataFrame(
        defect_rows,
        columns=[
            "detected_at",
            "machine_id",
            "work_order_id",
            "defect_type",
            "severity",
            "detection_stage",
            "quantity_affected",
            "disposition",
            "cost_usd",
        ],
    )
    d_df = d_df.sort_values(["detected_at", "machine_id", "work_order_id"], kind="mergesort").reset_index(drop=True)
    d_df["detected_at"] = d_df["detected_at"].dt.strftime("%Y-%m-%d %H:%M:%S")
    d_df.insert(0, "defect_id", [f"QD-{n:05d}" for n in range(1, len(d_df) + 1)])
    # The escape's quantity counts as scrap on its work order.
    mask = wo_df["work_order_id"] == escape_wo["work_order_id"]
    wo_df.loc[mask, "qty_scrap"] += 48
    wo_df.loc[mask, "qty_good"] -= 48
    return wo_df, d_df, daily


def build_plan(machines, daily, open_masks):
    machines_by_plant = {p: [m for m in machines if m.plant_id == p] for p in PLANT_IDS}
    plan_rows = []
    week = date(2025, 7, 7)
    last_week = date(2026, 9, 21)
    pf = {p["plant_id"]: p["plan_factor"] for p in PLANTS}
    while week <= last_week:
        iso = week.isocalendar()
        for pid in PLANT_IDS:
            days = [day_index(week + timedelta(days=k)) for k in range(7)]
            working = int(sum(open_masks[pid][d] for d in days))
            cap = 0.0
            for m in machines_by_plant[pid]:
                if day_index(m.commissioned) <= days[0]:
                    cap += m.tp["rate"]
            target = int(round(cap * 0.72 * pf[pid] * working / 100.0) * 100)
            actual = int(round(sum(daily.get((pid, d), 0.0) for d in days)))
            plan_rows.append(
                {
                    "plan_line_id": f"{pid}-{iso.year}-W{iso.week:02d}",
                    "week_start": week.isoformat(),
                    "plant_id": pid,
                    "working_days": working,
                    "target_units": target,
                    "actual_units": actual,
                }
            )
        week += timedelta(days=7)
    plan_df = pd.DataFrame(plan_rows)

    rows = []
    for pid in PLANT_IDS:
        for q, _lo, _hi in QUARTERS:
            rows.append(
                {
                    "target_id": f"{pid}-{q.replace('-', '')}",
                    "plant_id": pid,
                    "quarter": q,
                    "scrap_target_pct": SCRAP_TARGET_PCT[pid],
                    "unplanned_downtime_target_hours": DOWNTIME_TARGET_HOURS[pid],
                }
            )
    return plan_df, pd.DataFrame(rows)


# Document content -----------------------------------------------------------------------------------------------

COMPANY = "Atlas Precision Components"
Q3 = ("2026-07-01", "2026-10-01")
Q3_WEEKS = ("2026-07-06", "2026-09-21")  # week_start of the 12 full weeks of the quarter in the production plan
LAST12 = ("2025-10-01", "2026-10-01")
BORE_MACHINE = "MTY-LTH-04"
SCAN_MACHINE = "DAY-PRS-03"
SCAN_DATE = "2026-09-14"
PLANT_KEY = {"DAY": "dayton", "MTY": "monterrey", "BRN": "brno"}

# Condition-based triggers: vibration warning and alarm (mm/s RMS), maximum temperature warning and alarm (C).
THRESHOLDS = {
    "CNC Mill": (3.0, 4.0, 60, 70),
    "CNC Lathe": (3.5, 4.5, 62, 72),
    "Hydraulic Press": (4.0, 5.0, 62, 70),
    "Robot Cell": (2.5, 3.5, 55, 65),
    "Surface Grinder": (2.5, 3.5, 58, 68),
}
PM_TEXT = {  # PM-1 frequency, fluid and lubricant analysis, geometry and calibration check
    "CNC Mill": ("Daily", "Coolant concentration and pH weekly", "Ballbar circularity check at every PM-2"),
    "CNC Lathe": ("Daily", "Coolant concentration and pH weekly", "Chuck runout and turret indexing at every PM-2"),
    "Hydraulic Press": ("Daily", "Hydraulic oil sample every 90 days", "Platen parallelism at every second PM-2"),
    "Robot Cell": ("Weekly", "Reducer grease sample at every PM-2", "Tool center point calibration at every PM-2"),
    "Surface Grinder": ("Daily", "Coolant concentration weekly", "Wheel balance and spindle runout at every PM-2"),
}
LOTO_SOURCES = {
    "CNC Mill": (
        "Electrical main supply; pneumatic air; coolant and lubrication pumps; stored energy in servo drives",
        "Main disconnect; air supply valve with dump; pump breakers; wait 5 minutes for drive discharge",
    ),
    "CNC Lathe": (
        "Electrical main supply; hydraulic chuck and turret power pack; pneumatic air; coolant pump",
        "Main disconnect; hydraulic power pack breaker and bleed; air supply valve with dump",
    ),
    "Hydraulic Press": (
        "Electrical main supply; hydraulic power unit and accumulators; ram and platen (gravity); "
        "pneumatic die cushion",
        "Main disconnect; hydraulic supply and accumulator bleed valves; ram safety props; die cushion air valve",
    ),
    "Robot Cell": (
        "Electrical controller and weld power supply; pneumatic gripper or torch supply; stored energy in arm joints",
        "Controller and weld power disconnects; air supply valve; fixture clamps blocked; fence key removed",
    ),
    "Surface Grinder": (
        "Electrical main supply; coolant pump; magnetic chuck; rotating grinding wheel (inertia)",
        "Main disconnect; pump breaker; chuck switch; wait until the wheel has stopped turning",
    ),
}
SQA_TARGETS = [
    ("Defective parts per million (PPM) delivered", "500 or fewer", "Quarterly, from incoming and in-process findings"),
    ("Lot rejection rate at incoming inspection", "1.0 percent of lots or less", "Quarterly"),
    ("On-time delivery", "95 percent of line items or more", "Within one business day of the confirmed date"),
    ("Containment after notice of a nonconformance", "Within 24 hours", "Per event"),
    ("8D report after notice of a nonconformance", "Within 10 business days", "Per event"),
    ("Mill test certificate with each lot", "100 percent of lots", "Per delivery"),
]


def fmt_int(n: float) -> str:
    return f"{int(round(n)):,}"


def fmt1(x: float) -> str:
    return f"{x:,.1f}"


def long_date(d) -> str:
    if isinstance(d, str):
        d = date.fromisoformat(d[:10])
    return f"{d.day} {d.strftime('%B %Y')}"


def day_month(d) -> str:
    if isinstance(d, str):
        d = date.fromisoformat(d[:10])
    return f"{d.day} {d.strftime('%B')}"


def join_and(items: list[str]) -> str:
    return items[0] if len(items) == 1 else ", ".join(items[:-1]) + " and " + items[-1]


def compute_facts(t: dict) -> tuple[dict, dict, dict]:
    """The figures the documents and the README quote, computed from the generated tables."""
    m = t["machines_df"]
    wo = t["work_orders"]
    plan = t["plan"]
    sensor = t["sensor"]
    mp = m[["machine_id", "plant_id", "machine_type"]]
    ev = t["events"].merge(mp, on="machine_id")
    df = t["defects"].merge(mp, on="machine_id").merge(wo[["work_order_id", "part_number"]], on="work_order_id")
    targets = t["targets"]
    f: dict = {"machine_count": str(len(m))}
    raw: dict = {}
    tables: dict = {}
    charts: dict = {}

    unplanned = ev[ev["event_type"] == "unplanned"]
    q3_un = unplanned[(unplanned["started_at"] >= Q3[0]) & (unplanned["started_at"] < Q3[1])]
    q3_pm2 = ev[(ev["task_code"] == "PM-2") & (ev["started_at"] >= Q3[0]) & (ev["started_at"] < Q3[1])]
    done = wo[(wo["status"] == "completed") & (wo["completed_date"] >= Q3[0]) & (wo["completed_date"] < Q3[1])]
    q3_plan = plan[(plan["week_start"] >= Q3_WEEKS[0]) & (plan["week_start"] <= Q3_WEEKS[1])]
    q3_def = df[(df["detected_at"] >= Q3[0]) & (df["detected_at"] < Q3[1])]

    score_rows = []
    for p in PLANTS:
        pid, key = p["plant_id"], PLANT_KEY[p["plant_id"]]
        tgt = targets[(targets["plant_id"] == pid) & (targets["quarter"] == "2026-Q3")].iloc[0]
        pl = q3_plan[q3_plan["plant_id"] == pid]
        actual, target = int(pl["actual_units"].sum()), int(pl["target_units"].sum())
        dn = done[done["plant_id"] == pid]
        scrap = 100.0 * dn["qty_scrap"].sum() / dn["qty_started"].sum()
        un = q3_un[q3_un["plant_id"] == pid]
        dcost = float(q3_def[q3_def["plant_id"] == pid]["cost_usd"].sum())
        f[f"{key}_machines"] = str(int((m["plant_id"] == pid).sum()))
        f[f"{key}_actual"] = fmt_int(actual)
        f[f"{key}_target"] = fmt_int(target)
        f[f"{key}_attain"] = f"{100.0 * actual / target:.1f}"
        f[f"{key}_scrap"] = f"{scrap:.2f}"
        f[f"{key}_scrap_target"] = f"{tgt['scrap_target_pct']:.1f}"
        f[f"{key}_ue"] = str(len(un))
        f[f"{key}_dt"] = fmt1(un["downtime_hours"].sum())
        f[f"{key}_dt_target"] = fmt_int(tgt["unplanned_downtime_target_hours"])
        f[f"{key}_def_cost"] = fmt_int(dcost)
        f[f"{key}_pm2"] = str(int((q3_pm2["plant_id"] == pid).sum()))
        raw[pid] = {
            "attain": 100.0 * actual / target,
            "scrap": scrap,
            "dt": float(un["downtime_hours"].sum()),
            "ue": len(un),
        }
        raw[pid]["scrap_target"] = float(tgt["scrap_target_pct"])
        raw[pid]["dt_target"] = float(tgt["unplanned_downtime_target_hours"])
        top = un.groupby("failure_code")["downtime_hours"].sum().sort_values(ascending=False)
        f[f"{key}_top_code"] = str(top.index[0])
        f[f"{key}_top_code_desc"] = FAILURE_CODES[str(top.index[0])][0].lower()
        f[f"{key}_top_hours"] = fmt1(top.iloc[0])
        months = {}
        for mo, lo, hi in (
            ("July", "2026-07-01", "2026-08-01"),
            ("August", "2026-08-01", "2026-09-01"),
            ("September", "2026-09-01", "2026-10-01"),
        ):
            months[mo] = float(un[(un["started_at"] >= lo) & (un["started_at"] < hi)]["downtime_hours"].sum())
        raw[pid]["months"] = months
        longest = un.sort_values("downtime_hours").iloc[-1]
        f[f"{key}_longest"] = (
            f"{longest['failure_code']} on {longest['machine_id']} on {day_month(longest['started_at'])} "
            f"({longest['downtime_hours']:.1f} hours)"
        )
        score_rows.append(
            [
                p["plant_name"],
                f"{f[f'{key}_attain']}%",
                f"{f[f'{key}_scrap']}% (target {f[f'{key}_scrap_target']}%)",
                f[f"{key}_ue"],
                f"{f[f'{key}_dt']} h (target {f[f'{key}_dt_target']} h)",
                f"${f[f'{key}_def_cost']}",
            ]
        )
    a_tot, t_tot = int(q3_plan["actual_units"].sum()), int(q3_plan["target_units"].sum())
    s_tot = 100.0 * done["qty_scrap"].sum() / done["qty_started"].sum()
    f["q3_units_total"], f["q3_target_total"] = fmt_int(a_tot), fmt_int(t_tot)
    f["q3_attain_total"] = f"{100.0 * a_tot / t_tot:.1f}"
    f["q3_scrap_total"] = f"{s_tot:.2f}"
    f["q3_ue_total"] = str(len(q3_un))
    f["q3_dt_total"] = fmt1(q3_un["downtime_hours"].sum())
    f["q3_pm2_total"] = str(len(q3_pm2))
    f["q3_def_records"] = str(len(q3_def))
    f["q3_def_cost"] = fmt_int(q3_def["cost_usd"].sum())
    score_rows.append(
        [
            "All plants",
            f"{f['q3_attain_total']}%",
            f"{f['q3_scrap_total']}%",
            f["q3_ue_total"],
            f"{f['q3_dt_total']} h",
            f"${f['q3_def_cost']}",
        ]
    )
    tables["scorecard"] = (
        ["Plant", "Plan attainment", "Scrap rate", "Unplanned failures", "Unplanned downtime", "Defect cost"],
        score_rows,
    )
    charts["downtime"] = (
        ["July", "August", "September"],
        [(PLANT_NAME[pid], [raw[pid]["months"][mo] for mo in ("July", "August", "September")]) for pid in PLANT_IDS],
    )
    # Claims the deck makes in words: fail loudly if the data stops supporting them.
    assert raw["MTY"]["scrap"] > raw["MTY"]["scrap_target"] + 0.1
    assert (
        raw["DAY"]["scrap"] < raw["DAY"]["scrap_target"] - 0.1
        and raw["BRN"]["scrap"] < raw["BRN"]["scrap_target"] - 0.1
    )
    assert raw["DAY"]["dt"] > raw["DAY"]["dt_target"] + 10 and raw["MTY"]["dt"] > raw["MTY"]["dt_target"] + 10
    assert raw["BRN"]["dt"] < raw["BRN"]["dt_target"] - 10
    assert raw["DAY"]["dt"] > raw["MTY"]["dt"] * 1.05 > raw["BRN"]["dt"]

    by_code = q3_un.groupby("failure_code")["downtime_hours"].agg(["sum", "count"]).sort_values("sum", ascending=False)
    total_h = float(by_code["sum"].sum())
    f["q3_spn_share"] = f"{100.0 * float(by_code.loc['SPN-BRG', 'sum']) / total_h:.0f}"
    assert by_code.index[0] == "SPN-BRG"
    f["top_code_1"], f["top_code_2"] = str(by_code.index[0]), str(by_code.index[1])
    f["top2_share"] = f"{100.0 * float(by_code['sum'].iloc[:2].sum()) / total_h:.0f}"
    tables["failure_top"] = (
        ["Failure code", "Failure", "Events", "Downtime hours"],
        [[c, FAILURE_CODES[c][0], str(int(r["count"])), fmt1(r["sum"])] for c, r in by_code.head(5).iterrows()],
    )
    dc = q3_def.groupby("defect_type").agg(
        n=("defect_id", "count"), q=("quantity_affected", "sum"), c=("cost_usd", "sum")
    )
    dc = dc.sort_values("c", ascending=False)
    tables["defect_cost"] = (
        ["Defect type", "Records", "Pieces affected", "Cost (USD)"],
        [[str(k), str(int(r["n"])), fmt_int(r["q"]), fmt_int(r["c"])] for k, r in dc.head(5).iterrows()],
    )
    mat = q3_def[q3_def["defect_type"] == "material_flaw"]
    f["q3_mat_records"] = str(len(mat))
    f["q3_mat_qty"] = fmt_int(mat["quantity_affected"].sum())
    f["q3_mat_rts"] = str(int((mat["disposition"] == "return_to_supplier").sum()))

    # The 8D case ------------------------------------------------------------------------------------------------
    bore = df[(df["machine_id"] == BORE_MACHINE) & (df["defect_type"] == "dimension_oob")]
    bore = bore[(bore["detected_at"] >= Q3[0]) & (bore["detected_at"] < Q3[1])]
    assert (bore["part_number"] == BORE_PART).all() and bore["detected_at"].min() >= "2026-08-01"
    assert bore["detected_at"].max() < "2026-08-22"
    inplant = bore[bore["detection_stage"] != "customer_return"]
    escape = bore[bore["detection_stage"] == "customer_return"]
    assert len(escape) == 1
    assert set(inplant["disposition"]) <= {"scrap", "rework"}
    f["escape_qty"] = str(int(escape["quantity_affected"].sum()))
    f["n_inplant"] = str(len(inplant))
    f["inplant_qty"] = str(int(inplant["quantity_affected"].sum()))
    f["inplant_scrap_qty"] = str(int(inplant[inplant["disposition"] == "scrap"]["quantity_affected"].sum()))
    f["inplant_rework_qty"] = str(int(inplant[inplant["disposition"] == "rework"]["quantity_affected"].sum()))
    f["total_cost"] = f"{float(bore['cost_usd'].sum()):,.2f}"
    f["eight_d_cost"] = fmt_int(bore["cost_usd"].sum())
    f["cluster_first"] = day_month(inplant["detected_at"].min())
    f["cluster_last"] = day_month(inplant["detected_at"].max())
    wos = sorted(set(inplant["work_order_id"]))
    f["wo_list"] = join_and(wos)
    mach = m.set_index("machine_id").loc[BORE_MACHINE]
    f["bore_cell"] = str(mach["cell"])
    raw["eight_d"] = {"n_inplant": len(inplant), "escape_wo": str(escape["work_order_id"].iloc[0]), "wos": wos}

    ev_b = ev[ev["machine_id"] == BORE_MACHINE]
    failure = ev_b[(ev_b["event_type"] == "unplanned") & (ev_b["started_at"] >= "2026-08-01")].iloc[0]
    fail_day = date.fromisoformat(failure["started_at"][:10])
    assert failure["failure_code"] == "SPN-BRG" and fail_day == date(2026, 8, 24)
    pm_b = ev_b[ev_b["task_code"] == "PM-2"]["started_at"].str[:10]
    last_pm = date.fromisoformat(pm_b[pm_b < fail_day.isoformat()].max())
    after_pm = date.fromisoformat(pm_b[pm_b > fail_day.isoformat()].min())
    interval = TYPES["CNC Lathe"]["pm2"]
    due = last_pm + timedelta(days=interval)
    f["failure_date"] = long_date(fail_day)
    f["failure_hours"] = fmt1(float(failure["downtime_hours"]))
    f["last_pm2"] = long_date(last_pm)
    f["pm2_due"] = long_date(due)
    f["days_overdue_at_failure"] = str((fail_day - due).days)
    f["pm2_done_after"] = long_date(after_pm)
    assert (fail_day - due).days > 0
    sb = sensor[sensor["machine_id"] == BORE_MACHINE].copy()
    sb["d"] = pd.to_datetime(sb["reading_date"])
    base = sb[(sb["reading_date"] >= "2026-06-01") & (sb["reading_date"] <= "2026-07-31")]
    f["vib_base"] = f"{base['avg_vibration_mm_s'].mean():.1f}"
    f["temp_base"] = f"{base['max_temperature_c'].mean():.0f}"
    pk = sb[(sb["reading_date"] >= "2026-08-17") & (sb["reading_date"] <= "2026-08-21")]
    f["vib_week_peak"] = f"{pk['avg_vibration_mm_s'].mean():.1f}"
    f["temp_peak"] = f"{pk['max_temperature_c'].mean():.0f}"
    win = sb[(sb["reading_date"] >= "2026-08-03") & (sb["reading_date"] <= "2026-08-21")]
    f["alarms_ramp"] = str(int(win["alarm_count"].sum()))
    lw, la = THRESHOLDS["CNC Lathe"][0], THRESHOLDS["CNC Lathe"][1]
    f["lathe_warn"], f["lathe_alarm"] = f"{lw:.1f}", f"{la:.1f}"
    seq = sb[sb["reading_date"] <= fail_day.isoformat()].reset_index(drop=True)
    over = (seq["avg_vibration_mm_s"] > lw).to_numpy()
    first = next(
        i
        for i in range(len(over) - 2)
        if over[i] and over[i + 1] and over[i + 2] and seq["reading_date"][i] >= "2026-07-15"
    )
    warn_day = date.fromisoformat(seq["reading_date"][first])
    alarm_day = date.fromisoformat(
        seq[(seq["avg_vibration_mm_s"] > la) & (seq["reading_date"] >= "2026-07-15")]["reading_date"].iloc[0]
    )
    f["warn_date"], f["alarm_date"] = long_date(warn_day), long_date(alarm_day)
    f["warn_days_before"] = str((fail_day - warn_day).days)
    post = sb[(sb["reading_date"] >= "2026-09-01") & (sb["reading_date"] <= "2026-09-30")]
    f["vib_post"] = f"{post['avg_vibration_mm_s'].mean():.1f}"
    post_def = df[
        (df["machine_id"] == BORE_MACHINE) & (df["defect_type"] == "dimension_oob") & (df["detected_at"] > "2026-08-24")
    ]
    assert len(post_def) == 0
    f["post_defects_phrase"] = "No dimension_oob records were logged on MTY-LTH-04 between 25 August and 30 September."
    f["close_date"] = "18 September 2026"
    # Weekly vibration for the figure.
    s2 = sensor.merge(mp, on="machine_id")
    s2["week"] = pd.to_datetime(s2["reading_date"]).dt.to_period("W-SUN").dt.start_time
    weeks = list(pd.date_range("2026-06-01", "2026-09-21", freq="7D"))
    mine = s2[s2["machine_id"] == BORE_MACHINE].groupby("week")["avg_vibration_mm_s"].mean()
    others = s2[(s2["machine_type"] == "CNC Lathe") & (s2["machine_id"] != BORE_MACHINE)]
    others = others.groupby(["machine_id", "week"])["avg_vibration_mm_s"].mean().groupby("week").median()
    charts["vibration"] = (
        [f"{w.day} {w.strftime('%b')}" for w in weeks],
        [
            (BORE_MACHINE, [float(mine[w]) for w in weeks]),
            ("Other CNC lathes (median)", [float(others[w]) for w in weeks]),
        ],
    )
    f["fleet_med"] = f"{float(others[[w for w in weeks if w < pd.Timestamp('2026-08-01')]].mean()):.1f}"
    spec = (
        "Main bore of valve body AP-3312 (proportional valve body PV-12) oversize; "
        "bore specification 22.000 mm, +0.021 / 0 mm."
    )
    tables["eight_d_problem"] = (
        ["Question", "Answer"],
        [
            ["What", spec],
            [
                "Where",
                f"Monterrey Plant, CNC lathe {BORE_MACHINE} ({f['bore_cell']}). "
                "Found at in-process gauging and by the customer.",
            ],
            [
                "When",
                f"Customer return detected 14 August 2026; in-plant records from {f['cluster_first']} "
                f"to {f['cluster_last']} 2026.",
            ],
            ["Who", "Monterrey operators and quality inspectors; Brightline Hydraulics incoming inspection."],
            [
                "How many",
                f"{f['escape_qty']} pieces returned by the customer; {f['inplant_qty']} pieces in "
                f"{f['n_inplant']} in-plant records.",
            ],
            ["How detected", "Customer incoming inspection; in-process air gauge on one piece in 25."],
            ["Cost", f"${f['total_cost']} of scrap and rework, including the customer return."],
        ],
    )

    # Tables quoted by the manuals ---------------------------------------------------------------------------------
    tables["pm_intervals"] = (
        [
            "Machine type",
            "PM-1 operator check",
            "PM-2 interval (calendar days)",
            "Fluid and lubricant analysis",
            "Geometry and calibration",
        ],
        [[mt, PM_TEXT[mt][0], str(TYPES[mt]["pm2"]), PM_TEXT[mt][1], PM_TEXT[mt][2]] for mt in TYPE_ORDER],
    )
    tables["condition_thresholds"] = (
        [
            "Machine type",
            "Vibration warning (mm/s)",
            "Vibration alarm (mm/s)",
            "Max temperature warning (C)",
            "Max temperature alarm (C)",
        ],
        [[mt, f"{v[0]:.1f}", f"{v[1]:.1f}", str(v[2]), str(v[3])] for mt, v in THRESHOLDS.items()],
    )
    rows = []
    for code, (desc, _hrs, _w, _fix) in FAILURE_CODES.items():
        types = [mt for mt in TYPE_ORDER if code in TYPES[mt]["codes"]]
        rows.append([code, desc, ", ".join(types)])
    tables["failure_codes"] = (["Code", "Failure", "Machine types"], rows)
    tables["loto_sources"] = (
        ["Machine type", "Energy sources", "Isolating devices and checks"],
        [[mt, *LOTO_SOURCES[mt]] for mt in TYPE_ORDER],
    )
    tables["sqa_targets"] = (["Measure", "Target", "Basis"], [list(r) for r in SQA_TARGETS])
    return f, tables, {"charts": charts, "raw": raw}


def compute_scan_facts(t: dict) -> dict:
    """What the scanned round sheet says, taken from the tables so that paper and data agree."""
    s = t["sensor"]
    row = s[(s["machine_id"] == SCAN_MACHINE) & (s["reading_date"] == SCAN_DATE)].iloc[0]
    ev = t["events"]
    leak = ev[
        (ev["machine_id"] == SCAN_MACHINE) & (ev["failure_code"] == "HYD-LEAK") & (ev["started_at"] >= SCAN_DATE)
    ].iloc[0]
    return {
        "temp": f"{row['max_temperature_c']:.0f}",
        "leak_date": str(leak["started_at"])[:10],
        "leak_event": str(leak["event_id"]),
    }


# Writers: tables --------------------------------------------------------------------------------------------------


def write_tables(t: dict) -> dict[str, int]:
    out = FILES / "tables"
    out.mkdir(parents=True, exist_ok=True)
    counts = {}
    for name, key in [
        ("plants", "plants"),
        ("machines", "machines_df"),
        ("sensor_readings", "sensor"),
        ("maintenance_events", "events"),
        ("work_orders", "work_orders"),
        ("quality_defects", "defects"),
    ]:
        df = t[key]
        df.to_csv(out / f"{name}.csv", index=False, lineterminator="\n")
        counts[name] = len(df)
    write_plan_xlsx(out / "production_plan.xlsx", t["plan"], t["targets"])
    counts["production_plan_weekly_output"] = len(t["plan"])
    counts["production_plan_scrap_targets"] = len(t["targets"])
    return counts


def write_plan_xlsx(path: Path, plan: pd.DataFrame, targets: pd.DataFrame) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment
    from openpyxl.styles import Font
    from openpyxl.styles import PatternFill

    wb = Workbook()
    wb.properties.creator = COMPANY
    wb.properties.lastModifiedBy = COMPANY
    wb.properties.title = "Production plan"
    wb.properties.created = datetime(2026, 9, 30, 12, 0, 0)
    wb.properties.modified = datetime(2026, 9, 30, 12, 0, 0)
    head_fill = PatternFill("solid", fgColor="1F3A5F")
    for idx, (title, df, widths) in enumerate(
        [
            ("weekly_output", plan, [18, 12, 10, 14, 14, 14]),
            ("scrap_targets", targets, [16, 10, 10, 18, 32]),
        ]
    ):
        ws = wb.active if idx == 0 else wb.create_sheet()
        ws.title = title
        ws.append(list(df.columns))
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = head_fill
            cell.alignment = Alignment(horizontal="center")
        for row in df.itertuples(index=False):
            ws.append([date.fromisoformat(v) if c == "week_start" else v for c, v in zip(df.columns, row, strict=True)])
        for col_idx, width in enumerate(widths, start=1):
            ws.column_dimensions[chr(64 + col_idx)].width = width
        if "week_start" in df.columns:
            col = list(df.columns).index("week_start") + 1
            for r in range(2, len(df) + 2):
                ws.cell(row=r, column=col).number_format = "yyyy-mm-dd"
        ws.freeze_panes = "A2"
    wb.save(path)
    normalize_zip(path)


ZIP_STAMP = (2026, 9, 30, 12, 0, 0)
CORE_STAMP = b"2026-09-30T12:00:00Z"


def normalize_zip_bytes(raw: bytes) -> bytes:
    """Rewrite an OOXML zip with fixed entry timestamps and core-property dates, so that it is byte-identical from run
    to run. Embedded workbooks (the data behind a chart) are rewritten too."""
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(raw)) as zin, zipfile.ZipFile(out, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename.endswith(".xlsx"):
                data = normalize_zip_bytes(data)
            if item.filename == "docProps/core.xml":
                data = re.sub(
                    rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*(</dcterms:(?:created|modified)>)",
                    rb"\g<1>" + CORE_STAMP + rb"\g<2>",
                    data,
                )
            info = zipfile.ZipInfo(item.filename, date_time=ZIP_STAMP)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o644 << 16
            zout.writestr(info, data)
    return out.getvalue()


def normalize_zip(path: Path) -> None:
    path.write_bytes(normalize_zip_bytes(path.read_bytes()))


# Writers: documents -----------------------------------------------------------------------------------------------

INK = "#222B33"
NAVY = "#1F3A5F"
AMBER = "#D9822B"
STEEL = "#5B6B7B"
PALE = "#EEF2F6"


def parse_md(text: str, facts: dict | None, tables: dict) -> tuple[dict, list]:
    """A small markdown dialect: % meta, ## / ### headings, paragraphs, - and 1. lists, | tables, > notes,
    *captions*, !table:name and !chart:name tokens, and `subtitle:` / `notes:` lines for slides."""
    lines = (text if facts is None else text.format(**facts)).splitlines()
    meta: dict = {}
    blocks: list = []
    para: list[str] = []

    def flush() -> None:
        if para:
            blocks.append(("p", " ".join(para)))
            para.clear()

    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if line.startswith("<!--"):
            pass  # the SPDX header
        elif line.startswith("% "):
            key, _, val = line[2:].partition(":")
            meta[key.strip()] = val.strip()
        elif not line.strip():
            flush()
        elif line.startswith("### "):
            flush()
            blocks.append(("h3", line[4:]))
        elif line.startswith("## "):
            flush()
            blocks.append(("h2", line[3:]))
        elif line.startswith("# "):
            flush()
            blocks.append(("h1", line[2:]))
        elif line.startswith("> "):
            flush()
            blocks.append(("note", line[2:]))
        elif line.startswith("- "):
            flush()
            items = []
            while i < len(lines) and lines[i].startswith("- "):
                items.append(lines[i][2:].strip())
                i += 1
            blocks.append(("ul", items))
            continue
        elif re.match(r"\d+\. ", line):
            flush()
            items = []
            while i < len(lines) and re.match(r"\d+\. ", lines[i]):
                items.append(re.sub(r"^\d+\. ", "", lines[i]).strip())
                i += 1
            blocks.append(("ol", items))
            continue
        elif line.startswith("|"):
            flush()
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(set(c) <= set("-: ") for c in cells):
                    rows.append(cells)
                i += 1
            blocks.append(("table", (rows[0], rows[1:])))
            continue
        elif line.startswith("!table:"):
            flush()
            blocks.append(("table", tables[line[7:].strip()]))
        elif line.startswith("!chart:"):
            flush()
            blocks.append(("chart", line[7:].strip()))
        elif re.fullmatch(r"\*[^*].*[^*]\*", line):
            flush()
            blocks.append(("caption", line[1:-1]))
        elif re.match(r"(subtitle|notes):", line):
            flush()
            key, _, val = line.partition(":")
            blocks.append(("kv", (key, val.strip())))
        else:
            para.append(line.strip())
        i += 1
    flush()
    return meta, blocks


def rl_inline(text: str) -> str:
    from xml.sax.saxutils import escape

    out = escape(text)
    out = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", out)
    return re.sub(r"(?<![\w*])\*([^*]+?)\*(?![\w*])", r"<i>\1</i>", out)


def vibration_drawing(labels: list[str], series: list[tuple[str, list[float]]]):
    from reportlab.graphics.charts.legends import Legend
    from reportlab.graphics.charts.linecharts import HorizontalLineChart
    from reportlab.graphics.shapes import Drawing
    from reportlab.graphics.shapes import String
    from reportlab.lib import colors
    from reportlab.lib.units import inch

    width, height = 6.9 * inch, 3.0 * inch
    d = Drawing(width, height)
    lc = HorizontalLineChart()
    lc.x, lc.y, lc.width, lc.height = 36, 62, width - 52, height - 90
    warn = THRESHOLDS["CNC Lathe"][0]
    lc.data = [s[1] for s in series] + [[warn] * len(labels)]
    lc.categoryAxis.categoryNames = labels
    lc.categoryAxis.labels.fontName = "Helvetica"
    lc.categoryAxis.labels.fontSize = 7
    lc.categoryAxis.labels.angle = 40
    lc.categoryAxis.labels.boxAnchor = "ne"
    lc.categoryAxis.labels.dx = 2
    lc.categoryAxis.labels.dy = -3
    lc.valueAxis.labels.fontName = "Helvetica"
    lc.valueAxis.labels.fontSize = 8
    lc.valueAxis.valueMin = 0
    lc.valueAxis.valueMax = 6
    lc.valueAxis.valueStep = 1
    lc.valueAxis.gridStrokeColor = colors.HexColor("#D5DCE3")
    lc.valueAxis.visibleGrid = 1
    lc.joinedLines = 1
    styles = [(AMBER, 2.4, None), (STEEL, 1.6, None), ("#B03030", 1.2, [4, 3])]
    for i, (color, w, dash) in enumerate(styles):
        lc.lines[i].strokeColor = colors.HexColor(color)
        lc.lines[i].strokeWidth = w
        if dash:
            lc.lines[i].strokeDashArray = dash
        lc.lines[i].symbol = None
    d.add(lc)
    d.add(String(36, height - 14, "Weekly average vibration (mm/s RMS)", fontName="Helvetica-Bold", fontSize=9))
    legend = Legend()
    legend.x, legend.y = 70, 14
    legend.alignment = "right"
    legend.fontName = "Helvetica"
    legend.fontSize = 8
    legend.columnMaximum = 1
    legend.deltax = 120
    legend.colorNamePairs = [
        (colors.HexColor(c[0]), n)
        for c, n in zip(styles, [s[0] for s in series] + [f"CNC lathe warning level ({warn} mm/s)"], strict=True)
    ]
    legend.columnMaximum = 1
    legend.dxTextSpace = 4
    d.add(legend)
    return d


def build_pdf(md_name: str, out_path: Path, facts: dict, tables: dict, extra: dict) -> None:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import KeepTogether
    from reportlab.platypus import ListFlowable
    from reportlab.platypus import ListItem
    from reportlab.platypus import Paragraph
    from reportlab.platypus import SimpleDocTemplate
    from reportlab.platypus import Spacer
    from reportlab.platypus import Table
    from reportlab.platypus import TableStyle

    meta, blocks = parse_md((CONTENT / md_name).read_text(encoding="utf-8"), facts, tables)
    ink, navy, amber = colors.HexColor(INK), colors.HexColor(NAVY), colors.HexColor(AMBER)
    body = ParagraphStyle("body", fontName="Helvetica", fontSize=10, leading=14, spaceAfter=6, textColor=ink)
    h2 = ParagraphStyle(
        "h2",
        parent=body,
        fontName="Helvetica-Bold",
        fontSize=14,
        leading=18,
        spaceBefore=12,
        textColor=navy,
        keepWithNext=1,
    )
    h3 = ParagraphStyle(
        "h3",
        parent=body,
        fontName="Helvetica-Bold",
        fontSize=11.5,
        leading=15,
        spaceBefore=8,
        textColor=navy,
        keepWithNext=1,
    )
    title = ParagraphStyle(
        "title", parent=body, fontName="Helvetica-Bold", fontSize=22, leading=26, textColor=navy, spaceAfter=4
    )
    subtitle = ParagraphStyle(
        "subtitle", parent=body, fontSize=12, leading=16, textColor=colors.HexColor(STEEL), spaceAfter=10
    )
    cell = ParagraphStyle("cell", parent=body, fontSize=8.5, leading=11, spaceAfter=0)
    cell_head = ParagraphStyle("cell_head", parent=cell, fontName="Helvetica-Bold", textColor=colors.white)
    caption = ParagraphStyle(
        "caption",
        parent=body,
        fontName="Helvetica-Oblique",
        fontSize=8.5,
        leading=11,
        textColor=colors.HexColor(STEEL),
        spaceBefore=3,
        spaceAfter=10,
    )
    note = ParagraphStyle("note", parent=body, fontSize=9, leading=12.5, spaceAfter=0)

    page_w = letter[0] - 1.5 * inch

    def make_table(header: list[str], rows: list[list[str]]) -> Table:
        weights = []
        for c in range(len(header)):
            longest = max([len(header[c]), *[len(r[c]) for r in rows]])
            weights.append(min(max(longest, 7), 38) + 3)
        total = sum(weights)
        widths = [page_w * w / total for w in weights]
        data = [[Paragraph(rl_inline(h), cell_head) for h in header]] + [
            [Paragraph(rl_inline(c), cell) for c in r] for r in rows
        ]
        tb = Table(data, colWidths=widths, repeatRows=1, hAlign="LEFT")
        style = [
            ("BACKGROUND", (0, 0), (-1, 0), navy),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#B8C2CC")),
            ("LEFTPADDING", (0, 0), (-1, -1), 5),
            ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ("TOPPADDING", (0, 0), (-1, -1), 3.5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3.5),
        ]
        for r in range(2, len(data), 2):
            style.append(("BACKGROUND", (0, r), (-1, r), colors.HexColor(PALE)))
        tb.setStyle(TableStyle(style))
        return tb

    story: list = []
    story.append(Paragraph(rl_inline(meta["title"]), title))
    if meta.get("subtitle"):
        story.append(Paragraph(rl_inline(meta["subtitle"]), subtitle))
    labels = [
        ("doc_id", "Document ID"),
        ("revision", "Revision"),
        ("effective", "Effective"),
        ("owner", "Owner"),
        ("applies_to", "Applies to"),
    ]
    info = [[Paragraph(f"<b>{lab}</b>", cell), Paragraph(rl_inline(meta[k]), cell)] for k, lab in labels if k in meta]
    info_t = Table(info, colWidths=[1.2 * inch, page_w - 1.2 * inch], hAlign="LEFT")
    info_t.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (0, -1), colors.HexColor(PALE)),
                ("LINEABOVE", (0, 0), (-1, 0), 1.5, amber),
                ("LINEBELOW", (0, -1), (-1, -1), 0.4, colors.HexColor("#B8C2CC")),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]
        )
    )
    story += [info_t, Spacer(1, 6)]

    pending: list = []

    def emit(flowable, keep_next: bool = False) -> None:
        if keep_next:
            pending.append(flowable)
        elif pending:
            story.append(KeepTogether([*pending, flowable]))
            pending.clear()
        else:
            story.append(flowable)

    for idx, (kind, data) in enumerate(blocks):
        following = blocks[idx + 1][0] if idx + 1 < len(blocks) else ""
        if kind == "h2":
            emit(Paragraph(rl_inline(data), h2), keep_next=True)
        elif kind == "h3":
            emit(Paragraph(rl_inline(data), h3), keep_next=True)
        elif kind == "p":
            emit(Paragraph(rl_inline(data), body), keep_next=following == "table")
        elif kind == "note":
            box = Table([[Paragraph(rl_inline(data), note)]], colWidths=[page_w], hAlign="LEFT")
            box.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#FFF5E6")),
                        ("LINEBEFORE", (0, 0), (0, -1), 3, amber),
                        ("TOPPADDING", (0, 0), (-1, -1), 6),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                    ]
                )
            )
            emit(Spacer(1, 2))
            emit(box)
            emit(Spacer(1, 8))
        elif kind in ("ul", "ol"):
            items = [ListItem(Paragraph(rl_inline(x), body), leftIndent=16) for x in data]
            kwargs = (
                {"bulletType": "bullet", "start": "\u2022"}
                if kind == "ul"
                else {"bulletType": "1", "bulletFormat": "%s."}
            )
            emit(ListFlowable(items, bulletFontName="Helvetica", bulletFontSize=9, leftIndent=16, **kwargs))
            emit(Spacer(1, 4))
        elif kind == "table":
            header, rows = data
            emit(make_table(header, rows), keep_next=following == "caption")
            if following != "caption":
                emit(Spacer(1, 8))
        elif kind == "caption":
            emit(Paragraph(rl_inline(data), caption))
        elif kind == "chart":
            labels_c, series_c = extra["charts"][data]
            emit(vibration_drawing(labels_c, series_c), keep_next=following == "caption")
    if pending:
        story.extend(pending)

    def on_page(canvas, doc) -> None:
        canvas.saveState()
        canvas.setFont("Helvetica-Bold", 8)
        canvas.setFillColor(navy)
        canvas.drawString(0.75 * inch, letter[1] - 0.5 * inch, COMPANY.upper())
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.HexColor(STEEL))
        canvas.drawRightString(
            letter[0] - 0.75 * inch, letter[1] - 0.5 * inch, f"{meta.get('doc_id', '')}  |  {meta.get('revision', '')}"
        )
        canvas.setStrokeColor(amber)
        canvas.setLineWidth(1.2)
        canvas.line(0.75 * inch, letter[1] - 0.58 * inch, letter[0] - 0.75 * inch, letter[1] - 0.58 * inch)
        canvas.setStrokeColor(colors.HexColor("#B8C2CC"))
        canvas.setLineWidth(0.4)
        canvas.line(0.75 * inch, 0.62 * inch, letter[0] - 0.75 * inch, 0.62 * inch)
        canvas.drawString(
            0.75 * inch, 0.45 * inch, "Synthetic document for a software demonstration. Not for operational use."
        )
        canvas.drawRightString(letter[0] - 0.75 * inch, 0.45 * inch, f"Page {doc.page}")
        canvas.restoreState()

    doc = SimpleDocTemplate(
        str(out_path),
        pagesize=letter,
        leftMargin=0.75 * inch,
        rightMargin=0.75 * inch,
        topMargin=0.85 * inch,
        bottomMargin=0.85 * inch,
        title=meta["title"],
        author=COMPANY,
        subject=meta.get("subtitle", ""),
        creator=f"{COMPANY} (synthetic)",
        invariant=1,
    )
    doc.build(story, onFirstPage=on_page, onLaterPages=on_page)


def docx_inline(paragraph, text: str, size: float | None = None, bold: bool = False) -> None:
    from docx.shared import Pt

    for k, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if not part:
            continue
        run = paragraph.add_run(part)
        run.bold = bold or k % 2 == 1
        if size:
            run.font.size = Pt(size)


def build_docx(md_name: str, out_path: Path, facts: dict, tables: dict) -> None:
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches
    from docx.shared import Pt
    from docx.shared import RGBColor

    meta, blocks = parse_md((CONTENT / md_name).read_text(encoding="utf-8"), facts, tables)
    doc = Document()
    stamp = datetime(2026, 9, 30, 12, 0, 0)
    cp = doc.core_properties
    cp.author = COMPANY
    cp.last_modified_by = COMPANY
    cp.title = meta["title"]
    cp.subject = meta.get("subtitle", "")
    cp.comments = "Synthetic document for a software demonstration."
    cp.created = stamp
    cp.modified = stamp
    cp.revision = 1
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Inches(1)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    for name, size in (("Heading 1", 15), ("Heading 2", 12.5), ("Title", 24)):
        st = doc.styles[name]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor.from_string(NAVY[1:])
    sec.header.paragraphs[0].text = f"{COMPANY}  |  {meta.get('doc_id', '')}  |  {meta.get('revision', '')}"
    sec.footer.paragraphs[0].text = "Synthetic document for a software demonstration. Not for operational use."

    def shade(c, fill: str) -> None:
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fill)
        c._tc.get_or_add_tcPr().append(shd)

    def add_table(header: list[str], rows: list[list[str]]) -> None:
        tb = doc.add_table(rows=1 + len(rows), cols=len(header))
        tb.style = "Table Grid"
        for c, h in enumerate(header):
            cl = tb.rows[0].cells[c]
            cl.text = ""
            docx_inline(cl.paragraphs[0], h, size=9.5, bold=True)
            cl.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
            shade(cl, NAVY[1:])
        for r, row in enumerate(rows, start=1):
            for c, val in enumerate(row):
                cl = tb.rows[r].cells[c]
                cl.text = ""
                docx_inline(cl.paragraphs[0], val, size=9.5)
                if r % 2 == 0:
                    shade(cl, PALE[1:])
        doc.add_paragraph()

    doc.add_paragraph(meta["title"], style="Title")
    if meta.get("subtitle"):
        sp = doc.add_paragraph()
        docx_inline(sp, meta["subtitle"], size=12)
    keys = [
        ("doc_id", "Agreement ID"),
        ("revision", "Version"),
        ("effective", "Effective"),
        ("owner", "Owner"),
        ("applies_to", "Applies to"),
    ]
    add_table(["Field", "Value"], [[lab, meta[k]] for k, lab in keys if k in meta])
    for kind, data in blocks:
        if kind == "h2":
            doc.add_heading(data, level=1)
        elif kind == "h3":
            doc.add_heading(data, level=2)
        elif kind in ("p", "note"):
            p = doc.add_paragraph()
            docx_inline(p, data)
            if kind == "note":
                for r in p.runs:
                    r.italic = True
        elif kind == "ul":
            for item in data:
                docx_inline(doc.add_paragraph(style="List Bullet"), item)
        elif kind == "ol":
            for n, item in enumerate(data, start=1):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.3)
                p.paragraph_format.first_line_indent = Inches(-0.25)
                docx_inline(p, f"{n}.  {item}")
        elif kind == "table":
            add_table(*data)
        elif kind == "caption":
            p = doc.add_paragraph()
            docx_inline(p, data, size=9)
            for r in p.runs:
                r.italic = True
    doc.save(out_path)
    normalize_zip(out_path)


def build_pptx(md_name: str, out_path: Path, facts: dict, tables: dict, extra: dict) -> None:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.enum.chart import XL_LEGEND_POSITION
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR
    from pptx.enum.text import PP_ALIGN
    from pptx.oxml.ns import qn
    from pptx.util import Emu
    from pptx.util import Inches
    from pptx.util import Pt

    raw = (CONTENT / md_name).read_text(encoding="utf-8").format(**facts)
    head, _, rest = raw.partition("\n# ")
    meta, _ = parse_md(head, None, tables)
    parts = re.split(r"\n---\n", "# " + rest)

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    stamp = datetime(2026, 9, 30, 12, 0, 0)
    cp = prs.core_properties
    cp.author = COMPANY
    cp.last_modified_by = COMPANY
    cp.title = meta.get("title", "")
    cp.subject = meta.get("subtitle", "")
    cp.comments = "Synthetic document for a software demonstration."
    cp.created = stamp
    cp.modified = stamp
    cp.revision = 1
    navy, amber, ink = RGBColor.from_string(NAVY[1:]), RGBColor.from_string(AMBER[1:]), RGBColor.from_string(INK[1:])
    white = RGBColor(255, 255, 255)
    series_colors = [RGBColor.from_string(c) for c in ("1F3A5F", "D9822B", "7A8DA3")]

    def rect(slide, x, y, w, h, color) -> None:  # noqa: PLR0917
        shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
        shp.fill.solid()
        shp.fill.fore_color.rgb = color
        shp.line.fill.background()
        slide.shapes._spTree.remove(shp._element)
        slide.shapes._spTree.insert(2, shp._element)

    def bullets(slide, items, x, y, w, h, size=18, color=ink) -> None:  # noqa: PLR0917
        tb = slide.shapes.add_textbox(x, y, w, h)
        tf = tb.text_frame
        tf.word_wrap = True
        for n, item in enumerate(items):
            p = tf.paragraphs[0] if n == 0 else tf.add_paragraph()
            p.space_after = Pt(8)
            ppr = p._p.get_or_add_pPr()
            ppr.set("marL", "285750")
            ppr.set("indent", "-285750")
            bu = ppr.makeelement(qn("a:buChar"), {"char": "•"})
            ppr.append(bu)
            for k, part in enumerate(re.split(r"\*\*(.+?)\*\*", item)):
                if part:
                    r = p.add_run()
                    r.text = part
                    r.font.size = Pt(size)
                    r.font.bold = k % 2 == 1
                    r.font.color.rgb = color

    def set_title(slide, text, size, color, x, y, w, h) -> None:  # noqa: PLR0917
        ph = slide.shapes.title
        ph.left, ph.top, ph.width, ph.height = x, y, w, h
        tf = ph.text_frame
        tf.text = text
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        for r in p.runs:
            r.font.size = Pt(size)
            r.font.bold = True
            r.font.color.rgb = color

    layout = prs.slide_layouts[5]
    for number, part in enumerate(parts, start=1):
        _, blocks = parse_md(part, None, tables)
        title = next(d for k, d in blocks if k == "h1")
        kv = {d[0]: d[1] for k, d in blocks if k == "kv"}
        items = [x for k, d in blocks if k == "ul" for x in d]
        tbl = next((d for k, d in blocks if k == "table"), None)
        chart = next((d for k, d in blocks if k == "chart"), None)
        slide = prs.slides.add_slide(layout)
        if number == 1:
            rect(slide, 0, 0, prs.slide_width, prs.slide_height, navy)
            rect(slide, Inches(0.8), Inches(3.55), Inches(1.6), Inches(0.08), amber)
            set_title(slide, title, 40, white, Inches(0.8), Inches(1.7), Inches(11.5), Inches(1.7))
            sub = slide.shapes.add_textbox(Inches(0.8), Inches(3.8), Inches(11.5), Inches(0.6))
            sub.text_frame.word_wrap = True
            sub.text_frame.text = kv.get("subtitle", "")
            sub.text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT
            sub.text_frame.paragraphs[0].runs[0].font.size = Pt(22)
            sub.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xD5, 0xDC, 0xE3)
            bullets(
                slide,
                items,
                Inches(0.8),
                Inches(4.7),
                Inches(11.5),
                Inches(2.0),
                size=16,
                color=RGBColor(0xD5, 0xDC, 0xE3),
            )
        else:
            rect(slide, 0, 0, prs.slide_width, Inches(1.05), navy)
            rect(slide, 0, Inches(1.05), prs.slide_width, Inches(0.06), amber)
            set_title(slide, title, 30, white, Inches(0.6), Inches(0.1), Inches(12.1), Inches(0.85))
            top = Inches(1.45)
            if chart:
                labels, series = extra["charts"][chart]
                cd = CategoryChartData()
                cd.categories = labels
                for name, vals in series:
                    cd.add_series(name, [round(v, 1) for v in vals])
                gf = slide.shapes.add_chart(
                    XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.5), top, Inches(7.7), Inches(5.2), cd
                )
                ch = gf.chart
                ch.has_legend = True
                ch.legend.position = XL_LEGEND_POSITION.BOTTOM
                ch.legend.include_in_layout = False
                ch.legend.font.size = Pt(13)
                ch.category_axis.tick_labels.font.size = Pt(13)
                ch.value_axis.tick_labels.font.size = Pt(12)
                ch.value_axis.has_title = True
                ch.value_axis.axis_title.text_frame.text = "Unplanned downtime (hours)"
                ch.value_axis.axis_title.text_frame.paragraphs[0].runs[0].font.size = Pt(12)
                plot = ch.plots[0]
                plot.has_data_labels = True
                plot.data_labels.font.size = Pt(11)
                plot.data_labels.number_format = "0.0"
                plot.data_labels.number_format_is_linked = False
                for s, col in zip(plot.series, series_colors, strict=False):
                    s.format.fill.solid()
                    s.format.fill.fore_color.rgb = col
                bullets(slide, items, Inches(8.5), top + Inches(0.2), Inches(4.4), Inches(5.0), size=17)
            else:
                body_top = top
                if tbl:
                    header, rows = tbl
                    n_rows, n_cols = len(rows) + 1, len(header)
                    row_h = Inches(0.52)
                    gt = slide.shapes.add_table(n_rows, n_cols, Inches(0.6), top, Inches(12.1), row_h * n_rows)
                    t_ = gt.table
                    t_.first_row = True
                    weights = [max(len(header[c]), *[len(r[c]) for r in rows]) + 4 for c in range(n_cols)]
                    for c in range(n_cols):
                        t_.columns[c].width = Emu(int(Inches(12.1) * weights[c] / sum(weights)))
                    for r in range(n_rows):
                        t_.rows[r].height = row_h
                        for c in range(n_cols):
                            cell = t_.cell(r, c)
                            cell.text = header[c] if r == 0 else rows[r - 1][c]
                            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                            run = cell.text_frame.paragraphs[0].runs[0]
                            run.font.size = Pt(15)
                            run.font.bold = r == 0 or r == n_rows - 1 and rows[-1][0] == "All plants"
                            cell.fill.solid()
                            cell.fill.fore_color.rgb = (
                                navy if r == 0 else (RGBColor(0xEE, 0xF2, 0xF6) if r % 2 == 0 else white)
                            )
                            run.font.color.rgb = white if r == 0 else ink
                    body_top = top + row_h * n_rows + Inches(0.35)
                bullets(
                    slide,
                    items,
                    Inches(0.6),
                    body_top,
                    Inches(12.1),
                    Inches(7.0) - body_top,
                    size=18 if not tbl else 16,
                )
            foot = slide.shapes.add_textbox(Inches(0.6), Inches(7.05), Inches(12.1), Inches(0.3))
            foot.text_frame.word_wrap = True
            foot.text_frame.text = f"{COMPANY}  |  Q3 2026 Plant Performance Review  |  Synthetic data  |  {number}"
            foot.text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT
            fr = foot.text_frame.paragraphs[0].runs[0]
            fr.font.size = Pt(10)
            fr.font.color.rgb = RGBColor.from_string(STEEL[1:])
        if kv.get("notes"):
            slide.notes_slide.notes_text_frame.text = kv["notes"]
    prs.save(out_path)
    normalize_zip(out_path)


# Writers: the scanned round sheet -----------------------------------------------------------------------------------

SANS = [
    "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
]
SANS_BOLD = [
    "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf",
    "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
]
HAND = ["/usr/share/fonts/truetype/tlwg/Purisa.ttf", "/usr/share/fonts/opentype/urw-base35/Z003-MediumItalic.otf"]


def load_font(paths: list[str], size: int):
    from PIL import ImageFont

    for p in paths:
        if Path(p).exists():
            try:
                return ImageFont.truetype(p, size)
            except OSError:
                continue
    return ImageFont.load_default(size)


def render_scan(out_path: Path, sf: dict) -> None:
    """A photographed or scanned handwritten operator round sheet: grayscale, slightly rotated, noisy."""
    import textwrap

    from PIL import Image
    from PIL import ImageDraw
    from PIL import ImageFilter

    rng = np.random.default_rng(SEED + 77)
    w, h = 1275, 1650
    page = Image.new("L", (w, h), 250)
    d = ImageDraw.Draw(page)
    sans, sans_b = (lambda s: load_font(SANS, s)), (lambda s: load_font(SANS_BOLD, s))

    def hand(xy, text, size=30, ink=None, drift=1.2):
        font = load_font(HAND, size)
        x, y = xy
        for ch in text:
            if ch == " ":
                x += font.getlength(" ") * 0.85
                continue
            bbox = font.getbbox(ch)
            gw, gh = bbox[2] + 24, bbox[3] + 24
            tile = Image.new("L", (gw, gh), 0)
            ImageDraw.Draw(tile).text((12, 12), ch, font=font, fill=255)
            tile = tile.rotate(float(rng.normal(0, 3.2)), resample=Image.BICUBIC, expand=True)
            tone = int(ink if ink is not None else rng.uniform(22, 62))
            dy = float(rng.normal(0, drift)) + 2.0 * np.sin(x / 170.0)
            page.paste(tone, (int(x - 12), int(y - 12 + dy)), tile)
            x += font.getlength(ch) * float(rng.uniform(0.9, 1.0))
        return x

    def tick(cx, cy):
        pts = [(cx - 16, cy), (cx - 5, cy + 15), (cx + 20, cy - 20)]
        pts = [(px + float(rng.normal(0, 1.2)), py + float(rng.normal(0, 1.2))) for px, py in pts]
        d.line(pts, fill=45, width=5, joint="curve")

    def cross(cx, cy):
        for (a, b), (c, e) in [((-15, -15), (15, 16)), ((15, -15), (-14, 15))]:
            d.line(
                [(cx + a + float(rng.normal(0, 1)), cy + b), (cx + c, cy + e + float(rng.normal(0, 1)))],
                fill=40,
                width=5,
            )

    # Printed form.
    d.rectangle([60, 55, w - 60, 160], outline=30, width=3)
    d.text((82, 68), "ATLAS PRECISION COMPONENTS", font=sans_b(36), fill=20)
    d.text((82, 118), "PM-1 OPERATOR ROUND SHEET  -  HYDRAULIC PRESS", font=sans_b(24), fill=35)
    d.text((w - 345, 72), "Form MNT-F-012  Rev C", font=sans(21), fill=40)
    d.text((w - 345, 108), "Keep on machine; file 12 months", font=sans(17), fill=60)

    fields = [
        ("Plant", 60, 185, 560, "Dayton, OH"),
        ("Machine ID", 640, 185, w - 60, SCAN_MACHINE),
        ("Date", 60, 255, 560, "14 Sep 2026"),
        ("Shift / Operator", 640, 255, w - 60, "2   T. Brennan"),
    ]
    for label, x0, y0, x1, value in fields:
        d.rectangle([x0, y0, x1, y0 + 62], outline=50, width=2)
        d.text((x0 + 10, y0 + 4), label, font=sans(16), fill=70)
        hand((x0 + 190 if len(label) < 8 else x0 + 215, y0 + 12), value, size=31)

    cols = [60, 112, 640, 732, 840, w - 60]  # number, item, OK, NOT OK, remarks
    top, head_h, row_h = 350, 44, 100
    d.rectangle([cols[0], top, cols[-1], top + head_h], fill=215, outline=40, width=2)
    for label, x in [
        ("#", cols[0] + 16),
        ("CHECK ITEM", cols[1] + 14),
        ("OK", cols[2] + 28),
        ("NOT OK", cols[3] + 8),
        ("REMARKS / READING", cols[4] + 14),
    ]:
        d.text((x, top + 10), label, font=sans_b(19), fill=30)
    items = [
        ("Hydraulic oil level in sight glass (above minimum mark)", True, "topped up 2 L"),
        ("Oil temperature at working load (limit 62 C)", False, f"{sf['temp']} C  too hot"),
        ("Ram seal and cylinder rod free of leaks", False, "weep at ram seal, drip tray half full"),
        ("Pump noise and vibration normal", False, "pump whine louder than last wk"),
        ("Guard interlocks and light curtain test", True, ""),
        ("Two-hand control and E-stop test", True, ""),
        ("Accumulator pre-charge gauge in green band", True, "in green"),
        ("Work area clean, no slip hazards", True, ""),
    ]
    for n, (text, ok, note) in enumerate(items):
        y = top + head_h + n * row_h
        d.rectangle([cols[0], y, cols[-1], y + row_h], outline=40, width=2)
        for x in cols[1:-1]:
            d.line([(x, y), (x, y + row_h)], fill=40, width=2)
        d.text((cols[0] + 18, y + 34), str(n + 1), font=sans(24), fill=30)
        for k, line in enumerate(textwrap.wrap(text, 38)):
            d.text((cols[1] + 14, y + 18 + k * 30), line, font=sans(23), fill=30)
        (tick if ok else cross)((cols[2] + cols[3]) // 2 if ok else (cols[3] + cols[4]) // 2, y + row_h // 2)
        for k, line in enumerate(textwrap.wrap(note, 21)[:3]):
            hand((cols[4] + 14, y + 8 + k * 31), line, size=32 if note.endswith("hot") else 28)

    y = top + head_h + len(items) * row_h + 35
    d.text((60, y), "Action raised (maintenance request no. and what was asked for):", font=sans(21), fill=40)
    d.line([(60, y + 105), (w - 60, y + 105)], fill=90, width=2)
    d.line([(60, y + 175), (w - 60, y + 175)], fill=90, width=2)
    hand((78, y + 36), "MR-2287 raised - replace ram seal kit + check pump.", size=34)
    hand((78, y + 108), "Run light duty only until fixed. Told D. Whitfield.", size=34)
    y2 = y + 215
    for label, x0, x1, value in [
        ("Operator signature", 60, 600, "T. Brennan"),
        ("Supervisor review / time", 640, w - 60, "D. Whitfield  14:20"),
    ]:
        d.rectangle([x0, y2, x1, y2 + 100], outline=50, width=2)
        d.text((x0 + 10, y2 + 5), label, font=sans(16), fill=70)
        hand((x0 + 30, y2 + 36), value, size=34, drift=2.0)
    d.text(
        (60, h - 70),
        "Report any abnormal noise, heat, leak or alarm to the maintenance supervisor immediately.",
        font=sans(17),
        fill=90,
    )

    # Scanner artefacts: soft focus, uneven paper, noise, specks, a fold, then a slight rotation.
    page = page.filter(ImageFilter.GaussianBlur(0.8))
    arr = np.asarray(page, dtype=np.float32)
    yy, xx = np.mgrid[0:h, 0:w]
    arr *= 1.0 - 0.10 * (((xx - w / 2) / (w / 2)) ** 2 + ((yy - h / 2) / (h / 2)) ** 2) / 2
    coarse = rng.normal(0, 1, (9, 7)).astype(np.float32)
    blotch = np.asarray(
        Image.fromarray(((coarse - coarse.min()) / np.ptp(coarse) * 255).astype(np.uint8)).resize(
            (w, h), Image.BICUBIC
        ),
        dtype=np.float32,
    )
    arr += (blotch - 128.0) / 128.0 * 7.0
    arr[int(h * 0.47) : int(h * 0.47) + 3, :] -= 9.0
    arr += rng.normal(0, 4.2, arr.shape).astype(np.float32)
    for _ in range(170):
        cx, cy, r = int(rng.integers(0, w)), int(rng.integers(0, h)), int(rng.integers(1, 3))
        arr[max(cy - r, 0) : cy + r, max(cx - r, 0) : cx + r] -= float(rng.uniform(50, 130))
    page = Image.fromarray(np.clip(arr, 0, 255).astype(np.uint8), mode="L")
    page = page.rotate(-1.3, resample=Image.BICUBIC, fillcolor=198)
    page.save(out_path, optimize=True)


# Answer key and README ----------------------------------------------------------------------------------------------


def md_table(header: list[str], rows: list[list[str]]) -> str:
    out = ["| " + " | ".join(header) + " |", "|" + "|".join("---" for _ in header) + "|"]
    out += ["| " + " | ".join(str(c) for c in r) + " |" for r in rows]
    return "\n".join(out)


def compute_answers(t: dict, f: dict, extra: dict, scan: dict) -> dict[str, str]:
    m = t["machines_df"]
    mi = m.set_index("machine_id")
    ev = t["events"].merge(m[["machine_id", "plant_id", "machine_type"]], on="machine_id")
    plan, sensor, defects = t["plan"], t["sensor"], t["defects"]
    unplanned = ev[ev["event_type"] == "unplanned"]
    a: dict[str, str] = {}

    a["pm2-interval-hydraulic-press"] = (
        "The PM-2 interval for a hydraulic press is **60 calendar days** (PM-MAN-001, Table 1), and its hydraulic "
        "oil is sampled **every 90 days**. PM-1 operator checks on a press are daily."
    )
    scan_row = sensor[(sensor["machine_id"] == SCAN_MACHINE) & (sensor["reading_date"] == SCAN_DATE)].iloc[0]
    a["round-sheet-press-findings"] = (
        f"The operator recorded **{scan['temp']} C** for the oil temperature at working load, above the 62 C limit "
        "printed on the sheet, so the item is marked NOT OK. For the ram seal the sheet notes a weep at the ram seal "
        "with the drip tray half full (also NOT OK) and the action **maintenance request MR-2287: replace the ram "
        "seal kit and check the pump**, with the press on light duty until fixed. Three items are marked NOT OK: "
        "oil temperature, ram seal leak and pump noise. The tables agree with the sheet: "
        f"{SCAN_MACHINE} reached {scan_row['max_temperature_c']:.1f} C on 14 September and logged an unplanned "
        f"HYD-LEAK event ({scan['leak_event']}) on {long_date(scan['leak_date'])}."
    )
    a["lockout-lock-removal"] = (
        "Only the **Plant EHS Manager (or a deputy named in writing)** may remove it (SAF-SOP-007, section 7); "
        "supervisors and co-workers may not. Before the removal: (1) the Maintenance Supervisor verifies that the "
        "employee is not on site and makes at least **two documented attempts to reach them by phone**; (2) the EHS "
        "Manager confirms the machine is safe to re-energize and nobody is working on it; (3) the employee is told "
        "the lock was removed **before they return to work on that machine**; (4) the removal is entered in the "
        "lockout log with the names and the reason."
    )

    q3_un = unplanned[(unplanned["started_at"] >= Q3[0]) & (unplanned["started_at"] < Q3[1])]
    g = q3_un.groupby("plant_id")["downtime_hours"].agg(["count", "sum"]).sort_values("sum", ascending=False)
    rows = [[PLANT_NAME[p], str(int(r["count"])), f"{r['sum']:.1f}"] for p, r in g.iterrows()]
    a["unplanned-downtime-by-plant"] = "\n".join(
        [
            f"**{PLANT_NAME[g.index[0]]}**, with **{g.iloc[0]['sum']:.1f} hours** of unplanned downtime over "
            f"{int(g.iloc[0]['count'])} events.",
            "",
            md_table(["Plant", "Unplanned events", "Downtime hours"], rows),
            "",
            "```sql",
            "SELECT p.plant_name, COUNT(*) AS events, ROUND(SUM(e.downtime_hours), 1) AS downtime_hours",
            "FROM maintenance_events e",
            "JOIN machines m ON m.machine_id = e.machine_id",
            "JOIN plants p ON p.plant_id = m.plant_id",
            "WHERE e.event_type = 'unplanned' AND e.started_at >= '2026-07-01' AND e.started_at < '2026-10-01'",
            "GROUP BY p.plant_name ORDER BY downtime_hours DESC;",
            "```",
        ]
    )

    last12 = unplanned[(unplanned["started_at"] >= LAST12[0]) & (unplanned["started_at"] < LAST12[1])]
    c = last12.groupby("failure_code")["downtime_hours"].agg(["count", "sum"]).sort_values("sum", ascending=False)
    assert c["sum"].iloc[2] > c["sum"].iloc[3] * 1.1
    rows = [[k, FAILURE_CODES[k][0], str(int(r["count"])), f"{r['sum']:.1f}"] for k, r in c.head(3).iterrows()]
    a["top-failure-codes"] = "\n".join(
        [
            "The top three failure codes by unplanned downtime hours, 1 October 2025 to 30 September 2026:",
            "",
            md_table(["Failure code", "Failure", "Events", "Downtime hours"], rows),
            "",
            f"The fourth is {c.index[3]} with {c['sum'].iloc[3]:.1f} hours.",
            "",
            "```sql",
            "SELECT failure_code, COUNT(*) AS events, ROUND(SUM(downtime_hours), 1) AS downtime_hours",
            "FROM maintenance_events",
            "WHERE event_type = 'unplanned' AND started_at >= '2025-10-01' AND started_at < '2026-10-01'",
            "GROUP BY failure_code ORDER BY downtime_hours DESC LIMIT 3;",
            "```",
        ]
    )

    qp = plan[(plan["week_start"] >= Q3_WEEKS[0]) & (plan["week_start"] <= Q3_WEEKS[1])]
    gp = qp.groupby("plant_id")[["actual_units", "target_units"]].sum()
    gp["pct"] = 100.0 * gp["actual_units"] / gp["target_units"]
    gp = gp.sort_values("pct")
    assert gp["pct"].iloc[1] - gp["pct"].iloc[0] > 0.5
    rows = [
        [PLANT_NAME[p], fmt_int(r["actual_units"]), fmt_int(r["target_units"]), f"{r['pct']:.1f}%"]
        for p, r in gp.iterrows()
    ]
    a["plan-attainment-q3"] = "\n".join(
        [
            f"**{PLANT_NAME[gp.index[0]]}** had the lowest attainment, **{gp['pct'].iloc[0]:.1f}%**.",
            "",
            md_table(["Plant", "Actual units", "Target units", "Attainment"], rows),
            "",
            "```sql",
            "SELECT p.plant_name, SUM(w.actual_units) AS actual, SUM(w.target_units) AS target,",
            "       ROUND(100.0 * SUM(w.actual_units) / SUM(w.target_units), 1) AS attainment_pct",
            "FROM production_plan_weekly_output w JOIN plants p ON p.plant_id = w.plant_id",
            "WHERE w.week_start BETWEEN '2026-07-06' AND '2026-09-21'",
            "GROUP BY p.plant_name ORDER BY attainment_pct;",
            "```",
        ]
    )

    last_pm2 = ev[ev["task_code"] == "PM-2"].groupby("machine_id")["started_at"].max()
    rows = []
    for mid, last in last_pm2.items():
        mtype = str(mi.loc[mid, "machine_type"])
        interval = TYPES[mtype]["pm2"]
        since = (AS_OF - date.fromisoformat(last[:10])).days
        assert abs(since - interval) >= 5, mid
        if since > interval:
            plant = PLANT_NAME[str(mi.loc[mid, "plant_id"])]
            rows.append([str(mid), plant, mtype, last[:10], str(since), str(interval), str(since - interval)])
    rows.sort(key=lambda r: -int(r[6]))
    a["overdue-pm2"] = "\n".join(
        [
            f"**{len(rows)} machines** are past their PM-2 interval at 30 September 2026 (days since the last PM-2 "
            "minus the Table 1 interval for the machine type):",
            "",
            md_table(["Machine", "Plant", "Type", "Last PM-2", "Days since", "Interval", "Days overdue"], rows),
            "",
            "Intervals from PM-MAN-001, Table 1: CNC Mill 90, CNC Lathe 120, Hydraulic Press 60, Robot Cell 180 and "
            "Surface Grinder 90 days.",
            "",
            "```sql",
            "SELECT m.machine_id, m.machine_type, MAX(e.started_at)::DATE AS last_pm2,",
            "       DATE_DIFF('day', MAX(e.started_at)::DATE, DATE '2026-09-30') AS days_since",
            "FROM machines m",
            "JOIN maintenance_events e ON e.machine_id = m.machine_id AND e.task_code = 'PM-2'",
            "GROUP BY m.machine_id, m.machine_type;  -- then compare days_since with the interval for the type",
            "```",
        ]
    )

    mine = unplanned[(unplanned["machine_id"] == BORE_MACHINE) & (unplanned["started_at"] >= "2026-07-01")]
    assert len(mine) == 1
    dmine = defects[(defects["machine_id"] == BORE_MACHINE) & (defects["detected_at"] >= "2026-07-01")]
    a["eight-d-machine-history"] = "\n".join(
        [
            f"The 8D report (8D-2026-014) names **{BORE_MACHINE}**, a CNC lathe at the Monterrey Plant. Since 1 July "
            f"2026 it has logged **{len(mine)} unplanned maintenance event** ({mine['failure_code'].iloc[0]}, "
            f"{long_date(mine['started_at'].iloc[0])}) with **{mine['downtime_hours'].sum():.1f} hours of unplanned "
            f"downtime**, and its quality defects detected since 1 July cost **${dmine['cost_usd'].sum():,.2f}** "
            f"({len(dmine)} records).",
            "",
            "```sql",
            "SELECT COUNT(*), SUM(downtime_hours) FROM maintenance_events",
            f"WHERE machine_id = '{BORE_MACHINE}' AND event_type = 'unplanned' AND started_at >= '2026-07-01';",
            f"SELECT SUM(cost_usd) FROM quality_defects WHERE machine_id = '{BORE_MACHINE}'",
            "  AND detected_at >= '2026-07-01';",
            "```",
        ]
    )

    by_machine: dict[str, list[str]] = {}
    for mid, day, code, _ramp, _hours in extra["hidden"]:
        by_machine.setdefault(mid, []).append(f"{day} ({code})")
    sim_rows = [
        [mid, PLANT_NAME[str(mi.loc[mid, "plant_id"])], str(mi.loc[mid, "machine_type"]), ", ".join(events)]
        for mid, events in by_machine.items()
    ]
    a["failure-risk-30d"] = "\n".join(
        [
            "Kumo returns probabilities, so the check is overlap, not exact values. The generator simulates October "
            "2026 beyond the data (the tables stop at 30 September). The machines below are in a degradation ramp "
            "at the anchor time and fail within 30 days; at least three of the five highest-probability machines "
            "should come from this list.",
            "",
            md_table(["Machine", "Plant", "Type", "Simulated failure (code)"], sim_rows),
            "",
            "Signals a model can see: rising `avg_vibration_mm_s` and `max_temperature_c`, more `alarm_count`, a "
            "late PM-2 and earlier `unplanned` events.",
        ]
    )

    top_alarm = sorted(((v, k) for k, v in extra["fut_alarms"].items() if v > 10), reverse=True)
    last30 = sensor[sensor["reading_date"] >= "2026-09-01"].groupby("machine_id")["alarm_count"].sum()
    recent = ", ".join(f"{k} ({v})" for k, v in last30.sort_values(ascending=False).head(5).items())
    a["high-alarm-risk"] = (
        "Kumo returns probabilities; the check is overlap with the simulated truth. The generator continues the "
        f"simulation through 30 October 2026: {len(top_alarm)} machines exceed 10 alarms in the 30 days after the "
        "anchor time: "
        + ", ".join(f"{k} ({v})" for v, k in top_alarm)
        + ". At least three of the five highest-probability machines should come from this list. The plant and "
        "type of each come from `machines` (`SELECT machine_id, plant_id, machine_type FROM machines WHERE "
        f"machine_id IN (...)`). For reference, the machines with the most alarms in September 2026 are {recent}."
    )
    return a


def pdf_pages(path: Path) -> int:
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(str(path))
    n = len(pdf)
    pdf.close()
    return n


def write_readme(t: dict, f: dict, counts: dict, answers: dict) -> None:
    import yaml

    qs = yaml.safe_load((PACK / "questions.yaml").read_text(encoding="utf-8"))
    template = re.sub(r"^(<!--.*-->\n)+", "", (CONTENT / "readme_template.md").read_text(encoding="utf-8"))
    desc = {
        "plants": ("one row per plant", "plant_id", "-"),
        "machines": ("one row per machine", "machine_id", "-"),
        "sensor_readings": ("one row per machine per operating day", "reading_id", "reading_date"),
        "maintenance_events": ("one row per planned or unplanned maintenance event", "event_id", "started_at"),
        "work_orders": ("one row per production work order", "work_order_id", "start_date"),
        "quality_defects": ("one row per quality defect record", "defect_id", "detected_at"),
        "production_plan_weekly_output": (
            "one row per plant per week (sheet `weekly_output`)",
            "plan_line_id",
            "week_start",
        ),
        "production_plan_scrap_targets": ("one row per plant per quarter (sheet `scrap_targets`)", "target_id", "-"),
    }
    table_rows = [[f"`{k}`", f"{v:,}", *desc[k]] for k, v in counts.items()]
    docs = []
    for sub, name, what in [
        (
            "procedures",
            "preventive_maintenance_manual.pdf",
            "Preventive maintenance manual PM-MAN-001 with the PM-2 interval table per machine type",
        ),
        ("procedures", "lockout_tagout_sop.pdf", "Lockout/tagout SOP SAF-SOP-007"),
        (
            "procedures",
            "pm1_round_sheet_scan.png",
            "Scanned handwritten PM-1 round sheet for press DAY-PRS-03 (grayscale, rotated 1.3 degrees, noise)",
        ),
        ("reports", "supplier_quality_agreement.docx", "Supplier quality agreement with Halvorsen Alloys Ltd."),
        ("reports", "q3_2026_plant_performance_review.pptx", "Q3 2026 plant performance review deck (8 slides)"),
        (
            "reports",
            "8d_report_valve_body_bore_oversize.pdf",
            "8D report 8D-2026-014 on the bore-diameter escape from MTY-LTH-04",
        ),
    ]:
        p = FILES / sub / name
        size = f"{p.stat().st_size / 1024:,.0f} KB"
        pages = f"{pdf_pages(p)} pages, " if name.endswith(".pdf") else ""
        docs.append([f"`files/{sub}/{name}`", what, pages + size])
    qrows = []
    key_blocks = []
    featured = {q["id"] for q in qs["questions"] if q.get("featured")}
    for q in qs["questions"]:
        qrows.append(
            [
                f"`{q['id']}`",
                q["tag"],
                "yes" if q["id"] in featured else "",
                ", ".join(q["tools"]),
                ", ".join(q["sources"]),
            ]
        )
        key_blocks.append(f"### `{q['id']}` ({q['tag']})\n\n> {q['question']}\n\n{answers[q['id']]}")
    out = template.format(
        row_total=f"{sum(counts.values()):,}",
        machine_count=f["machine_count"],
        tables_block=md_table(["Table", "Rows", "Grain", "Key", "Time column"], table_rows),
        documents_block=md_table(["File", "What it is", "Size"], docs),
        questions_block=md_table(["Question", "Tag", "Featured", "Tools", "Sources"], qrows),
        answer_key="\n\n".join(key_blocks),
        failures=str(int((t["events"]["event_type"] == "unplanned").sum())),
        planned=str(int((t["events"]["event_type"] == "planned").sum())),
        n_overdue=str(N_OVERDUE),
    )
    (PACK / "README.md").write_text(out, encoding="utf-8")


def render_previews(out_dir: Path) -> None:
    import pypdfium2 as pdfium

    out_dir.mkdir(parents=True, exist_ok=True)
    for pdf_path in sorted(FILES.glob("*/*.pdf")):
        pdf = pdfium.PdfDocument(str(pdf_path))
        for n in range(len(pdf)):
            pdf[n].render(scale=1.2).to_pil().save(out_dir / f"{pdf_path.stem}-p{n + 1}.png")
        pdf.close()


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--stats", action="store_true", help="print table row counts and stop")
    parser.add_argument("--preview", metavar="DIR", help="also render every PDF page to PNG files in DIR")
    args = parser.parse_args()
    t = build_tables()
    if args.stats:
        for name in ["plants", "machines_df", "sensor", "events", "work_orders", "defects", "plan", "targets"]:
            print(name, len(t[name]))
        return 0
    for sub in ("tables", "procedures", "reports"):
        (FILES / sub).mkdir(parents=True, exist_ok=True)
    counts = write_tables(t)
    f, tables, extra = compute_facts(t)
    extra["hidden"] = t["hidden_failures"]
    extra["fut_alarms"] = t["future_alarms"]
    scan = compute_scan_facts(t)
    build_pdf("pm_manual.md", FILES / "procedures" / "preventive_maintenance_manual.pdf", f, tables, extra)
    build_pdf("loto_sop.md", FILES / "procedures" / "lockout_tagout_sop.pdf", f, tables, extra)
    build_pdf("eight_d_report.md", FILES / "reports" / "8d_report_valve_body_bore_oversize.pdf", f, tables, extra)
    build_docx("supplier_quality_agreement.md", FILES / "reports" / "supplier_quality_agreement.docx", f, tables)
    build_pptx("plant_review_q3.md", FILES / "reports" / "q3_2026_plant_performance_review.pptx", f, tables, extra)
    render_scan(FILES / "procedures" / "pm1_round_sheet_scan.png", scan)
    answers = compute_answers(t, f, extra, scan)
    write_readme(t, f, counts, answers)
    if args.preview:
        render_previews(Path(args.preview))
    total = sum(p.stat().st_size for p in FILES.rglob("*") if p.is_file())
    print(f"wrote {len(counts)} tables ({sum(counts.values()):,} rows) and 6 documents; files/ is {total / 1e6:.2f} MB")
    return 0


if __name__ == "__main__":
    sys.exit(main())
