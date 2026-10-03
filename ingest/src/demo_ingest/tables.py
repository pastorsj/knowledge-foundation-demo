# SPDX-FileCopyrightText: Copyright (c) 2026, NVIDIA CORPORATION & AFFILIATES. All rights reserved.
# SPDX-License-Identifier: Apache-2.0
"""Tables: spreadsheets and data files into one DuckDB file per structured source, then profiled.

Loading: CSV/TSV with ``read_csv`` (the whole file sampled), XLSX one table per sheet with the excel extension's
``read_xlsx``, Parquet with ``read_parquet``, JSON and JSON Lines with ``read_json_auto``. Table names are the
snake-cased file stem (plus ``_<sheet>`` for a workbook of several sheets), unique in the database; column names are
snake-cased, non-empty and unique. ``read_xlsx`` types each column from its first data row and drops a column whose
header cell is empty, so a sheet where either happens is read with openpyxl instead and typed by ``read_csv``.

Profiling fills the source manifest's TableInfo: row counts, column types, nullability, a primary key, a time column
and foreign keys. Declarations from a pack's ``pack.yaml`` win over what profiling infers.
"""

from __future__ import annotations

import csv
import datetime
import logging
import os
import re
import tempfile
import time
import unicodedata
from collections.abc import Callable
from collections.abc import Iterable
from contextlib import contextmanager
from dataclasses import dataclass
from dataclasses import field
from pathlib import Path
from typing import Any

import duckdb

from .models import IngestError

logger = logging.getLogger(__name__)

PARSERS = {
    ".csv": "duckdb-csv",
    ".tsv": "duckdb-csv",
    ".xlsx": "duckdb-xlsx",
    ".parquet": "duckdb-parquet",
    ".json": "duckdb-json",
    ".jsonl": "duckdb-json",
}
TIME_TYPES = ("DATE", "TIMESTAMP")
# Types a guessed (undeclared, not id-named) primary key may have: never a measure, a date or a flag.
KEY_TYPES = ("VARCHAR", "BIGINT", "INTEGER", "HUGEINT", "UBIGINT", "UINTEGER", "SMALLINT", "UUID")
LOCK_WAIT_SECONDS = 30.0

Progress = Callable[[int, int], None]


@dataclass
class LoadResult:
    tables: list[str]
    parser: str
    warnings: list[str] = field(default_factory=list)


# Names


def snake_case(name: str) -> str:
    """ASCII snake case: accents dropped, camelCase split, every other run of symbols one underscore."""
    text = unicodedata.normalize("NFKD", str(name))
    text = "".join(c for c in text if not unicodedata.combining(c))
    text = re.sub(r"([a-z0-9])([A-Z])", r"\1_\2", text)
    text = re.sub(r"([A-Z]+)([A-Z][a-z])", r"\1_\2", text)
    return re.sub(r"[^a-z0-9]+", "_", text.lower()).strip("_")


def _identifier(name: str, *, fallback: str, digit_prefix: str) -> str:
    snake = snake_case(name)[:100].strip("_") or fallback
    return f"{digit_prefix}_{snake}" if snake[0].isdigit() else snake


def unique(name: str, taken: Iterable[str]) -> str:
    taken = set(taken)
    if name not in taken:
        return name
    number = 2
    while f"{name}_{number}" in taken:
        number += 1
    return f"{name}_{number}"


def column_names(raw: list[Any]) -> list[str]:
    names: list[str] = []
    for value in raw:
        text = "" if value is None else str(value)
        names.append(unique(_identifier(text, fallback="col", digit_prefix="col"), names))
    return names


def table_name(stem: str, sheet: str | None, existing: Iterable[str]) -> str:
    base = _identifier(stem, fallback="table", digit_prefix="t")
    if sheet is not None:
        base = f"{base}_{_identifier(sheet, fallback='sheet', digit_prefix='s')}"
    return unique(base, existing)


def _undo_duckdb_dedup(names: list[str]) -> list[str]:
    """DuckDB renames a repeated header ``x`` to ``x_1``, ``x_2``: give it back its own name, so it is ``x_2``."""
    raw: list[str] = []
    for name in names:
        match = re.fullmatch(r"(.*)_(\d+)", name)
        raw.append(match[1] if match and match[1] in raw else name)
    return raw


def _quote(identifier: str) -> str:
    return '"' + identifier.replace('"', '""') + '"'


# Connections


@contextmanager
def connect(db_path: Path, *, read_only: bool = False):
    """A short-lived connection. Readers in other services open the file read-only, so ingest never holds it long,
    and waits (rather than fails) while one of them has it open."""
    config = {}
    if directory := os.environ.get("DUCKDB_EXTENSION_DIRECTORY"):
        config["extension_directory"] = directory
    db_path.parent.mkdir(parents=True, exist_ok=True)
    deadline = time.monotonic() + LOCK_WAIT_SECONDS
    while True:
        try:
            connection = duckdb.connect(str(db_path), read_only=read_only, config=config)
            break
        except duckdb.IOException as error:
            if "lock" not in str(error).lower() or time.monotonic() > deadline:
                raise
            time.sleep(0.5)
    try:
        yield connection
    finally:
        connection.close()


def _load_excel(connection: duckdb.DuckDBPyConnection) -> None:
    try:
        connection.execute("LOAD excel")
    except duckdb.Error:  # the image installs it at build time; a dev machine downloads it once
        connection.execute("INSTALL excel")
        connection.execute("LOAD excel")


# Loading


def load_table_file(
    db_path: Path, file_path: Path, file_name: str, existing: set[str], on_progress: Progress | None = None
) -> LoadResult:
    """Load one file into ``db_path`` as one table (or one per sheet); names avoid ``existing``."""
    suffix = Path(file_name).suffix.lower()
    if suffix not in PARSERS:
        raise IngestError("unsupported_type", f"{suffix or 'A file without an extension'} is not a table format.")
    if file_path.stat().st_size == 0:
        raise IngestError("empty_file")
    stem = Path(file_name).stem
    taken = set(existing)
    with connect(db_path) as connection:
        if suffix == ".xlsx":
            return _load_workbook(
                connection, file_path, file_name=file_name, stem=stem, taken=taken, on_progress=on_progress
            )
        name = table_name(stem, None, taken)
        if on_progress:
            on_progress(1, 1)
        if suffix in (".csv", ".tsv"):
            _create_from_csv(connection, name, file_path, delimiter="\t" if suffix == ".tsv" else None)
        else:
            reader = "read_parquet" if suffix == ".parquet" else "read_json_auto"
            _create(connection, name, f"{reader}(?)", [str(file_path)])
        _finish(connection, name, file_name)
        return LoadResult(tables=[name], parser=PARSERS[suffix])


def _create(connection: duckdb.DuckDBPyConnection, name: str, relation: str, params: list[Any]) -> None:
    """CREATE OR REPLACE the table from ``relation``, its columns renamed to clean, unique names."""
    try:
        described = connection.execute(f"DESCRIBE SELECT * FROM {relation}", params).fetchall()
    except duckdb.Error as error:
        raise IngestError("load_failed", _first_line(error)) from error
    _create_renamed(
        connection, name, relation, params, source_columns=[row[0] for row in described], raw_names=_undo_duckdb_dedup
    )


def _create_renamed(
    connection: duckdb.DuckDBPyConnection,
    name: str,
    relation: str,
    params: list[Any],
    *,
    source_columns: list[str],
    raw_names: Callable[[list[str]], list[Any]] | list[Any],
) -> None:
    raw = raw_names(source_columns) if callable(raw_names) else raw_names
    targets = column_names(raw)
    select = ", ".join(f"{_quote(src)} AS {_quote(dst)}" for src, dst in zip(source_columns, targets, strict=True))
    try:
        connection.execute(f"CREATE OR REPLACE TABLE {_quote(name)} AS SELECT {select} FROM {relation}", params)
    except duckdb.Error as error:
        raise IngestError("load_failed", _first_line(error)) from error


def _create_from_csv(
    connection: duckdb.DuckDBPyConnection, name: str, path: Path, *, delimiter: str | None = None
) -> None:
    options = ", delim = ?" if delimiter else ""
    params: list[Any] = [str(path), *([delimiter] if delimiter else [])]
    try:
        _create(connection, name, f"read_csv(?, sample_size = -1{options})", params)
    except IngestError as error:
        if "not utf-8 encoded" not in str(error.__cause__):
            raise
        # Spreadsheet exports are often Windows-1252; latin-1 decodes every byte, so the text survives.
        _create(connection, name, f"read_csv(?, sample_size = -1, encoding = 'latin-1'{options})", params)


def _finish(connection: duckdb.DuckDBPyConnection, name: str, origin: str) -> None:
    rows = connection.execute(f"SELECT count(*) FROM {_quote(name)}").fetchone()[0]
    if rows == 0:
        connection.execute(f"DROP TABLE {_quote(name)}")
        raise IngestError("empty_file", "The file has no rows.")
    literal = "'" + origin[:300].replace("'", "''") + "'"  # COMMENT ON takes no parameters
    connection.execute(f"COMMENT ON TABLE {_quote(name)} IS {literal}")


def _load_workbook(
    connection: duckdb.DuckDBPyConnection,
    path: Path,
    *,
    file_name: str,
    stem: str,
    taken: set[str],
    on_progress: Progress | None,
) -> LoadResult:
    import openpyxl

    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as error:  # zipfile, XML and openpyxl errors alike
        raise IngestError("load_failed", f"The workbook could not be opened: {_first_line(error)}") from error
    try:
        sheets = list(workbook.sheetnames)
        _load_excel(connection)
        loaded: list[str] = []
        warnings: list[str] = []
        for number, sheet in enumerate(sheets, start=1):
            if on_progress:
                on_progress(number, len(sheets))
            name = table_name(stem, sheet if len(sheets) > 1 else None, taken)
            try:
                _load_sheet(connection, workbook[sheet], path, sheet, name)
                origin = f"{file_name} › {sheet}" if len(sheets) > 1 else file_name
                _finish(connection, name, origin)
            except IngestError as error:
                warnings.append(f"Sheet {sheet!r} was not loaded: {error.message}")
                continue
            loaded.append(name)
            taken.add(name)
    finally:
        workbook.close()
    if not loaded:
        raise IngestError("empty_file", "No sheet of the workbook has a header and rows.")
    return LoadResult(tables=loaded, parser=PARSERS[".xlsx"], warnings=warnings)


def _load_sheet(connection: duckdb.DuckDBPyConnection, worksheet: Any, path: Path, sheet: str, name: str) -> None:
    header = _header_row(worksheet)
    if header is None:
        raise IngestError("empty_file", "The sheet is empty.")
    relation = "read_xlsx(?, sheet = ?, header = true)"
    params = [str(path), sheet]
    try:
        described = connection.execute(f"DESCRIBE SELECT * FROM {relation}", params).fetchall()
        columns = [row[0] for row in described]
        if len(columns) != len(header):
            raise ValueError("read_xlsx dropped a column whose header cell is empty")
        _create_renamed(connection, name, relation, params, source_columns=columns, raw_names=header)
    except (duckdb.Error, ValueError, IngestError) as error:
        logger.info("sheet %r: %s; reading it with openpyxl", sheet, _first_line(error))
        _load_sheet_as_csv(connection, worksheet, len(header), name)


def _header_row(worksheet: Any) -> list[Any] | None:
    """The first row with a value, as wide as the sheet's dimension."""
    for row in worksheet.iter_rows(values_only=True):
        if any(value not in (None, "") for value in row):
            return list(row)
    return None


def _load_sheet_as_csv(connection: duckdb.DuckDBPyConnection, worksheet: Any, width: int, name: str) -> None:
    """Every cell through openpyxl, so mixed columns become text and nothing is dropped; read_csv types the rest.

    Columns with neither a header nor a value (formatted cells past the data) are left out.
    """
    rows = worksheet.iter_rows(values_only=True)
    header: list[Any] = []
    for row in rows:
        if any(value not in (None, "") for value in row):
            header = list(row)
            break
    used = max((i + 1 for i, value in enumerate(header) if value not in (None, "")), default=0)
    count = 0
    with tempfile.TemporaryDirectory(prefix="demo-ingest-") as directory:
        path = Path(directory) / "sheet.csv"
        with path.open("w", newline="", encoding="utf-8") as out:
            writer = csv.writer(out)
            writer.writerow([f"c{i}" for i in range(width)])
            for row in rows:
                values = [_cell(value) for value in list(row)[:width]]
                if any(values):
                    writer.writerow(values + [""] * (width - len(values)))
                    used = max(used, *(i + 1 for i, value in enumerate(values) if value))
                    count += 1
        if count == 0:
            raise IngestError("empty_file", "The sheet has a header but no rows.")
        relation = "read_csv(?, header = true, sample_size = -1)"
        columns = [f"c{i}" for i in range(used)]
        _create_renamed(connection, name, relation, [str(path)], source_columns=columns, raw_names=header[:used])


def _cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime.datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, (datetime.date, datetime.time)):
        return value.isoformat()
    return str(value)


def drop_tables(db_path: Path, names: Iterable[str]) -> None:
    if not db_path.exists():
        return
    with connect(db_path) as connection:
        for name in names:
            connection.execute(f"DROP TABLE IF EXISTS {_quote(name)}")


def table_names(db_path: Path) -> list[str]:
    if not db_path.exists():
        return []
    with connect(db_path, read_only=True) as connection:
        return [row[0] for row in connection.execute("SELECT table_name FROM duckdb_tables() ORDER BY table_name")]


# Profiling


def profile_database(db_path: Path, declarations: dict[str, dict] | None) -> list[dict[str, Any]]:
    """TableInfo for every table, in name order."""
    declarations = declarations or {}
    if not db_path.exists():
        return []
    with connect(db_path, read_only=True) as connection:
        tables = connection.execute(
            "SELECT table_name, comment FROM duckdb_tables() WHERE schema_name = 'main' ORDER BY table_name"
        ).fetchall()
        infos = [_profile_table(connection, name, comment, declarations.get(name, {})) for name, comment in tables]
        keys = {info["name"]: info["primary_key"] for info in infos if info["primary_key"]}
        for info in infos:
            info["foreign_keys"] = _foreign_keys(connection, info, keys, declarations.get(info["name"], {}))
    return infos


def _profile_table(
    connection: duckdb.DuckDBPyConnection, name: str, comment: str | None, declared: dict[str, Any]
) -> dict[str, Any]:
    described = connection.execute(f"DESCRIBE {_quote(name)}").fetchall()
    columns = [(row[0], row[1]) for row in described]
    stats = ", ".join(f"count({_quote(c)}), count(DISTINCT {_quote(c)})" for c, _ in columns)
    row = connection.execute(f"SELECT count(*), {stats} FROM {_quote(name)}").fetchone()
    row_count = row[0]
    non_null = {c: row[1 + 2 * i] for i, (c, _) in enumerate(columns)}
    distinct = {c: row[2 + 2 * i] for i, (c, _) in enumerate(columns)}
    descriptions = declared.get("columns", {})
    names = [c for c, _ in columns]
    types = dict(columns)

    primary_key = declared.get("primary_key") if declared.get("primary_key") in types else None
    if primary_key is None and row_count:
        is_key = [c for c in names if non_null[c] == row_count and distinct[c] == row_count]
        preferred = [c for c in is_key if c == "id" or c.endswith("_id")]
        fallback = [c for c in is_key if types[c].startswith(KEY_TYPES)]
        primary_key = (preferred or fallback or [None])[0]

    time_column = declared.get("time_column") if declared.get("time_column") in types else None
    if time_column is None:
        timed = [c for c in names if types[c].startswith(TIME_TYPES)]
        preferred = [c for c in timed if c.endswith(("_at", "_date", "date"))]
        time_column = (preferred or timed or [None])[0]

    info: dict[str, Any] = {
        "name": name,
        "description": declared.get("description", ""),
        "row_count": row_count,
        "primary_key": primary_key,
        "time_column": time_column,
        "columns": [
            {
                "name": c,
                "type": t[:64],
                "description": descriptions.get(c, ""),
                "nullable": non_null[c] < row_count,
            }
            for c, t in columns
        ],
        "foreign_keys": [],
    }
    if comment:
        info["origin_file"] = comment
    return info


def _foreign_keys(
    connection: duckdb.DuckDBPyConnection, info: dict[str, Any], keys: dict[str, str], declared: dict[str, Any]
) -> list[dict[str, Any]]:
    found = []
    for fk in declared.get("foreign_keys", []):
        table, column = fk["references"].split(".", 1)
        found.append(
            {"column": fk["column"], "references_table": table, "references_column": column, "inferred": False}
        )
    declared_columns = {fk["column"] for fk in found}
    for column in info["columns"]:
        name = column["name"]
        if name in declared_columns or name == info["primary_key"]:
            continue
        for table, key in keys.items():
            if table == info["name"] or not _names_agree(name, table, key):
                continue
            if _contained(connection, info["name"], name, table, key):
                found.append({"column": name, "references_table": table, "references_column": key, "inferred": True})
                break
    return found


def _names_agree(column: str, table: str, key: str) -> bool:
    """Names that may join: customer_id or billing_customer_id to customers.customer_id; customer_id to customers.id."""
    if column == key or (key.endswith("_id") and column.endswith(f"_{key}")):
        return True
    if key == "id" and column.endswith("_id"):
        stem = column.removesuffix("_id")
        return table in (stem, f"{stem}s", f"{stem}es") or table.removesuffix("s") == stem
    return False


def _contained(connection: duckdb.DuckDBPyConnection, table: str, column: str, ref_table: str, ref_key: str) -> bool:
    """Every distinct non-null value of table.column appears in ref_table.ref_key (and there is at least one)."""
    t, c, rt, rk = map(_quote, (table, column, ref_table, ref_key))
    try:
        values, missing = connection.execute(
            f"SELECT count(*), count(*) FILTER (WHERE v NOT IN (SELECT {rk} FROM {rt})) "
            f"FROM (SELECT DISTINCT {c} AS v FROM {t} WHERE {c} IS NOT NULL)"
        ).fetchone()
    except duckdb.Error:  # types that do not compare
        return False
    return values > 0 and missing == 0


def _first_line(error: BaseException) -> str:
    return (str(error).strip().splitlines() or [type(error).__name__])[0][:500]
