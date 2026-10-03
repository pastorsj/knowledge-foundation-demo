# SPDX-FileCopyrightText: Copyright (c) 2026, NVIDIA CORPORATION & AFFILIATES. All rights reserved.
# SPDX-License-Identifier: Apache-2.0
"""Offline fixtures: settings over a temporary knowledge root, generated documents and workbooks, a fake embedder,
a small whitespace tokenizer, and a fake Nemotron Parse server.

The fake Parse server is OpenAI-compatible, as the parse service (vLLM) is. It replays the Nemotron Parse 2.0
responses in ``fixtures/parse/``, recorded from the real model, chosen by the sha256 of the page image docling sends.
To record them again against a running parse service (concurrency 1, so pages arrive in order):

    INGEST_RECORD_PARSE_URL=http://127.0.0.1:8340/v1 uv run --directory ingest pytest tests/test_documents.py -k parse
"""

from __future__ import annotations

import base64
import hashlib
import json
import math
import os
import re
import threading
from collections.abc import Callable
from collections.abc import Iterator
from dataclasses import dataclass
from dataclasses import field
from dataclasses import replace
from http.server import BaseHTTPRequestHandler
from http.server import ThreadingHTTPServer
from pathlib import Path
from typing import Any

import pytest
import requests

from demo_ingest.catalog import Catalog
from demo_ingest.settings import Settings

REPO = Path(__file__).resolve().parents[2]
CATALOG_FIXTURES = REPO / "contracts" / "fixtures" / "catalog"
FIXTURES = Path(__file__).resolve().parent / "fixtures"
PARSE_FIXTURES = FIXTURES / "parse"
PARSE_MODEL = "nvidia/NVIDIA-Nemotron-Parse-2.0"
RECORD_FROM = os.environ.get("INGEST_RECORD_PARSE_URL", "").rstrip("/")
DIMENSION = 2048  # nvidia/nemotron-3-embed-1b


@pytest.fixture
def knowledge_dir(tmp_path: Path) -> Path:
    path = tmp_path / "knowledge"
    path.mkdir()
    return path


@pytest.fixture
def settings(tmp_path: Path, knowledge_dir: Path) -> Settings:
    return Settings.from_env(
        {
            "KNOWLEDGE_DIR": str(knowledge_dir),
            "PACKS_DIR": str(tmp_path / "packs"),
            "MILVUS_URI": str(tmp_path / "milvus.db"),
            "PARSE_BASE_URL": "",
            "RETRIEVER_API_KEY": "nvapi-test-dummy",
        },
        secrets_dir=tmp_path / "no-secrets",
    )


@pytest.fixture
def catalog(knowledge_dir: Path) -> Catalog:
    return Catalog(knowledge_dir)


# Workbooks

AWKWARD_HEADER = ["Store #", "Net Sales ($)", "Net Sales ($)", "Région"]


def write_awkward_xlsx(path: Path) -> Path:
    """Two sheets with spaces, symbols, a duplicate and a non-ASCII header; the second starts with blank rows."""
    import openpyxl

    workbook = openpyxl.Workbook()
    q1 = workbook.active
    q1.title = "Q1 Sales"
    q1.append(AWKWARD_HEADER)
    q1.append([1, 100.5, 90.0, "Nord"])
    q1.append([2, 200.0, 180.0, "Sud"])
    q2 = workbook.create_sheet("Q2 Sales")
    q2.append([None] * 4)
    q2.append([None] * 4)
    q2.append(AWKWARD_HEADER)
    q2.append([1, 110.5, 95.0, "Nord"])
    workbook.save(path)
    return path


@pytest.fixture
def awkward_xlsx(tmp_path: Path) -> Path:
    return write_awkward_xlsx(tmp_path / "awkward.xlsx")


# Documents

POLICY_TABLE = [
    ["Category", "Return window", "Restocking fee"],
    ["Electronics", "15 days", "15%"],
    ["Furniture", "30 days", "10%"],
    ["Apparel", "60 days", "None"],
]


def write_policy_pdf(path: Path) -> Path:
    """Two pages: a title, a heading, a paragraph and a 4x3 table, then a second page. Byte-for-byte reproducible."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import PageBreak
    from reportlab.platypus import Paragraph
    from reportlab.platypus import SimpleDocTemplate
    from reportlab.platypus import Spacer
    from reportlab.platypus import Table
    from reportlab.platypus import TableStyle

    styles = getSampleStyleSheet()
    table = Table(POLICY_TABLE)
    table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black)]))
    story = [
        Paragraph("Northwind Return Policy", styles["Title"]),
        Paragraph("Restocking fees", styles["Heading1"]),
        Paragraph(
            "Northwind Retail accepts returns of unused merchandise with a receipt. A restocking fee applies to "
            "opened electronics and assembled furniture, as the table below shows. This policy is synthetic.",
            styles["BodyText"],
        ),
        Spacer(1, 12),
        table,
        PageBreak(),
        Paragraph("Exceptions", styles["Heading1"]),
        Paragraph(
            "Clearance items are final sale. Gift cards cannot be returned. Store managers may waive the "
            "restocking fee for loyalty members of the gold tier.",
            styles["BodyText"],
        ),
    ]
    SimpleDocTemplate(str(path), pagesize=letter, invariant=1, title="Northwind Return Policy").build(story)
    return path


def write_scan_png(path: Path, pdf: Path) -> Path:
    """The PDF's second page as an image, as a scanner would produce it."""
    import pypdfium2

    document = pypdfium2.PdfDocument(str(pdf))
    try:
        document[1].render(scale=1.5).to_pil().convert("RGB").save(path, format="PNG")
    finally:
        document.close()
    return path


def write_docx(path: Path) -> Path:
    import docx

    document = docx.Document()
    document.add_heading("Store Operations", level=1)
    document.add_paragraph("Stores open at 9 a.m. and close at 9 p.m. on weekdays. This document is synthetic.")
    document.add_heading("Cash handling", level=2)
    document.add_paragraph("Count the drawer at every shift change and record the total in the log.")
    table = document.add_table(rows=2, cols=2)
    for row, values in zip(table.rows, [["Shift", "Count"], ["Morning", "Twice"]], strict=True):
        for cell, value in zip(row.cells, values, strict=True):
            cell.text = value
    document.save(path)
    return path


@pytest.fixture
def policy_pdf(tmp_path: Path) -> Path:
    return write_policy_pdf(tmp_path / "policy.pdf")


@pytest.fixture
def scan_png(tmp_path: Path, policy_pdf: Path) -> Path:
    return write_scan_png(tmp_path / "scan.png", policy_pdf)


@pytest.fixture
def docx_file(tmp_path: Path) -> Path:
    return write_docx(tmp_path / "operations.docx")


# Embeddings and tokens


def words(text: str) -> list[str]:
    return re.findall(r"[a-z0-9]+", text.lower())


def embed(text: str) -> list[float]:
    """A hashed bag of words, so texts that share words get similar vectors."""
    vector = [0.0] * DIMENSION
    for word in words(text) or ["empty"]:
        vector[int(hashlib.sha256(word.encode()).hexdigest(), 16) % DIMENSION] += 1.0
    norm = math.sqrt(sum(value * value for value in vector))
    return [value / norm for value in vector]


@dataclass
class FakeEmbedder:
    """Embedder's interface, offline: deterministic vectors of the embed model's dimension."""

    model: str = "nvidia/nemotron-3-embed-1b"
    calls: list[list[str]] = field(default_factory=list)
    fail_with: Exception | None = None

    def embed_documents(self, texts: list[str], on_progress: Callable[[int, int], None] | None = None):
        if self.fail_with is not None:
            raise self.fail_with
        self.calls.append(list(texts))
        if on_progress:
            on_progress(len(texts), len(texts))
        return [embed(text) for text in texts]


@pytest.fixture
def embedder() -> FakeEmbedder:
    return FakeEmbedder()


def whitespace_tokenizer(max_tokens: int = 128):
    """A tokenizer that needs no download: one token per word or symbol."""
    from docling_core.transforms.chunker.tokenizer.huggingface import HuggingFaceTokenizer
    from tokenizers import Tokenizer
    from tokenizers import models
    from tokenizers import pre_tokenizers
    from transformers import PreTrainedTokenizerFast

    backend = Tokenizer(models.WordLevel(vocab={"[UNK]": 0}, unk_token="[UNK]"))
    backend.pre_tokenizer = pre_tokenizers.Whitespace()
    return HuggingFaceTokenizer(
        tokenizer=PreTrainedTokenizerFast(tokenizer_object=backend, unk_token="[UNK]"), max_tokens=max_tokens
    )


@pytest.fixture(scope="session")
def tokenizer():
    return whitespace_tokenizer()


# Nemotron Parse


@dataclass
class FakeParse:
    """An OpenAI-compatible Nemotron Parse that replays recorded responses (or records them, see the module doc)."""

    label: str = "unlabelled"
    delay: float = 0.0
    fail_status: int | None = None
    requests: list[dict[str, Any]] = field(default_factory=list)
    server: ThreadingHTTPServer | None = None
    _served: dict[str, int] = field(default_factory=dict)
    _lock: threading.Lock = field(default_factory=threading.Lock)

    @property
    def url(self) -> str:
        assert self.server is not None
        return f"http://127.0.0.1:{self.server.server_address[1]}/v1"

    def index(self) -> dict[str, str]:
        path = PARSE_FIXTURES / "index.json"
        return json.loads(path.read_text()) if path.exists() else {}

    def answer(self, body: dict[str, Any]) -> tuple[int, dict[str, Any]]:
        image = body["messages"][0]["content"][0]["image_url"]["url"]
        digest = hashlib.sha256(base64.b64decode(image.split(",", 1)[1])).hexdigest()
        with self._lock:
            self.requests.append({**body, "image_sha256": digest})
            if self.fail_status is not None:
                return self.fail_status, {"error": {"message": "injected failure"}}
            if RECORD_FROM:
                return 200, self._record(body, digest)
            index = self.index()
            name = index.get(digest)
            if name is None:  # an image this version renders differently: replay this label's pages in order
                names = sorted(n for n in set(index.values()) if n.startswith(f"{self.label}-"))
                assert names, f"no recorded Parse response for {self.label}"
                name = names[self._served.get(self.label, 0) % len(names)]
                self._served[self.label] = self._served.get(self.label, 0) + 1
            return 200, json.loads((PARSE_FIXTURES / name).read_text())

    def _record(self, body: dict[str, Any], digest: str) -> dict[str, Any]:
        response = requests.post(f"{RECORD_FROM}/chat/completions", json=body, timeout=600)
        response.raise_for_status()
        payload = response.json()
        index = self.index()
        number = sum(1 for name in index.values() if name.startswith(f"{self.label}-")) + 1
        name = index.get(digest) or f"{self.label}-{number}.json"
        PARSE_FIXTURES.mkdir(parents=True, exist_ok=True)
        (PARSE_FIXTURES / name).write_text(json.dumps(payload, indent=2) + "\n")
        index[digest] = name
        (PARSE_FIXTURES / "index.json").write_text(json.dumps(index, indent=2, sort_keys=True) + "\n")
        return payload


def _handler(fake: FakeParse) -> type[BaseHTTPRequestHandler]:
    class Handler(BaseHTTPRequestHandler):
        def log_message(self, format: str, *args: Any) -> None:
            pass

        def _send(self, status: int, payload: dict[str, Any]) -> None:
            data = json.dumps(payload).encode()
            self.send_response(status)
            self.send_header("Content-Type", "application/json")
            self.send_header("Content-Length", str(len(data)))
            self.end_headers()
            self.wfile.write(data)

        def do_GET(self) -> None:
            if self.path.rstrip("/") == "/v1/models":
                self._send(200, {"object": "list", "data": [{"id": PARSE_MODEL, "object": "model"}]})
            else:
                self._send(404, {"error": {"message": "not found"}})

        def do_POST(self) -> None:
            body = json.loads(self.rfile.read(int(self.headers["Content-Length"])))
            if fake.delay:
                threading.Event().wait(fake.delay)
            try:
                self._send(*fake.answer(body))
            except (BrokenPipeError, ConnectionResetError):
                pass

    return Handler


@pytest.fixture
def fake_parse() -> Iterator[FakeParse]:
    fake = FakeParse()
    fake.server = ThreadingHTTPServer(("127.0.0.1", 0), _handler(fake))
    fake.server.daemon_threads = True
    thread = threading.Thread(target=fake.server.serve_forever, daemon=True)
    thread.start()
    yield fake
    fake.server.shutdown()
    fake.server.server_close()


@pytest.fixture
def parse_settings(settings: Settings, fake_parse: FakeParse) -> Settings:
    """Settings that point at the fake Parse server (one page at a time while recording, so pages keep their order)."""
    return replace(settings, parse_base_url=fake_parse.url, parse_concurrency=1 if RECORD_FROM else 4)


def closed_port_url() -> str:
    """A URL nothing listens on."""
    import socket

    with socket.socket() as probe:
        probe.bind(("127.0.0.1", 0))
        port = probe.getsockname()[1]
    return f"http://127.0.0.1:{port}/v1"
