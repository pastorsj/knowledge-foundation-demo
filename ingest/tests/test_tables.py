# SPDX-FileCopyrightText: Copyright (c) 2026, NVIDIA CORPORATION & AFFILIATES. All rights reserved.
# SPDX-License-Identifier: Apache-2.0

from __future__ import annotations

import datetime
import json
from pathlib import Path

import duckdb
import openpyxl
import pytest
from conftest import CATALOG_FIXTURES
from conftest import FIXTURES
from jsonschema import Draft202012Validator

from demo_ingest import tables
from demo_ingest.models import IngestError

TABLES = CATALOG_FIXTURES / "tables"
SOURCE_SCHEMA = json.loads((CATALOG_FIXTURES.parents[1] / "catalog" / "source-manifest.schema.json").read_text())
TABLE_INFO = Draft202012Validator({"$defs": SOURCE_SCHEMA["$defs"], "$ref": "#/$defs/TableInfo"})


def columns(db: Path, table: str) -> list[str]:
    with duckdb.connect(str(db), read_only=True) as connection:
        return [row[0] for row in connection.execute(f'DESCRIBE "{table}"').fetchall()]


@pytest.fixture
def db(tmp_path: Path) -> Path:
    return tmp_path / "tables.duckdb"


@pytest.fixture
def retail(db: Path) -> Path:
    existing: set[str] = set()
    for name in ("customers.csv", "orders.csv"):
        result = tables.load_table_file(db, TABLES / name, name, existing)
        existing.update(result.tables)
    return db


def test_csv_files_load_into_tables_named_after_their_stem(db: Path):
    result = tables.load_table_file(db, TABLES / "orders.csv", "orders.csv", set())

    assert (result.tables, result.parser, result.warnings) == (["orders"], "duckdb-csv", [])
    assert columns(db, "orders") == ["order_id", "customer_id", "ordered_at", "net_amount"]


def test_profiling_finds_keys_time_columns_and_foreign_keys(retail: Path):
    profiles = {info["name"]: info for info in tables.profile_database(retail, None)}

    assert set(profiles) == {"customers", "orders"}
    customers, orders = profiles["customers"], profiles["orders"]
    assert (customers["row_count"], customers["primary_key"]) == (3, "customer_id")
    assert (orders["row_count"], orders["primary_key"], orders["time_column"]) == (5, "order_id", "ordered_at")
    assert orders["foreign_keys"] == [
        {"column": "customer_id", "references_table": "customers", "references_column": "customer_id", "inferred": True}
    ]
    assert customers["foreign_keys"] == []
    assert {c["name"]: c["type"] for c in orders["columns"]} == {
        "order_id": "VARCHAR",
        "customer_id": "VARCHAR",
        "ordered_at": "TIMESTAMP",
        "net_amount": "DOUBLE",
    }
    assert all(c["nullable"] is False and c["description"] == "" for c in orders["columns"])
    assert (orders["origin_file"], orders["description"]) == ("orders.csv", "")
    for info in profiles.values():
        assert list(TABLE_INFO.iter_errors(info)) == []


def test_declared_keys_and_descriptions_override_profiling(retail: Path):
    declarations = {
        "orders": {
            "description": "One row per order.",
            "primary_key": "customer_id",
            "time_column": "ordered_at",
            "foreign_keys": [{"column": "customer_id", "references": "customers.customer_id"}],
            "columns": {"net_amount": "Order value after discounts, in USD."},
        },
        "customers": {"time_column": "joined_at", "primary_key": "tier"},
    }

    profiles = {info["name"]: info for info in tables.profile_database(retail, declarations)}

    orders = profiles["orders"]
    assert orders["description"] == "One row per order."
    assert orders["primary_key"] == "customer_id"
    assert orders["foreign_keys"] == [
        {
            "column": "customer_id",
            "references_table": "customers",
            "references_column": "customer_id",
            "inferred": False,
        }
    ]
    assert {c["name"]: c["description"] for c in orders["columns"]}[
        "net_amount"
    ] == "Order value after discounts, in USD."
    assert (profiles["customers"]["primary_key"], profiles["customers"]["time_column"]) == ("tier", "joined_at")


def test_an_awkward_workbook_loads_one_table_per_sheet_with_clean_columns(db: Path, awkward_xlsx: Path):
    progress: list[tuple[int, int]] = []

    result = tables.load_table_file(
        db, awkward_xlsx, "awkward.xlsx", set(), on_progress=lambda i, n: progress.append((i, n))
    )

    assert result.tables == ["awkward_q1_sales", "awkward_q2_sales"]
    assert result.parser == "duckdb-xlsx"
    assert progress == [(1, 2), (2, 2)]
    for table in result.tables:
        assert columns(db, table) == ["store", "net_sales", "net_sales_2", "region"]
    profiles = {info["name"]: info for info in tables.profile_database(db, None)}
    assert profiles["awkward_q2_sales"]["row_count"] == 1
    assert profiles["awkward_q1_sales"]["origin_file"] == "awkward.xlsx › Q1 Sales"


def test_a_single_sheet_workbook_is_named_after_the_file_alone(db: Path, tmp_path: Path):
    workbook = openpyxl.Workbook()
    workbook.active.append(["Customer ID", "2024 Revenue", ""])
    workbook.active.append(["C1", 10, "x"])
    path = tmp_path / "Customer List (final).xlsx"
    workbook.save(path)

    result = tables.load_table_file(db, path, path.name, set())

    assert result.tables == ["customer_list_final"]
    assert columns(db, "customer_list_final") == ["customer_id", "col_2024_revenue", "col"]


def test_an_empty_sheet_fails_alone(db: Path, awkward_xlsx: Path):
    workbook = openpyxl.load_workbook(awkward_xlsx)
    workbook.create_sheet("Notes")
    workbook.create_sheet("Header only").append(["a", "b"])
    workbook.save(awkward_xlsx)

    result = tables.load_table_file(db, awkward_xlsx, "awkward.xlsx", set())

    assert result.tables == ["awkward_q1_sales", "awkward_q2_sales"]
    assert len(result.warnings) == 2
    assert "Notes" in result.warnings[0] and "Header only" in result.warnings[1]


def test_table_names_stay_unique_within_the_database(db: Path):
    first = tables.load_table_file(db, TABLES / "orders.csv", "orders.csv", set())
    second = tables.load_table_file(db, TABLES / "orders.csv", "Orders.csv", set(first.tables))
    third = tables.load_table_file(db, TABLES / "orders.csv", "orders.csv", {"orders", "orders_2"})

    assert (first.tables, second.tables, third.tables) == (["orders"], ["orders_2"], ["orders_3"])


@pytest.mark.parametrize("name", ["empty.csv"])
def test_a_zero_byte_file_fails_with_empty_file(db: Path, name: str):
    with pytest.raises(IngestError) as error:
        tables.load_table_file(db, FIXTURES / name, name, set())

    assert error.value.code == "empty_file"


def test_a_csv_with_a_header_and_no_rows_fails_with_empty_file(db: Path, tmp_path: Path):
    path = tmp_path / "header.csv"
    path.write_text("a,b\n")

    with pytest.raises(IngestError) as error:
        tables.load_table_file(db, path, path.name, set())

    assert error.value.code == "empty_file"
    assert not db.exists() or "header" not in {info["name"] for info in tables.profile_database(db, None)}


def test_a_latin_1_csv_still_loads(db: Path, tmp_path: Path):
    path = tmp_path / "cities.csv"
    path.write_bytes("name,city\nJosé,Zürich\n".encode("latin-1"))

    result = tables.load_table_file(db, path, path.name, set())

    with duckdb.connect(str(db), read_only=True) as connection:
        assert connection.execute(f'SELECT * FROM "{result.tables[0]}"').fetchall() == [("José", "Zürich")]


def test_tsv_parquet_and_json_files_load(db: Path, tmp_path: Path):
    (tmp_path / "a.tsv").write_text("x\ty\n1\t2\n")
    (tmp_path / "b.jsonl").write_text('{"x": 1}\n{"x": 2}\n')
    (tmp_path / "c.json").write_text('[{"x": 1, "y": "a"}]')
    with duckdb.connect() as connection:
        connection.execute(f"COPY (SELECT 1 AS x) TO '{tmp_path / 'd.parquet'}' (FORMAT parquet)")

    parsers = {}
    for name in ("a.tsv", "b.jsonl", "c.json", "d.parquet"):
        result = tables.load_table_file(db, tmp_path / name, name, set())
        parsers[result.tables[0]] = result.parser

    assert parsers == {"a": "duckdb-csv", "b": "duckdb-json", "c": "duckdb-json", "d": "duckdb-parquet"}
    assert columns(db, "a") == ["x", "y"]


def test_drop_tables_removes_them(retail: Path):
    tables.drop_tables(retail, ["orders", "missing"])

    assert [info["name"] for info in tables.profile_database(retail, None)] == ["customers"]


def test_snake_case():
    assert tables.snake_case("Store #") == "store"
    assert tables.snake_case("Net Sales ($)") == "net_sales"
    assert tables.snake_case("Région") == "region"
    assert tables.snake_case("  ") == ""
    assert tables.snake_case("camelCaseName") == "camel_case_name"


def test_a_sheet_whose_column_mixes_types_keeps_every_value(db: Path, tmp_path: Path):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["Order ID", "Amount", "Shipped"])
    sheet.append([1, 10.5, datetime.date(2026, 1, 2)])
    sheet.append([2, "TBD", datetime.date(2026, 1, 3)])
    path = tmp_path / "mixed.xlsx"
    workbook.save(path)

    result = tables.load_table_file(db, path, path.name, set())

    [info] = tables.profile_database(db, None)
    types = {c["name"]: c["type"] for c in info["columns"]}
    # openpyxl reads date cells as datetimes, so the fallback types them TIMESTAMP
    assert types == {"order_id": "BIGINT", "amount": "VARCHAR", "shipped": "TIMESTAMP"}
    assert (info["row_count"], info["primary_key"], info["time_column"]) == (2, "order_id", "shipped")
    assert result.warnings == []
